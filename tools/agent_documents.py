"""Bounded source inspection and explicit publication of selected document tables.

Uploads enter through trusted application code. Model tools accept source IDs or
public URLs, never local paths. Extracted text is untrusted source data.
"""

from __future__ import annotations

import csv
from concurrent.futures import ThreadPoolExecutor, TimeoutError as FutureTimeout
import calendar
from datetime import date
from decimal import Decimal, InvalidOperation
import hashlib
from html.parser import HTMLParser
import http.client
import io
import ipaddress
import json
import mimetypes
import os
from pathlib import Path
import re
import socket
import ssl
import tempfile
import threading
import time
import unicodedata
from urllib import error, parse, request
import xml.etree.ElementTree as ET
import zipfile

from tools.lakehouse_store import StoreError


_DNS_WORKERS = ThreadPoolExecutor(max_workers=4, thread_name_prefix="document-dns")
_DNS_SLOTS = threading.BoundedSemaphore(4)


class DocumentError(ValueError):
    def __init__(self, message, code="DOCUMENT_ERROR"):
        super().__init__(message)
        self.code = code


def _canonical(value):
    return json.dumps(value, sort_keys=True, ensure_ascii=False, allow_nan=False, separators=(",", ":")).encode()


def _write_json(path, value):
    fd, name = tempfile.mkstemp(dir=path.parent, suffix=".tmp")
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(_canonical(value))
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(name, path)
    finally:
        Path(name).unlink(missing_ok=True)


def _public_destination(url, *, timeout=5):
    if not isinstance(url, str) or len(url) > 4096 or any(ord(char) < 32 for char in url):
        raise DocumentError("Invalid public URL.", "UNSAFE_URL")
    parsed = parse.urlsplit(url)
    if parsed.scheme not in {"https", "http"} or not parsed.hostname or parsed.username or parsed.password:
        raise DocumentError("Only public HTTP(S) URLs without credentials are allowed.", "UNSAFE_URL")
    try:
        port = parsed.port or (443 if parsed.scheme == "https" else 80)
        if port not in {80, 443}:
            raise DocumentError("Only ports 80 and 443 are allowed.", "UNSAFE_URL")
        if not _DNS_SLOTS.acquire(blocking=False):
            raise DocumentError("Source resolver is busy; retry within the request budget.", "FETCH_BUSY")
        try:
            pending = _DNS_WORKERS.submit(socket.getaddrinfo, parsed.hostname, port, type=socket.SOCK_STREAM)
        except Exception:
            _DNS_SLOTS.release()
            raise
        pending.add_done_callback(lambda future: _DNS_SLOTS.release())
        try:
            addresses = pending.result(timeout=timeout)
        except FutureTimeout as exc:
            raise DocumentError("Source DNS resolution exceeded its time budget.", "FETCH_TIMEOUT") from exc
    except DocumentError:
        raise
    except (ValueError, OSError) as exc:
        raise DocumentError("URL host could not be safely resolved.", "UNSAFE_URL") from exc
    resolved = []
    for entry in addresses:
        address = ipaddress.ip_address(entry[4][0])
        if not address.is_global or isinstance(address, ipaddress.IPv6Address) and address.ipv4_mapped and not address.ipv4_mapped.is_global:
            raise DocumentError("Private, loopback, reserved and local destinations are blocked.", "UNSAFE_URL")
        resolved.append(str(address))
    if not resolved:
        raise DocumentError("URL host has no public address.", "UNSAFE_URL")
    return parsed, resolved[0], port


class _NoRedirect(request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        return None


def fetch_public_url(url, *, max_bytes=16 * 1024**2, timeout=20, max_redirects=3):
    """Revalidate each hop and pin the validated IP for the actual connection."""
    deadline = time.monotonic() + timeout
    current = url
    for hop in range(max_redirects + 1):
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            raise DocumentError("Download time budget exceeded.", "FETCH_TIMEOUT")
        parsed, address, port = _public_destination(current, timeout=min(remaining, 5))
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            raise DocumentError("Download time budget exceeded.", "FETCH_TIMEOUT")

        class PinnedHTTP(http.client.HTTPConnection):
            def connect(self):
                self.sock = socket.create_connection((address, port), timeout=min(remaining, 5))

        class PinnedHTTPS(http.client.HTTPSConnection):
            def connect(self):
                raw = socket.create_connection((address, port), timeout=min(remaining, 5))
                try:
                    self.sock = ssl.create_default_context().wrap_socket(raw, server_hostname=parsed.hostname)
                except Exception:
                    raw.close()
                    raise

        class HTTPHandler(request.HTTPHandler):
            def http_open(self, req):
                return self.do_open(PinnedHTTP, req)

        class HTTPSHandler(request.HTTPSHandler):
            def https_open(self, req):
                return self.do_open(PinnedHTTPS, req)

        opener = request.build_opener(request.ProxyHandler({}), _NoRedirect(), HTTPHandler(), HTTPSHandler())
        req = request.Request(current, headers={"User-Agent": "AgenticMinds-SourceReader/1", "Accept-Encoding": "identity"})
        try:
            response = opener.open(req, timeout=min(remaining, 5))
        except error.HTTPError as exc:
            try:
                if exc.code in {301, 302, 303, 307, 308} and exc.headers.get("Location") and hop < max_redirects:
                    current = parse.urljoin(current, exc.headers["Location"])
                    continue
                raise DocumentError(f"Source returned HTTP {exc.code}.", "FETCH_FAILED") from exc
            finally:
                exc.close()
        except (OSError, error.URLError) as exc:
            raise DocumentError("Source download failed or timed out.", "FETCH_FAILED") from exc
        with response:
            declared = response.headers.get("Content-Length")
            if declared and (not declared.isdecimal() or int(declared) > max_bytes):
                raise DocumentError("Source exceeds the download size limit.", "SOURCE_TOO_LARGE")
            blocks, total = [], 0
            while True:
                if time.monotonic() >= deadline:
                    raise DocumentError("Download time budget exceeded.", "FETCH_TIMEOUT")
                chunk = response.read(min(65536, max_bytes - total + 1))
                if not chunk:
                    break
                total += len(chunk)
                if total > max_bytes:
                    raise DocumentError("Source exceeds the download size limit.", "SOURCE_TOO_LARGE")
                blocks.append(chunk)
            return b"".join(blocks), response.headers.get_content_type(), current
    raise DocumentError("Too many redirects.", "FETCH_FAILED")


class _HTMLTables(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables, self.text, self.current, self.row, self.cell = [], [], None, None, None
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.hidden += 1
        if tag == "table":
            if self.current is not None:
                raise DocumentError("Nested HTML tables require explicit review.", "AMBIGUOUS_TABLE")
            self.current = []
        elif tag == "tr" and self.current is not None:
            self.row = []
        elif tag in {"td", "th"} and self.row is not None:
            spans = dict(attrs)
            if spans.get("colspan", "1") != "1" or spans.get("rowspan", "1") != "1":
                raise DocumentError("HTML tables with merged cells need explicit review.", "AMBIGUOUS_TABLE")
            self.cell = []

    def handle_data(self, data):
        if not self.hidden:
            self.text.append(data)
            if self.cell is not None:
                self.cell.append(data)

    def handle_endtag(self, tag):
        if tag in {"script", "style"}:
            self.hidden = max(0, self.hidden - 1)
        elif tag in {"td", "th"} and self.cell is not None:
            self.row.append(" ".join(self.cell).strip())
            self.cell = None
        elif tag == "tr" and self.row is not None:
            self.current.append(self.row)
            self.row = None
        elif tag == "table" and self.current is not None:
            self.tables.append(self.current)
            self.current = None


def _column_name(value, index, used):
    ascii_value = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode()
    base = re.sub(r"[^A-Za-z0-9_]+", "_", ascii_value.strip()).strip("_")[:100] or f"column_{index + 1}"
    if base[0].isdigit():
        base = "c_" + base
    name, suffix = base, 2
    while name in used:
        name, suffix = f"{base}_{suffix}", suffix + 1
    used.add(name)
    return name


class DocumentTools:
    def __init__(self, store, workspace_id, *, upload_root=None, ocr_callback=None, searxng_url=None,
                 max_source_bytes=16 * 1024**2, max_rows=5000, max_columns=64, max_pages=30):
        self.store, self.workspace_id = store, workspace_id
        self.store.workspace(workspace_id)
        self.root = store.root / "document_sources" / workspace_id
        self.root.mkdir(parents=True, exist_ok=True)
        self.upload_root = Path(upload_root or store.root / "uploads").resolve()
        self.upload_root.mkdir(parents=True, exist_ok=True)
        self.ocr_callback, self.searxng_url = ocr_callback, searxng_url
        self.max_source_bytes, self.max_rows = max_source_bytes, max_rows
        self.max_columns, self.max_pages = max_columns, max_pages

    def _directory(self, source_id):
        if not isinstance(source_id, str) or not re.fullmatch(r"source_[a-f0-9]{64}", source_id):
            raise DocumentError("Invalid source identifier.")
        directory = self.root / source_id
        if directory.is_symlink() or not directory.resolve().is_relative_to(self.root.resolve()):
            raise DocumentError("Source path is invalid.")
        return directory

    def source(self, source_id):
        directory = self._directory(source_id)
        manifest = json.loads((directory / "manifest.json").read_text())
        if hashlib.sha256((directory / "raw.bin").read_bytes()).hexdigest() != manifest["raw_sha256"]:
            raise DocumentError("Raw source hash mismatch.", "SOURCE_HASH_MISMATCH")
        return manifest

    def list_sources(self):
        return [self.source(path.name) for path in sorted(self.root.glob("source_*")) if path.is_dir()]

    def raw_source_bytes(self, source_id):
        """Trusted UI raw-download helper, with no caller-supplied filesystem path."""
        manifest = self.source(source_id)
        data = (self._directory(source_id) / "raw.bin").read_bytes()
        if hashlib.sha256(data).hexdigest() != manifest["raw_sha256"]:
            raise DocumentError("Raw source hash mismatch.", "SOURCE_HASH_MISMATCH")
        return data

    def review_candidate(self, source_id, table_id):
        """Full bounded candidate for trusted UI review, never a model tool."""
        manifest = self.source(source_id)
        self.inspect_source(source_id=source_id)
        inspection = json.loads((self._directory(source_id) / "inspection.json").read_text())
        table = next((item for item in inspection["tables"] if item["table_id"] == table_id), None)
        if table is None:
            raise DocumentError("Unknown candidate table.")
        return {"status": "ok", "source_id": source_id, "raw_sha256": manifest["raw_sha256"],
                **self._reviewed_table(source_id, table), "preview_truncated": False}

    def _register(self, data, filename, mime_type, source_url=None):
        if not data or len(data) > self.max_source_bytes:
            raise DocumentError("Source is empty or exceeds its byte limit.", "SOURCE_TOO_LARGE")
        filename = Path(filename).name[:200]
        metadata = {"filename": filename, "mime_type": mime_type, "source_url": source_url,
                    "raw_sha256": hashlib.sha256(data).hexdigest(), "size_bytes": len(data),
                    "workspace_id": self.workspace_id}
        source_id = "source_" + hashlib.sha256(_canonical(metadata)).hexdigest()
        directory = self._directory(source_id)
        directory.mkdir(exist_ok=True)
        manifest = {**metadata, "source_id": source_id, "artifact_ref": source_id,
                    "source_namespace": "external:" + (parse.urlsplit(source_url).hostname if source_url else "upload") + ":" + source_id,
                    "status": "registered", "content_is_untrusted_data": True}
        if not (directory / "manifest.json").exists():
            fd, name = tempfile.mkstemp(dir=directory)
            try:
                with os.fdopen(fd, "wb") as handle:
                    handle.write(data)
                os.replace(name, directory / "raw.bin")
            finally:
                Path(name).unlink(missing_ok=True)
            _write_json(directory / "manifest.json", manifest)
        return self.source(source_id)

    def register_upload(self, path, filename=None):
        """Trusted server entry point; deliberately absent from model tool schemas."""
        source = Path(path).resolve(strict=True)
        if not source.is_relative_to(self.upload_root) or not source.is_file():
            raise DocumentError("Upload is outside the server upload directory.", "UNSAFE_UPLOAD")
        if source.stat().st_size > self.max_source_bytes:
            raise DocumentError("Upload exceeds its byte limit.", "SOURCE_TOO_LARGE")
        with source.open("rb") as handle:
            data = handle.read(self.max_source_bytes + 1)
        name = filename or source.name
        return self._register(data, name, mimetypes.guess_type(name)[0] or "application/octet-stream")

    def _table(self, rows, *, page=None, sheet=None, units=None, origin="parsed"):
        if not rows or len(rows) < 2:
            return None
        if len(rows) - 1 > self.max_rows or not 1 <= len(rows[0]) <= self.max_columns:
            raise DocumentError("Candidate table exceeds row or column limits.", "TABLE_TOO_LARGE")
        width, used = len(rows[0]), set()
        if any(len(row) != width for row in rows):
            raise DocumentError("Candidate contains inconsistent row widths.", "AMBIGUOUS_TABLE")
        original = [str(value or "").strip() for value in rows[0]]
        columns = [_column_name(value, index, used) for index, value in enumerate(original)]
        values = [[None if cell is None or str(cell).strip() == "" else str(cell).strip() for cell in row] for row in rows[1:]]
        if any(len(cell) > 20000 for row in values for cell in row if cell is not None):
            raise DocumentError("A candidate cell exceeds its text limit.", "TABLE_TOO_LARGE")
        return {"columns": columns, "original_columns": dict(zip(columns, original)), "rows": values,
                "row_count": len(values), "page": page, "sheet": sheet, "units": units or {}, "origin": origin,
                "source_header_quotes": dict(zip(columns, original))}

    def _extract_image(self, data, mime_type, *, page=1):
        from PIL import Image
        if len(data) > 8 * 1024**2:
            raise DocumentError("Image exceeds the 8 MiB extraction limit.", "SOURCE_TOO_LARGE")
        with Image.open(io.BytesIO(data)) as image:
            if image.width * image.height > 20_000_000:
                raise DocumentError("Image exceeds its pixel limit.", "SOURCE_TOO_LARGE")
            image.verify()
        if self.ocr_callback is None:
            raise DocumentError("Image extraction is not configured.", "OCR_UNAVAILABLE")
        extracted = self.ocr_callback(data, mime_type)
        if not isinstance(extracted, dict) or not isinstance(extracted.get("text", ""), str):
            raise DocumentError("Image extractor returned an invalid envelope.", "OCR_INVALID_OUTPUT")
        if extracted.get("finish_reason") == "length" or extracted.get("status", "ok") != "ok":
            raise DocumentError("Image extraction was incomplete or failed.", "OCR_INVALID_OUTPUT")
        if not isinstance(extracted.get("tables", []), list):
            raise DocumentError("Image extractor tables must be a list.", "OCR_INVALID_OUTPUT")
        tables = []
        for table in extracted.get("tables", [])[:10]:
            rows, columns = table.get("rows", []), table.get("columns", [])
            if rows and isinstance(rows[0], dict):
                rows = [[row.get(column) for column in columns] for row in rows]
            tables.append(self._table([columns, *rows], page=page, units=table.get("units"), origin="ocr"))
        return extracted.get("text", ""), tables

    def _parse(self, manifest, data):
        suffix, mime = Path(manifest["filename"]).suffix.lower(), manifest["mime_type"]
        tables, pages, text, warnings = [], [], "", []
        if suffix == ".csv" or mime == "text/csv":
            text = data.decode("utf-8-sig")
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(io.StringIO(text), dialect=dialect, strict=True)
            rows = []
            for row in reader:
                rows.append(row)
                if len(rows) > self.max_rows + 1:
                    raise DocumentError("CSV exceeds its row limit.", "TABLE_TOO_LARGE")
            tables.append(self._table(rows))
        elif suffix == ".xlsx" or mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist()) > 10000 or sum(item.file_size for item in archive.infolist()) > 64 * 1024**2:
                    raise DocumentError("Expanded workbook exceeds its size limit.", "SOURCE_TOO_LARGE")
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                for sheet in workbook.worksheets[:10]:
                    if sheet.max_row > self.max_rows + 1 or sheet.max_column > self.max_columns:
                        raise DocumentError("Workbook sheet exceeds table limits.", "TABLE_TOO_LARGE")
                    rows = list(sheet.iter_rows(values_only=True))
                    if any(isinstance(value, str) and value.startswith("=") for row in rows for value in row):
                        warnings.append({"code": "FORMULAS_REQUIRE_VALUES", "sheet": sheet.title})
                        continue
                    tables.append(self._table(rows, sheet=sheet.title))
                    text += "\n" + "\n".join(" | ".join(str(value or "") for value in row) for row in rows)
            finally:
                workbook.close()
        elif suffix in {".html", ".htm"} or mime == "text/html":
            parser = _HTMLTables()
            parser.feed(data.decode("utf-8-sig"))
            text = " ".join(parser.text)
            tables = [self._table(rows) for rows in parser.tables[:10]]
        elif suffix == ".pdf" or mime == "application/pdf":
            import pdfplumber
            image_pages, deferred_pages = 0, []
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                if len(pdf.pages) > self.max_pages:
                    raise DocumentError("PDF exceeds the page limit.", "PAGE_LIMIT")
                for number, page in enumerate(pdf.pages, 1):
                    page_text = (page.extract_text() or "")[:30000]
                    page_method = "pdf_text"
                    if not page_text.strip() and self.ocr_callback is not None and image_pages < 3:
                        if (page.width * 110 / 72) * (page.height * 110 / 72) > 20_000_000:
                            raise DocumentError("Rendered PDF page exceeds its pixel limit.", "SOURCE_TOO_LARGE")
                        buffer = io.BytesIO()
                        page.to_image(resolution=110).original.convert("RGB").save(buffer, format="PNG")
                        page_text, extracted_tables = self._extract_image(buffer.getvalue(), "image/png", page=number)
                        tables.extend(extracted_tables)
                        image_pages += 1
                        page_method = "ocr"
                    elif not page_text.strip():
                        deferred_pages.append(number)
                    pages.append({"page": number, "text": page_text[:30000], "extraction_method": page_method})
                    text += "\n" + page_text
                    if page_method == "pdf_text":
                        for rows in page.extract_tables()[:10]:
                            tables.append(self._table(rows, page=number))
            if image_pages:
                warnings.append({"code": "OCR_REVIEW_REQUIRED", "pages_processed": image_pages,
                                 "message": "Rendered PDF image cells require independent review before publication."})
            if deferred_pages:
                warnings.append({"code": "OCR_PAGE_LIMIT" if self.ocr_callback else "SCANNED_PDF_NEEDS_OCR",
                                 "pages": deferred_pages, "message": "Some image-only pages were not extracted; at most three OCR pages are processed per document."})
        elif suffix in {".png", ".jpg", ".jpeg"} or mime in {"image/png", "image/jpeg"}:
            text, tables = self._extract_image(data, "image/png" if data.startswith(b"\x89PNG") else "image/jpeg")
            pages = [{"page": 1, "text": text[:30000]}]
            warnings.append({"code": "OCR_REVIEW_REQUIRED", "message": "Image cells require independent numeric verification before publication."})
        elif suffix == ".txt" or mime == "text/plain":
            text = data.decode("utf-8-sig")
        else:
            raise DocumentError("Supported sources are CSV, XLSX, HTML, PDF, UTF-8 text, PNG and JPEG.", "UNSUPPORTED_FORMAT")
        tables = [table for table in tables if table is not None]
        if len(tables) > 10:
            raise DocumentError("Document exceeds the ten-table inspection limit.", "TABLE_LIMIT")
        for index, table in enumerate(tables, 1):
            table["table_id"] = f"table_{index:03d}"
        return {"text": text[:200000], "text_truncated": len(text) > 200000, "pages": pages,
                "tables": tables, "warnings": warnings}

    def inspect_source(self, source_id=None, url=None):
        if bool(source_id) == bool(url):
            raise DocumentError("Provide exactly one source_id or public URL.")
        if url:
            data, mime, final_url = fetch_public_url(url, max_bytes=self.max_source_bytes)
            name = Path(parse.unquote(parse.urlsplit(final_url).path)).name or "source"
            manifest = self._register(data, name, mime, final_url)
            source_id = manifest["source_id"]
        manifest = self.source(source_id)
        directory = self._directory(source_id)
        cache = directory / "inspection.json"
        if cache.exists():
            inspection = json.loads(cache.read_text())
        else:
            inspection = self._parse(manifest, (directory / "raw.bin").read_bytes())
            _write_json(cache, inspection)
        inspection["tables"] = [self._reviewed_table(source_id, table) for table in inspection["tables"]]
        previews = [{**{key: value for key, value in table.items() if key != "rows"},
                     "source_header_quotes": table.get("source_header_quotes", table["original_columns"]),
                     "preview": [dict(zip(table["columns"], row)) for row in table["rows"][:8]],
                     "preview_truncated": table["row_count"] > 8} for table in inspection["tables"]]
        return {**manifest, "status": "ok", "tables": previews, "text": inspection["text"][:12000],
                "text_truncated": inspection["text_truncated"] or len(inspection["text"]) > 12000,
                "pages": [{**page, "text": page["text"][:2000]} for page in inspection["pages"]],
                "warnings": inspection["warnings"], "publication_requires_explicit_contract": True}

    def _reviewed_table(self, source_id, table):
        review_path = self._directory(source_id) / (table["table_id"] + "_review.json")
        if not review_path.exists():
            return table
        review = json.loads(review_path.read_text())
        content = {key: value for key, value in review.items() if key != "review_sha256"}
        if hashlib.sha256(_canonical(content)).hexdigest() != review.get("review_sha256") or review["raw_sha256"] != self.source(source_id)["raw_sha256"]:
            raise DocumentError("Reviewed table hash does not match its source.", "SOURCE_HASH_MISMATCH")
        return {**table, "rows": review["rows"], "row_count": len(review["rows"]),
                "review": {key: value for key, value in review.items() if key != "rows"}}

    def review_table(self, source_id, table_id, reviewed_rows, unit_evidence):
        """Trusted UI-only human review; never registered as a model tool."""
        manifest = self.source(source_id)
        self.inspect_source(source_id=source_id)
        inspection = json.loads((self._directory(source_id) / "inspection.json").read_text())
        table = next((item for item in inspection["tables"] if item["table_id"] == table_id), None)
        if table is None or not isinstance(reviewed_rows, list) or not 1 <= len(reviewed_rows) <= self.max_rows:
            raise DocumentError("Select a known table and provide reviewed rows within the row limit.")
        if not isinstance(unit_evidence, dict) or any(column not in table["columns"] or not isinstance(quote, str) or not 1 <= len(quote) <= 1000 for column, quote in unit_evidence.items()):
            raise DocumentError("Reviewed unit declarations must name candidate columns and bounded source quotes or explicit user declarations.")
        rows = []
        for row in reviewed_rows:
            if isinstance(row, dict):
                if set(row) != set(table["columns"]):
                    raise DocumentError("Every reviewed row must contain exactly the candidate columns.")
                row = [row[column] for column in table["columns"]]
            if not isinstance(row, (list, tuple)) or len(row) != len(table["columns"]):
                raise DocumentError("Reviewed row width must equal the candidate column count.")
            if any(value is not None and (isinstance(value, (dict, list, bool)) or len(str(value)) > 20000) for value in row):
                raise DocumentError("Reviewed cells must be bounded scalar values.")
            rows.append([None if value is None or str(value).strip() == "" else str(value).strip() for value in row])
        review = {"source_id": source_id, "table_id": table_id, "raw_sha256": manifest["raw_sha256"],
                  "rows": rows, "unit_evidence": unit_evidence, "verification": "explicit_user_cell_and_unit_review"}
        review["review_sha256"] = hashlib.sha256(_canonical(review)).hexdigest()
        _write_json(self._directory(source_id) / (table_id + "_review.json"), review)
        return {"status": "ok", "source_id": source_id, "table_id": table_id, "row_count": len(rows),
                "review_sha256": review["review_sha256"], "verification": review["verification"]}

    @staticmethod
    def _unit_supported(unit, quote):
        lowered = unicodedata.normalize("NFKC", quote).casefold()
        aliases = {"TRY": ["tl", "try", "türk lirası"], "TRY_million": ["milyon tl", "million tl", "million try"],
                   "TRY_billion": ["milyar tl", "billion tl", "billion try"], "USD": ["usd", "dollar", "dolar"],
                   "percent": ["%", "yüzde", "percent"], "persons": ["kişi", "persons", "people"],
                   "visits": ["ziyaret", "visits"], "count": ["adet", "count", "sayı"]}
        if unit == "TRY" and any(word in lowered for word in ["milyon", "million", "milyar", "billion", "bin tl", "thousand"]):
            return False
        return any(re.search(r"(?<!\w)" + re.escape(value) + r"(?!\w)", lowered) for value in aliases.get(unit, [unit.casefold()]))

    def _publication_key(self, arguments):
        content = {"workspace_id": self.workspace_id, **arguments}
        source_id, table_id = arguments.get("source_id", ""), arguments.get("table_id", "")
        if re.fullmatch(r"source_[a-f0-9]{64}", source_id) and re.fullmatch(r"table_\d{3}", table_id):
            review = self._directory(source_id) / (table_id + "_review.json")
            if review.exists():
                content["review_sha256"] = json.loads(review.read_text())["review_sha256"]
        return hashlib.sha256(_canonical(content)).hexdigest()

    @staticmethod
    def _normalize_source_date(value, frequency, source_format):
        """Relabel a declared native period boundary without aggregating observations."""
        if source_format not in {"native", "iso_period_start", "iso_period_end"}:
            raise DocumentError("source_date_format must be native, iso_period_start or iso_period_end.", "INVALID_SOURCE_DATE_FORMAT")
        coarse = frequency in {"monthly", "quarterly", "annual"}
        if source_format != "native" and not coarse:
            raise DocumentError("ISO period-boundary normalization applies only to declared monthly, quarterly or annual data.", "INVALID_SOURCE_DATE_FORMAT")
        if value is None:
            return value
        is_iso = bool(re.fullmatch(r"\d{4}-\d{2}-\d{2}", value))
        if source_format == "native":
            if coarse and is_iso:
                raise DocumentError("Source dates use YYYY-MM-DD for declared " + frequency +
                    " data. Set source_date_format=iso_period_start or iso_period_end only if these dates label native period boundaries; this never aggregates daily observations.",
                    "SOURCE_DATE_FORMAT_REQUIRED")
            return value
        if not is_iso:
            raise DocumentError("The declared source_date_format requires every nonmissing date to use YYYY-MM-DD.", "INVALID_SOURCE_DATE_FORMAT")
        try:
            observed = date.fromisoformat(value)
        except ValueError as exc:
            raise DocumentError("Source contains an invalid calendar date.", "INVALID_SOURCE_DATE_FORMAT") from exc
        if frequency == "monthly":
            first = date(observed.year, observed.month, 1)
            last = date(observed.year, observed.month, calendar.monthrange(observed.year, observed.month)[1])
            label = f"{observed.year:04d}-{observed.month:02d}"
        elif frequency == "quarterly":
            quarter = (observed.month - 1) // 3 + 1
            first_month, last_month = (quarter - 1) * 3 + 1, quarter * 3
            first = date(observed.year, first_month, 1)
            last = date(observed.year, last_month, calendar.monthrange(observed.year, last_month)[1])
            label = f"{observed.year:04d}-Q{quarter}"
        else:
            first, last = date(observed.year, 1, 1), date(observed.year, 12, 31)
            label = f"{observed.year:04d}"
        expected = first if source_format == "iso_period_start" else last
        if observed != expected:
            raise DocumentError(f"Source date {value} is not the declared {source_format} boundary for {frequency} data; daily observations cannot be converted by relabeling.",
                                "NON_BOUNDARY_SOURCE_DATE")
        return label

    def publish_selected_table(self, source_id, table_id, contract, expected_version, column_mapping=None, unit_evidence=None):
        contract = self.store._validate_contract(contract)
        manifest = self.source(source_id)
        self.inspect_source(source_id=source_id)
        inspection = json.loads((self._directory(source_id) / "inspection.json").read_text())
        table = next((item for item in inspection["tables"] if item["table_id"] == table_id), None)
        if table is None:
            raise DocumentError("Unknown candidate table.")
        table = self._reviewed_table(source_id, table)
        if table["origin"] == "ocr" and "review" not in table:
            raise DocumentError("OCR candidates need independent cell verification before publication.", "OCR_REVIEW_REQUIRED")
        arguments = {"source_id": source_id, "table_id": table_id, "contract": contract,
                     "expected_version": expected_version, "column_mapping": column_mapping, "unit_evidence": unit_evidence}
        key = self._publication_key(arguments)
        mapping = column_mapping or {column: column for column in table["columns"]}
        if not isinstance(mapping, dict) or set(mapping) != set(table["columns"]) or len(set(mapping.values())) != len(mapping):
            raise DocumentError("Map every candidate column exactly once.")
        contract = json.loads(_canonical(contract))
        if set(contract.get("columns", {})) != set(mapping.values()):
            raise DocumentError("Contract columns must exactly match the mapped candidate columns.")
        reviewed_evidence = {mapping[column]: quote for column, quote in table.get("review", {}).get("unit_evidence", {}).items()}
        evidence = unit_evidence or reviewed_evidence
        for column, spec in contract["columns"].items():
            if spec.get("dtype") not in {"integer", "float"}:
                continue
            quote = evidence.get(column, "")
            unit = spec.get("unit", "")
            if unit == "TRY" and spec.get("scale", 1) in {1000000, 1000000000}:
                unit = "TRY_million" if spec["scale"] == 1000000 else "TRY_billion"
            explicitly_reviewed = quote == reviewed_evidence.get(column) and column in reviewed_evidence
            if not isinstance(quote, str) or not quote.strip() or not (quote in inspection["text"] or explicitly_reviewed) or not self._unit_supported(unit, quote):
                original_column = next(name for name, mapped in mapping.items() if mapped == column)
                original_header = table["original_columns"].get(original_column, "")
                raise DocumentError("Numeric unit needs an exact source quote consistent with the declared unit: " + column +
                    ". Copy only a verbatim source substring, without explanatory additions. Candidate header: " + repr(original_header), "UNITS_REVIEW_REQUIRED")
        number_format = contract.get("number_format", "decimal_dot")
        if number_format not in {"decimal_dot", "decimal_comma"}:
            raise DocumentError("Choose decimal_dot or decimal_comma numeric parsing explicitly.")
        source_date_format = contract.get("source_date_format", "native")
        self._normalize_source_date(None, contract["frequency"], source_date_format)
        rows, date_proof = [], []
        for row_index, row in enumerate(table["rows"], 1):
            output = []
            for original, value in zip(table["columns"], row):
                spec = contract["columns"][mapping[original]]
                if mapping[original] == contract.get("date_column"):
                    original_value = value
                    value = self._normalize_source_date(value, contract["frequency"], source_date_format)
                    if value != original_value:
                        date_proof.append({"candidate_row": row_index, "original": original_value, "normalized": value})
                if value is not None and spec.get("dtype") in {"integer", "float"}:
                    if number_format == "decimal_comma":
                        if not re.fullmatch(r"[+-]?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?", value):
                            raise DocumentError("Numeric cell does not match decimal_comma format.", "INVALID_NUMERIC_CELL")
                        value = value.replace(".", "").replace(",", ".")
                    try:
                        if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", value) or not Decimal(value).is_finite():
                            raise InvalidOperation
                    except InvalidOperation as exc:
                        raise DocumentError("Numeric cell needs explicit format or correction.", "INVALID_NUMERIC_CELL") from exc
                output.append(value if value is not None else "")
            rows.append(output)
        contract["source_namespace"] = manifest["source_namespace"]
        contract["document_provenance"] = {"source_id": source_id, "table_id": table_id, "raw_sha256": manifest["raw_sha256"],
            "source_url": manifest["source_url"], "page": table["page"], "sheet": table["sheet"],
            "publication_key": key, "unit_evidence": evidence, "numeric_verification": "explicit_format_and_store_contract"}
        if table.get("review"):
            contract["document_provenance"]["human_review"] = table["review"]
        if date_proof:
            contract["document_provenance"]["date_normalization"] = {
                "operation": "normalize_native_period_label", "source_date_format": source_date_format,
                "declared_frequency": contract["frequency"], "date_column": contract["date_column"],
                "aggregation_performed": False, "rows": date_proof}
        fd, name = tempfile.mkstemp(dir=self.root, suffix=".csv")
        try:
            with os.fdopen(fd, "w", encoding="utf-8", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow([mapping[column] for column in table["columns"]])
                writer.writerows(rows)
            workspace = self.store.ingest_csv(self.workspace_id, name, contract, expected_version=expected_version)
        finally:
            Path(name).unlink(missing_ok=True)
        result = self._find_publication(key, workspace)
        _write_json(self.root / ("publication_" + key + ".json"), result)
        return result

    def _find_publication(self, key, workspace=None):
        workspace = workspace or self.store.workspace(self.workspace_id)
        for dataset_id in workspace["datasets"]:
            manifest = self.store.dataset_manifest(dataset_id)
            provenance = manifest.get("contract", {}).get("document_provenance", {})
            if provenance.get("publication_key") == key:
                self.store.overlay_path(dataset_id)
                return {"status": "ok", "dataset_id": dataset_id, "artifact_ref": dataset_id,
                        "source_id": provenance["source_id"], "table_id": provenance["table_id"],
                        "workspace_version": workspace["version"], "row_count": manifest["row_count"],
                        "source_namespace": manifest["contract"]["source_namespace"], "provenance": provenance}
        return {"status": "blocked", "code": "WRITE_OUTCOME_UNKNOWN", "message": "No committed publication was found; the write was not repeated."}

    def recover_publication(self, arguments, intent):
        canonical = {"column_mapping": None, "unit_evidence": None, **arguments}
        return self._find_publication(self._publication_key(canonical))

    def web_search(self, query, limit=5):
        if self.searxng_url is False:
            return {"status": "unavailable", "code": "WEB_SEARCH_UNCONFIGURED", "results": []}
        if not isinstance(query, str) or not 1 <= len(query) <= 500 or type(limit) is not int or not 1 <= limit <= 10:
            raise DocumentError("Search query and limit exceed their bounds.")
        backend = "SearXNG" if self.searxng_url else "Bing RSS"
        try:
            if self.searxng_url:
                url = self.searxng_url.rstrip("/") + "/search?" + parse.urlencode({"q": query, "format": "json"})
                raw, _, _ = fetch_public_url(url, max_bytes=1024**2, timeout=15)
                items = json.loads(raw).get("results", [])
            else:
                url = "https://www.bing.com/search?" + parse.urlencode({"format": "rss", "q": query})
                raw, _, _ = fetch_public_url(url, max_bytes=1024**2, timeout=15)
                if b"<!doctype" in raw.lower() or b"<!entity" in raw.lower():
                    raise DocumentError("Search response contains unsupported XML declarations.", "SEARCH_INVALID_RESPONSE")
                rss = ET.fromstring(raw)
                if rss.tag != "rss":
                    raise DocumentError("Search did not return RSS results.", "SEARCH_INVALID_RESPONSE")
                items = [{"title": item.findtext("title", ""), "url": item.findtext("link", ""),
                          "content": item.findtext("description", "")} for item in rss.findall("./channel/item")]
            results = []
            for item in items[:limit]:
                target = item.get("url", "")
                if parse.urlsplit(target).scheme in {"http", "https"}:
                    results.append({"title": str(item.get("title", ""))[:300], "url": target[:4096], "snippet": str(item.get("content", ""))[:1200]})
            return {"status": "ok", "query": query, "results": results, "source_backend": backend,
                    "content_is_untrusted_data": True, "sources_verified": False,
                    "next_step": "Inspect result URLs before relying on them as citation evidence."}
        except (DocumentError, ValueError, OSError, ET.ParseError) as exc:
            return {"status": "unavailable", "code": getattr(exc, "code", "SEARCH_INVALID_RESPONSE"),
                    "message": str(exc), "source_backend": backend, "results": []}

    def _contract_schema(self):
        """Describe the store contract without restricting its metadata extension fields."""
        nonblank = {"type": "string", "minLength": 1, "pattern": r"\S"}
        identifier = {"type": "string", "pattern": r"^[A-Za-z_][A-Za-z0-9_]{0,127}$"}
        kinds = ["stock", "flow", "count", "count_stock", "count_flow", "ratio", "rate", "index", "price", "dimension", "unknown"]
        column = {"type": "object", "properties": {
            "dtype": {"type": "string", "enum": ["string", "integer", "float", "boolean", "date"],
                      "description": "Stored type: date for the calendar key, integer/float for numeric metrics, string for labels; boolean cells must be true/false."},
            "unit": {**nonblank, "description": "Explicit measurement unit, e.g. TRY, TRY_million, visits, persons, percent, calendar or label. Numeric columns need matching unit_evidence."},
            "kind": {"anyOf": [{"type": "string", "enum": kinds}, nonblank],
                     "description": "stock/count_stock are point-in-time balances; flow/count_flow are amounts/counts over a period; count is a generic count; ratio/rate are ratios/rates; index is an index; price is a unit price; dimension is a key/label. Other nonempty kinds are stored but require semantic review before analytics."},
            "nullable": {"type": "boolean", "description": "Whether empty cells are allowed; key cells can never be null."},
            "scale": {"type": "number", "default": 1, "description": "Number of base units represented by a stored number, e.g. unit TRY with scale 1000000 for million TL. Do not double-apply a unit conversion."},
            "currency": {"type": ["string", "null"], "description": "Currency code for monetary metrics, e.g. TRY or USD; needed for monetary deflation."},
            "aggregation": {"type": "string", "description": "Declared aggregation semantics; describe native stock/flow meaning, never infer additivity across populations."},
            "index_role": {"type": ["string", "null"], "description": "price_deflator only when source evidence establishes a suitable price index; a production index is not an inflation deflator."},
            "deflator_currency": {"type": ["string", "null"], "description": "Currency whose monetary values this reviewed price deflator applies to."},
            "price_scope": {"type": ["string", "null"], "description": "Documented price scope such as consumer_prices."},
        }, "required": ["dtype", "unit", "kind", "nullable"], "additionalProperties": True}
        keys = {"type": "array", "items": identifier, "minItems": 1, "uniqueItems": True}
        return {"type": "object", "properties": {
            "name": {**nonblank, "description": "Human-readable dataset name; used in discovered metric titles."},
            "columns": {"type": "object", "propertyNames": identifier, "minProperties": 1,
                        "maxProperties": min(self.max_columns, self.store.limits["max_columns"]), "additionalProperties": column,
                        "description": "Exactly all mapped candidate columns; each definition requires dtype/unit/kind/nullable."},
            "key": {**keys, "description": "Existing columns forming a unique nonmissing row key, including date_column for time data."},
            "grain": {**keys, "description": "Exactly the same set of columns as key, declaring what one row represents."},
            "frequency": {"type": "string", "enum": ["daily", "business_daily", "weekly", "monthly", "quarterly", "annual", "event", "static"],
                          "description": "Native frequency: monthly dates YYYY-MM; quarterly YYYY-Qn; annual YYYY; daily/business_daily/weekly/event YYYY-MM-DD. Static data has no date axis. Storage accepts event/static; regular time-series analytics require a supported calendar frequency."},
            "date_column": {"type": ["string", "null"], "description": "Required for non-static data: an existing dtype=date column included in key and grain. Static data must omit it or use null."},
            "source_date_format": {"type": "string", "enum": ["native", "iso_period_start", "iso_period_end"], "default": "native",
                                   "description": "Use native when dates already have the declared frequency's labels. For monthly/quarterly/annual source rows labeled with YYYY-MM-DD, explicitly choose iso_period_start or iso_period_end only when every date is that native period's first/last calendar day. Values and row counts are preserved; this never resamples daily data. Examples: monthly 2025-01-01 -> 2025-01 with iso_period_start; quarterly 2025-03-31 -> 2025-Q1 with iso_period_end."},
            "expected_rows": {"type": "integer", "minimum": 1, "maximum": self.store.limits["max_rows"], "description": "Optional exact expected row count; use the full candidate count, never the preview count."},
            "expected_periods": {"type": "array", "items": {"type": "string"}, "minItems": 1, "uniqueItems": True,
                                 "description": "Optional exact native period labels required for every entity; not allowed for static data."},
            "number_format": {"type": "string", "enum": ["decimal_dot", "decimal_comma"], "default": "decimal_dot",
                              "description": "Default decimal_dot uses 1234.56 without thousands separators. Explicit decimal_comma accepts 1.234,56 or 1234,56. Never guess a separator convention."},
        }, "required": ["name", "columns", "key", "grain", "frequency"], "additionalProperties": True,
            "allOf": [{"if": {"properties": {"frequency": {"const": "static"}}, "required": ["frequency"]},
                       "then": {"properties": {"date_column": {"type": "null"}}, "not": {"required": ["expected_periods"]}},
                       "else": {"required": ["date_column"], "properties": {"date_column": identifier}}}]}

    def extra_tools(self):
        import jsonschema
        definitions = {
            "inspect_source": (self.inspect_source, "Inspect an uploaded source ID or public URL; returns untrusted source text and candidate tables.",
                {"source_id": {"type": "string"}, "url": {"type": "string"}}, []),
            "publish_selected_table": (self.publish_selected_table, "Publish one inspected table with explicit dtype/unit/kind/grain contract and source unit quotes. Example monthly columns: month={dtype:date,unit:calendar,kind:dimension,nullable:false}, visits={dtype:integer,unit:visits,kind:count_flow,nullable:false}; name=clinic_visits, key=[month], grain=[month], date_column=month, frequency=monthly. Use actual candidate column names or explicit column_mapping, not these example names when different.",
                {"source_id": {"type": "string"}, "table_id": {"type": "string"}, "contract": self._contract_schema(),
                 "expected_version": {"type": "integer", "minimum": 0}, "column_mapping": {"type": "object", "additionalProperties": {"type": "string"}},
                 "unit_evidence": {"type": "object", "description": "Keys are mapped numeric column names; values must be exact source substrings such as staff (persons), copied from source_header_quotes or source text. Do not add a prefix, translation or explanation. Trusted user-reviewed unit declarations are reused automatically when omitted.",
                                   "additionalProperties": {"type": "string", "minLength": 1}}}, ["source_id", "table_id", "contract", "expected_version"]),
            "web_search": (self.web_search, "Find public sources with Bing RSS or configured SearXNG. Inspect result URLs before using them as citation evidence.",
                {"query": {"type": "string", "maxLength": 500}, "limit": {"type": "integer", "minimum": 1, "maximum": 10}}, ["query"]),
        }
        registry = {}
        for name, (function, description, properties, required) in definitions.items():
            parameters = {"type": "object", "properties": properties, "required": required, "additionalProperties": False}
            validator = jsonschema.Draft202012Validator(parameters)
            def handler(arguments, function=function, validator=validator):
                try:
                    validator.validate(arguments)
                    return function(**arguments)
                except jsonschema.ValidationError as exc:
                    return {"status": "blocked", "code": "INVALID_ARGUMENTS", "message": exc.message}
                except ImportError:
                    return {"status": "blocked", "code": "DEPENDENCY_UNAVAILABLE", "message": "Required document parser is unavailable."}
                except (DocumentError, StoreError, ValueError, TypeError, KeyError, OSError, csv.Error, zipfile.BadZipFile) as exc:
                    return {"status": "blocked", "code": getattr(exc, "code", "DOCUMENT_ERROR"), "message": str(exc)}
            registry[name] = {"schema": {"type": "function", "function": {"name": name, "description": description,
                "parameters": parameters}},
                "handler": handler, "mutating": name == "publish_selected_table"}
        registry["publish_selected_table"]["recover"] = self.recover_publication
        return registry
