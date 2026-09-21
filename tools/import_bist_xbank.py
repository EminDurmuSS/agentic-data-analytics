"""Operator-only import of Borsa Istanbul's anonymous monthly XBANK workbook.

The downloaded ZIP remains the immutable source, not a generated CSV presented
as a download. A bounded, source-specific parser records workbook cells and
actual observation dates. No daily resampling, interpolation, price adjustment
or substitution of XUMAL/XU100 is performed. Not exposed as an agent tool.
"""
from __future__ import annotations

import argparse
import calendar
import csv
from datetime import date, datetime
from decimal import Decimal
import hashlib
import io
import json
from pathlib import Path
import tempfile
import zipfile

import httpx
from openpyxl import load_workbook

from agentic_analytics.agent.tools.documents import DocumentTools, OFFICIAL_SOURCE_REGISTRY
from agentic_analytics.lakehouse.shared import SharedLakehouse
from agentic_analytics.lakehouse.store import LakehouseStore, _write_json

SOURCE_URL = "https://www.borsaistanbul.com/datum/TR_PayEndeksleriFiyat.zip"
LANDING_URL = "https://www.borsaistanbul.com/veriler/konsolide-veriler"
SERIES_URL = "https://www.borsaistanbul.com/endeks/xbank"
MEMBER = "TR_PayEndeksleriFiyat.xlsx"
SHEET = "Pay Piyasası Fiyat Endeksleri K"
HEADERS = ("Tarih (GG.AA.YYYY)", "Endeks Kodu", "Endeksler", "Kur Türü", "Kapanış Değeri")
PARSER_VERSION = "bist-monthly-xbank-v1"
MAX_DOWNLOAD = 8 * 1024**2
MAX_EXPANDED = 64 * 1024**2


def download() -> bytes:
    # No credentials, arbitrary URLs or redirects. HTML/login/error responses
    # cannot pass the archive/header checks in parse_archive.
    with httpx.Client(timeout=60, follow_redirects=False) as client:
        with client.stream("GET", SOURCE_URL) as response:
            response.raise_for_status()
            data = bytearray()
            for chunk in response.iter_bytes():
                data.extend(chunk)
                if len(data) > MAX_DOWNLOAD:
                    raise ValueError("Official archive exceeds download limit")
    return bytes(data)


def parse_archive(raw: bytes, year: int) -> dict:
    if not 1997 <= year < date.today().year:
        raise ValueError("Choose a completed calendar year from 1997 onward")
    if not raw or len(raw) > MAX_DOWNLOAD:
        raise ValueError("Invalid archive size")
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        members = archive.infolist()
        if (len(members) != 1 or members[0].filename != MEMBER
                or members[0].flag_bits & 1 or members[0].file_size > MAX_DOWNLOAD):
            raise ValueError("Unexpected or oversized official workbook archive")
        workbook_bytes = archive.read(MEMBER)
    with zipfile.ZipFile(io.BytesIO(workbook_bytes)) as workbook_zip:
        parts = workbook_zip.infolist()
        if (len(parts) > 1000 or sum(part.file_size for part in parts) > MAX_EXPANDED
                or len({part.filename for part in parts}) != len(parts)
                or any(part.flag_bits & 1 for part in parts)):
            raise ValueError("Unsafe or oversized workbook")
    workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=False)
    observations = []
    try:
        if SHEET not in workbook.sheetnames:
            raise ValueError("Official sheet is missing")
        sheet = workbook[SHEET]
        if sheet.max_column != 5 or sheet.max_row > 100_000:
            raise ValueError("Official table shape has changed")
        rows = sheet.iter_rows(values_only=True)
        if tuple(next(rows)) != HEADERS:
            raise ValueError("Official headers have changed")
        for physical_row, values in enumerate(rows, 2):
            observed, code, name, currency, closing = values
            if code != "XBANK":
                continue
            if not isinstance(observed, (date, datetime)):
                raise ValueError("XBANK observation has an invalid or formula date")
            if observed.year != year:
                continue
            if name != "BIST BANKA" or currency != "TL":
                raise ValueError("XBANK identity/currency does not match the requested price index")
            if isinstance(observed, datetime):
                if observed.time() != datetime.min.time():
                    raise ValueError("Unexpected intraday timestamp")
                observed = observed.date()
            # This report contains native monthly snapshots. Reject a changed
            # daily feed, incomplete months or unexpected midmonth observations;
            # do not silently take the last of several dates in a month.
            last_day = calendar.monthrange(year, observed.month)[1]
            if observed.day < last_day - 7:
                raise ValueError("Unexpected non-month-end observation in monthly report")
            if type(closing) not in {int, float}:
                raise ValueError("Closing must be an original numeric cell, not text/formula")
            amount = Decimal(str(closing))
            if not amount.is_finite() or amount <= 0:
                raise ValueError("Closing must be finite and positive")
            observations.append({
                "month": observed.strftime("%Y-%m"), "closing_date": observed.isoformat(),
                "index_code": code, "index_name": name, "quotation_currency": currency,
                "closing": str(amount), "source_row": physical_row,
            })
    finally:
        workbook.close()
    observations.sort(key=lambda row: row["month"])
    expected = [f"{year}-{month:02d}" for month in range(1, 13)]
    if [row["month"] for row in observations] != expected:
        raise ValueError("Expected exactly one source observation for each of the 12 months")
    return {
        "parser": PARSER_VERSION, "year": year, "source_url": SOURCE_URL,
        "landing_url": LANDING_URL, "series_url": SERIES_URL,
        "archive_sha256": hashlib.sha256(raw).hexdigest(),
        "member": MEMBER, "workbook_sha256": hashlib.sha256(workbook_bytes).hexdigest(),
        "sheet": SHEET, "observations": observations,
    }


def publish(store: LakehouseStore, workspace_id: str, raw: bytes, year: int) -> dict:
    """Validate, ingest and explicitly promote; never change the base snapshot."""
    parsed = parse_archive(raw, year)  # All source checks happen before writes.
    workspace = store.workspace(workspace_id)
    import_key = f"{PARSER_VERSION}:{year}:{parsed['archive_sha256']}"
    for dataset_id in workspace["datasets"]:
        provenance = store.dataset_manifest(dataset_id)["contract"].get("document_provenance", {})
        if provenance.get("import_key") == import_key:
            return SharedLakehouse(store).promote(
                workspace_id, dataset_id, official_sources=OFFICIAL_SOURCE_REGISTRY,
                reason="Repeat of the same validated official XBANK import",
            )
    documents = DocumentTools(store, workspace_id, searxng_url=False)
    source = documents._register(raw, Path(SOURCE_URL).name, "application/zip", SOURCE_URL)
    observations = parsed["observations"]
    source_columns = ["closing_date", "index_code", "index_name", "quotation_currency", "closing"]
    origins = []
    for row in observations:
        cells = {column: {"sheet": SHEET, "cell": f"{letter}{row['source_row']}",
                          "archive_member": MEMBER, "candidate_row": len(origins) + 1,
                          "candidate_column": column}
                 for column, letter in zip(source_columns, "ABCDE")}
        origins.append({**cells, "month": {**cells["closing_date"], "transform": "calendar_month_label"}})
    # Source inspection contains original dates/values; derived month labels
    # live only in the dataset and are described explicitly in its provenance.
    inspection = {"source_id": source["source_id"], "parser": PARSER_VERSION,
                  "pages": [], "processed_pages": [], "text_truncated": False,
                  "text": SHEET + "\n" + " | ".join(HEADERS) + "\n" + "\n".join(
                      " | ".join(row[column] for column in source_columns) for row in observations),
                  "warnings": [f"Operator parser selected only the 12 XBANK TL observations for {year}; other series were not imported."],
                  "archive_member": MEMBER, "workbook_sha256": parsed["workbook_sha256"],
                  "tables": [{"table_id": "table_001", "origin": "parsed", "sheet": SHEET,
                              "page": None, "columns": source_columns,
                              "original_columns": dict(zip(source_columns, HEADERS)),
                              "rows": [[row[column] for column in source_columns] for row in observations],
                              "row_count": 12, "selection": {"index_code": "XBANK", "year": year, "currency": "TL"},
                              "row_origins": [{"sheet": SHEET, "source_row": row["source_row"]} for row in observations],
                              "cell_origins": [{key: item[key] for key in source_columns} for item in origins]}]}
    _write_json(documents._directory(source["source_id"]) / "inspection.json", inspection)
    dimension = {"dtype": "string", "unit": "label", "kind": "dimension", "nullable": False}
    contract = {
        "name": f"Borsa İstanbul BIST Banka Endeksi XBANK aylık kapanış {year} (TL fiyat endeksi)",
        "frequency": "monthly", "date_column": "month", "key": ["month"], "grain": ["month"],
        "expected_rows": 12, "expected_periods": [row["month"] for row in observations],
        "source_namespace": source["source_namespace"],
        "columns": {
            "month": {"dtype": "date", "unit": "calendar", "kind": "dimension", "nullable": False},
            "closing_date": {**dimension, "unit": "calendar"},
            "index_code": dimension, "index_name": dimension, "quotation_currency": dimension,
            "closing": {"dtype": "float", "unit": "index", "kind": "index", "scale": 1,
                        "currency": "TRY", "aggregation": "last", "additive_over_time": False,
                        "nullable": False, "source_semantics": "Monthly source-reported closing level, TL price index; not a return or money amount"},
        },
        "document_provenance": {
            "source_id": source["source_id"], "table_id": "table_001", "source_url": SOURCE_URL,
            "raw_sha256": parsed["archive_sha256"], "sheet": SHEET, "page": None,
            "archive_member": MEMBER, "workbook_sha256": parsed["workbook_sha256"],
            "landing_url": LANDING_URL, "series_metadata_url": SERIES_URL,
            "import_key": import_key, "parser": PARSER_VERSION,
            "numeric_verification": "original_numeric_workbook_cells_and_explicit_store_contract",
            "cell_origins": origins,
            "unit_evidence": {"closing": SHEET},
            "date_normalization": {
                "operation": "calendar_month_label_of_source_monthly_snapshot", "aggregation_performed": False,
                "original_date_column": "closing_date", "date_column": "month",
                "rows": [{"original": row["closing_date"], "normalized": row["month"]} for row in observations],
            },
        },
    }
    with tempfile.TemporaryDirectory(prefix="bist-xbank-") as temporary:
        csv_path = Path(temporary) / "monthly.csv"
        with csv_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(contract["columns"]), extrasaction="ignore")
            writer.writeheader()
            writer.writerows(observations)
        updated = store.ingest_csv(workspace_id, csv_path, contract, expected_version=workspace["version"])
    added = set(updated["datasets"]) - set(workspace["datasets"])
    if len(added) != 1:
        raise ValueError("Expected exactly one new verified dataset")
    result = SharedLakehouse(store).promote(
        workspace_id, added.pop(), official_sources=OFFICIAL_SOURCE_REGISTRY,
        reason=f"User-requested official XBANK {year} monthly closing import; 12 original cells verified",
    )
    return {**result, "row_count": 12, "workbook_sha256": parsed["workbook_sha256"], "observations": observations}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--year", type=int, default=2025)
    parser.add_argument("--publish", action="store_true", help="Explicitly publish to the shared release (default: inspect only)")
    parser.add_argument("--store", type=Path, help="Lakehouse runtime store, not the base DuckDB path")
    parser.add_argument("--workspace-id", help="Existing operator import workspace")
    args = parser.parse_args()
    if args.publish and (args.store is None or not args.workspace_id):
        parser.error("--publish requires --store and --workspace-id")
    raw = download()
    result = (publish(LakehouseStore(args.store), args.workspace_id, raw, args.year)
              if args.publish else parse_archive(raw, args.year))
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
