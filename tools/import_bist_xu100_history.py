"""Operator-only, source-specific import of official historical XU100 closes.

The anonymous Borsa Istanbul ZIP is the immutable primary evidence. This is
native MONTHLY data, not a daily history. Original dates and numeric cells are
preserved; no rebasing, daily resampling, holiday filling or EVDS replacement is
performed. Reviewed, hash-pinned official PDF evidence accompanies publication.
Default execution only downloads and validates. Not exposed as an agent tool.
"""
from __future__ import annotations

import argparse
import base64
import calendar
import csv
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
import gzip
import hashlib
import io
import json
from pathlib import Path
import tempfile
import zipfile

import httpx
from openpyxl import load_workbook

from agentic_analytics.agent.tools.documents import DocumentTools, OFFICIAL_SOURCE_REGISTRY
from agentic_analytics.lakehouse.shared import SharedLakehouse
from agentic_analytics.lakehouse.store import LakehouseStore, _canonical, _write_json

SOURCE_URL = "https://www.borsaistanbul.com/datum/TR_PayEndeksleriFiyat.zip"
LANDING_URL = "https://www.borsaistanbul.com/veriler/konsolide-veriler"
SERIES_URL = "https://www.borsaistanbul.com/endeks/xu100"
MEMBER = "TR_PayEndeksleriFiyat.xlsx"
SHEET = "Pay Piyasası Fiyat Endeksleri K"
HEADERS = ("Tarih (GG.AA.YYYY)", "Endeks Kodu", "Endeksler", "Kur Türü", "Kapanış Değeri")
PARSER_VERSION = "bist-monthly-xu100-history-v1"
MAX_DOWNLOAD = 8 * 1024**2
MAX_EXPANDED = 64 * 1024**2
MAX_SUPPORT_TOTAL = 16 * 1024**2
NAME_CHANGE_DATE = "2013-04-05"
SCALE_CHANGE_DATE = "2020-07-27"
SOURCE_COLUMNS = ("closing_date", "index_code", "index_name", "quotation_currency", "closing")
SEMANTICS = (
    "İMKB 100 / BIST 100 XU100 TL fiyat endeksi, kaynakta yayımlanan aylık kapanış; "
    "endeks puanıdır, TL para tutarı veya getiri değildir. Güncel kaynak geçmişi "
    "27.07.2020 iki sıfır atılması sonrası ölçeğinde yayımlar. Aktarımda ölçek "
    "dönüşümü yapılmaz. Günlük veri, enterpolasyon veya sentetik tatil doldurma yoktur. "
    "BIST 100 kaynak adı korunur; 05.04.2013 öncesi tarihsel adı İMKB 100'dür. "
    "Eski EVDS TP.MK.F.BILESIK kayıtlarıyla birleştirilmez; 2010-12 için "
    "660.04 (Borsa İstanbul) / 670.26 (yerel EVDS) uyuşmazlığı çözülmemiştir."
)
# Reviewed on 2026-09-21. A changed document requires renewed operator review,
# not a permissive fallback. Page numbers below are physical PDF pages (1-based).
SUPPORTING_DOCUMENTS = {
    "annual_2010": {
        "url": "https://www.borsaistanbul.com/files/IMKB_FINAL.pdf",
        "sha256": "7e504c98720d79f32edae3c2ce5870bd720b4cc5e45f647d3ff23fac6539f7b2",
        "evidence": [{"page": 18, "printed_page": "16", "locator": "right column, second paragraph; lower-left chart endpoint",
                      "finding": "İMKB 100 31.12.2010 closing: 66004.48 original-scale index points."}],
    },
    "name_change_2013": {
        "url": "https://www.borsaistanbul.com/datum/duyuru_ekleri/GenelMektup_4030_Endeks_Adlari.pdf",
        "sha256": "e9c7a8a086b1ccf4f45df09e86fe6d86b3c06cbd460f5fb774e14064f6309185",
        "evidence": [{"page": 1, "locator": "letter 4030, dated 04.04.2013, final substantive paragraph",
                      "finding": "Index-name change takes effect on 05.04.2013; not a new historical series."},
                     {"page": 2, "locator": "first data row in equity-index name-change annex",
                      "finding": "XU100: IMKB 100 -> BIST 100. Visually reviewed scanned table, not OCR-derived numerical data."}],
    },
    "scale_rules_2020": {
        "url": "https://www.borsaistanbul.com/files/bist-pay-endekslerinden-iki-sifir-atilmasi-ve-yapilacak-uyarlamalar-hakkinda-2020-12.pdf",
        "sha256": "9c0f894ec230027720d4eba3c3e6049431f07a2cb5509331dc19afadafdb0da3",
        "evidence": [{"page": 2, "printed_page": "1/2", "locator": "price-index annex first row and bottom comparison note",
                      "finding": "XU100 is the TL price index; old values are divided by 100 for comparison with the new scale."}],
    },
    "scale_go_live_2020": {
        "url": "https://www.borsaistanbul.com/files/2020-46_Removal_of_Zero_from_the_Index_Go_Live_Date_Announcement.pdf",
        "sha256": "63386df45d54c3a5758fda963b53e3d21298448eea1bf15774893eda75b509bb",
        "evidence": [{"page": 1, "locator": "announcement 2020/46, first two body paragraphs",
                      "finding": "Two-zero removal, including BIST 100, went live on 27.07.2020."}],
    },
    "historical_restatement_2020": {
        "url": "https://www.borsaistanbul.com/files/BIST2020EntegreFaaliyetRaporuveFinansalTablolar.pdf",
        "sha256": "f26d7dfb680235956d3a6534271edf2d01eed82037bb56c5fadff6a42b94d8b7",
        "evidence": [{"page": 24, "printed_page": "22", "locator": "Endeks Hesaplama / Yenilikler, first paragraph",
                      "finding": "The July 27 two-zero change covers TL price/return indexes; publicly shared historical data was updated."}],
    },
}


def _download(url: str) -> bytes:
    # Callers supply only the fixed URLs above. No redirects or credentials.
    with httpx.Client(timeout=60, follow_redirects=False) as client:
        with client.stream("GET", url) as response:
            response.raise_for_status()
            raw = bytearray()
            for chunk in response.iter_bytes():
                raw.extend(chunk)
                if len(raw) > MAX_DOWNLOAD:
                    raise ValueError("Official source exceeds download limit")
    return bytes(raw)


def download() -> bytes:
    return _download(SOURCE_URL)


def download_supporting_documents() -> dict[str, bytes]:
    raw = {key: _download(spec["url"]) for key, spec in SUPPORTING_DOCUMENTS.items()}
    _supporting_evidence(raw)
    return raw


def _supporting_evidence(raw: dict[str, bytes]) -> list[dict]:
    if not isinstance(raw, dict) or set(raw) != set(SUPPORTING_DOCUMENTS):
        raise ValueError("All reviewed supporting PDFs are required, with no extra documents")
    if any(not isinstance(value, bytes) for value in raw.values()):
        raise ValueError("Supporting documents must be immutable bytes")
    if sum(map(len, raw.values())) > MAX_SUPPORT_TOTAL:
        raise ValueError("Supporting PDFs exceed total size limit")
    evidence = []
    for key, spec in SUPPORTING_DOCUMENTS.items():
        content = raw[key]
        if (not content.startswith(b"%PDF-") or len(content) > MAX_DOWNLOAD
                or hashlib.sha256(content).hexdigest() != spec["sha256"]):
            raise ValueError(f"Supporting PDF changed or invalid: {key}; operator review required")
        evidence.append({"document_key": key, "source_url": spec["url"], "content_type": "application/pdf",
                         "raw_sha256": spec["sha256"], "size_bytes": len(content),
                         "reviewed_on": "2026-09-21", "review_method": "text_extraction_and_visual_page_review",
                         "page_numbering": "physical_1_based", "evidence": spec["evidence"],
                         "encoding": "gzip+base64", "raw_gzip_base64": base64.b64encode(
                             gzip.compress(content, mtime=0)).decode("ascii")})
    return evidence


def parse_archive(raw: bytes, year: int = 2010) -> dict:
    if type(year) is not int or year != 2010:
        raise ValueError("This reviewed import supports the 2010 calendar year only")
    if not raw or len(raw) > MAX_DOWNLOAD:
        raise ValueError("Invalid archive size")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            members = archive.infolist()
            if (len(members) != 1 or members[0].filename != MEMBER
                    or members[0].flag_bits & 1 or members[0].file_size > MAX_DOWNLOAD):
                raise ValueError("Unexpected or oversized official workbook archive")
            workbook_bytes = archive.read(MEMBER)
        with zipfile.ZipFile(io.BytesIO(workbook_bytes)) as workbook_zip:
            parts = workbook_zip.infolist()
            if (len(parts) > 1000 or sum(part.file_size for part in parts) > MAX_EXPANDED
                    or len({part.filename for part in parts}) != len(parts)
                    or any(part.flag_bits & 1 or "vbaProject" in part.filename
                           or "externalLinks/" in part.filename for part in parts)):
                raise ValueError("Unsafe or oversized workbook")
    except zipfile.BadZipFile as exc:
        raise ValueError("Expected the official ZIP/XLSX, not HTML or another response") from exc
    workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=False, keep_links=False)
    observations = []
    try:
        # The official export also contains a non-data generator metadata tab.
        if workbook.sheetnames not in ([SHEET], [SHEET, "Mozart Reports"]):
            raise ValueError("Official sheets have changed")
        sheet = workbook[SHEET]
        if sheet.max_column != 5 or not 2 <= sheet.max_row <= 100_000:
            raise ValueError("Official table shape has changed")
        rows = sheet.iter_rows(values_only=True)
        if tuple(next(rows)) != HEADERS:
            raise ValueError("Official headers have changed")
        for physical_row, values in enumerate(rows, 2):
            observed, code, name, currency, closing = values
            if code != "XU100":
                continue
            if not isinstance(observed, (date, datetime)):
                raise ValueError("XU100 observation has an invalid or formula date")
            if observed.year != year:
                continue
            if name != "BIST 100" or currency != "TL":
                raise ValueError("XU100 identity/currency does not match the TL price index")
            if isinstance(observed, datetime):
                if observed.time() != datetime.min.time():
                    raise ValueError("Unexpected intraday timestamp")
                observed = observed.date()
            if observed.day < calendar.monthrange(observed.year, observed.month)[1] - 7:
                raise ValueError("Unexpected non-month-end observation in monthly report")
            if type(closing) not in {int, float}:
                raise ValueError("Closing must be an original numeric cell, not text/formula/bool")
            amount = Decimal(str(closing))
            if not amount.is_finite() or amount <= 0:
                raise ValueError("Closing must be finite and positive")
            observations.append({
                "month": observed.strftime("%Y-%m"), "closing_date": observed.isoformat(),
                "index_code": code, "index_name": name, "quotation_currency": currency,
                "historical_index_name": "İMKB 100" if observed.isoformat() < NAME_CHANGE_DATE else "BIST 100",
                "closing": str(amount), "source_row": physical_row,
            })
    finally:
        workbook.close()
    observations.sort(key=lambda row: row["month"])
    expected = [f"{year}-{month:02d}" for month in range(1, 13)]
    if [row["month"] for row in observations] != expected:
        raise ValueError("Expected exactly one original source observation for each of the 12 months")
    # A separate contemporary official report supplies the review anchor. This
    # is validation only: imported values remain the workbook's original cells.
    annual_value = Decimal("66004.48")
    comparable = (annual_value / 100).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if observations[-1]["closing_date"] != "2010-12-31" or Decimal(observations[-1]["closing"]) != comparable:
        raise ValueError("2010 year-end conflicts with independently reviewed annual-report/scale evidence")
    return {"parser": PARSER_VERSION, "year": year, "source_url": SOURCE_URL,
            "landing_url": LANDING_URL, "series_url": SERIES_URL,
            "archive_sha256": hashlib.sha256(raw).hexdigest(), "member": MEMBER,
            "workbook_sha256": hashlib.sha256(workbook_bytes).hexdigest(), "sheet": SHEET,
            "observations": observations, "semantics": SEMANTICS,
            "cross_checks": [{"date": "2010-12-31", "document_key": "annual_2010", "page": 18,
                              "original_scale_value": str(annual_value), "comparison_divisor": 100,
                              "rounded_comparable_value": str(comparable), "match": True,
                              "changes_imported_values": False}]}


def publish(store: LakehouseStore, workspace_id: str, raw: bytes, year: int = 2010,
            *, supporting_raw: dict[str, bytes]) -> dict:
    """Validate all evidence first, then ingest/promote without touching EVDS."""
    parsed = parse_archive(raw, year)
    supporting = _supporting_evidence(supporting_raw)
    # Preserve complete original PDFs without exceeding the store's 8-MiB
    # metadata-file limit. SharedLakehouse already verifies/copies content-
    # addressed extraction artifacts; no new shared-package mechanism is needed.
    annual_evidence = {"artifact_type": "reviewed_supporting_pdf_evidence", "parser": PARSER_VERSION,
                       "primary_raw_sha256": parsed["archive_sha256"],
                       "supporting_documents": [item for item in supporting if item["document_key"].startswith("annual_")]}
    extraction_ref = "extraction_" + hashlib.sha256(_canonical(annual_evidence)).hexdigest()
    annual_artifact = {**annual_evidence, "extraction_id": extraction_ref}
    _canonical(annual_artifact)  # Validate exact encoded size before any writes.
    inspection_support = [item if not item["document_key"].startswith("annual_") else
                          {**{key: value for key, value in item.items() if key not in {"raw_gzip_base64", "encoding"}},
                           "bytes_artifact_ref": extraction_ref} for item in supporting]
    workspace = store.workspace(workspace_id)
    evidence_hash = hashlib.sha256(json.dumps(
        supporting, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()
    import_key = f"{PARSER_VERSION}:{year}:{parsed['archive_sha256']}:{evidence_hash}"
    for dataset_id in workspace["datasets"]:
        provenance = store.dataset_manifest(dataset_id)["contract"].get("document_provenance", {})
        if provenance.get("import_key") == import_key:
            return SharedLakehouse(store).promote(
                workspace_id, dataset_id, official_sources=OFFICIAL_SOURCE_REGISTRY,
                reason="Repeat of the same reviewed official XU100 history import")
    documents = DocumentTools(store, workspace_id, searxng_url=False)
    source = documents._register(raw, Path(SOURCE_URL).name, "application/zip", SOURCE_URL)
    observations = parsed["observations"]
    origins = []
    for row in observations:
        cells = {column: {"sheet": SHEET, "cell": f"{letter}{row['source_row']}", "archive_member": MEMBER,
                          "candidate_row": len(origins) + 1, "candidate_column": column}
                 for column, letter in zip(SOURCE_COLUMNS, "ABCDE")}
        origins.append({**cells, "month": {**cells["closing_date"], "transform": "calendar_month_label"},
                        "historical_index_name": {**cells["index_name"], "transform": "dated_name_alias",
                                                  "supporting_document": "name_change_2013"}})
    inspection = {
        "source_id": source["source_id"], "parser": PARSER_VERSION, "pages": [], "processed_pages": [],
        "text_truncated": False, "text": SEMANTICS + "\n" + SHEET + "\n" + " | ".join(HEADERS) + "\n" +
        "\n".join(" | ".join(row[column] for column in SOURCE_COLUMNS) for row in observations),
        "warnings": ["Native monthly XU100/TL closes only; not daily data or a merged EVDS series.",
                     "2010-12: source 660.04 conflicts with existing EVDS 670.26; EVDS is unchanged."],
        "archive_member": MEMBER, "workbook_sha256": parsed["workbook_sha256"],
        "supporting_documents": inspection_support, "cross_checks": parsed["cross_checks"],
        "tables": [{"table_id": "table_001", "origin": "parsed", "sheet": SHEET, "page": None,
                    "columns": list(SOURCE_COLUMNS), "original_columns": dict(zip(SOURCE_COLUMNS, HEADERS)),
                    "rows": [[row[column] for column in SOURCE_COLUMNS] for row in observations],
                    "row_count": len(observations), "selection": {"index_code": "XU100", "year": year, "currency": "TL"},
                    "row_origins": [{"sheet": SHEET, "source_row": row["source_row"]} for row in observations],
                    "cell_origins": [{key: item[key] for key in SOURCE_COLUMNS} for item in origins]}],
    }
    _canonical(inspection)
    evidence_directory = store._path("document_sources", workspace_id, "extractions")
    evidence_directory.mkdir(exist_ok=True)
    extraction_path = evidence_directory / (extraction_ref + ".json")
    if extraction_path.exists():
        if extraction_path.is_symlink() or extraction_path.read_bytes() != _canonical(annual_artifact) + b"\n":
            raise ValueError("Existing supporting evidence artifact differs")
    else:
        _write_json(extraction_path, annual_artifact)
    _write_json(documents._directory(source["source_id"]) / "inspection.json", inspection)
    dimension = {"dtype": "string", "unit": "label", "kind": "dimension", "nullable": False}
    evidence_refs = [{key: value for key, value in item.items() if key not in {"raw_gzip_base64", "encoding"}}
                     for item in supporting]
    contract = {
        "name": f"İMKB 100 / BIST 100 XU100 {year} aylık kapanış (TL fiyat endeksi, güncel ölçek)",
        "frequency": "monthly", "date_column": "month", "key": ["month"], "grain": ["month"],
        "expected_rows": len(observations), "expected_periods": [row["month"] for row in observations],
        "source_namespace": source["source_namespace"],
        "columns": {"month": {"dtype": "date", "unit": "calendar", "kind": "dimension", "nullable": False},
                    "closing_date": {**dimension, "unit": "calendar"}, "index_code": dimension,
                    "index_name": dimension, "quotation_currency": dimension, "historical_index_name": dimension,
                    "closing": {"dtype": "float", "unit": "index", "kind": "index", "scale": 1,
                                "currency": "TRY", "aggregation": "last", "additive_over_time": False,
                                "nullable": False, "source_semantics": SEMANTICS}},
        "document_provenance": {
            "source_id": source["source_id"], "table_id": "table_001", "source_url": SOURCE_URL,
            "raw_sha256": parsed["archive_sha256"], "sheet": SHEET, "page": None, "archive_member": MEMBER,
            "workbook_sha256": parsed["workbook_sha256"], "landing_url": LANDING_URL, "series_metadata_url": SERIES_URL,
            "parser": PARSER_VERSION, "import_key": import_key,
            "numeric_verification": "original_numeric_workbook_cells_plus_independent_annual_report_scale_check",
            "cell_origins": origins, "unit_evidence": {"closing": SHEET + "; scale_rules_2020 PDF p2"},
            "supporting_documents": evidence_refs, "supporting_evidence_sha256": evidence_hash,
            "extraction_artifact_ref": extraction_ref,
            "supporting_evidence_location": "source/inspection.json:supporting_documents; annual reports in source/extractions/ (gzip+base64 original PDF bytes)",
            "cross_checks": parsed["cross_checks"],
            "index_identity": {"code": "XU100", "type": "price", "quotation_currency": "TL",
                               "source_name": "BIST 100", "historical_name": "İMKB 100",
                               "name_change_effective_date": NAME_CHANGE_DATE, "splicing_performed": False},
            "scale_history": {"event": "removal_of_two_zeros", "effective_date": SCALE_CHANGE_DATE,
                              "source_values": "historical_values_already_republished_on_post_2020_scale",
                              "import_conversion_applied": False, "import_multiplier": 1,
                              "comparison_only_old_to_new_divisor": 100},
            "known_conflicts": [{"other_source": "EVDS TP.MK.F.BILESIK, existing local 2010-12 record",
                                 "month": "2010-12", "other_value": "670.26", "official_workbook_value": "660.04",
                                 "resolution": "unresolved; separate official series; no overwrite or merge"}],
            "date_normalization": {"operation": "calendar_month_label_of_source_monthly_snapshot",
                                   "aggregation_performed": False, "interpolation_performed": False,
                                   "synthetic_dates_created": False, "original_date_column": "closing_date", "date_column": "month",
                                   "rows": [{"original": row["closing_date"], "normalized": row["month"]} for row in observations]},
        },
    }
    with tempfile.TemporaryDirectory(prefix="bist-xu100-history-") as temporary:
        csv_path = Path(temporary) / "monthly.csv"
        with csv_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(contract["columns"]), extrasaction="ignore")
            writer.writeheader()
            writer.writerows(observations)
        updated = store.ingest_csv(workspace_id, csv_path, contract, expected_version=workspace["version"])
    added = set(updated["datasets"]) - set(workspace["datasets"])
    if len(added) != 1:
        raise ValueError("Expected exactly one new verified dataset")
    result = SharedLakehouse(store).promote(
        workspace_id, added.pop(), official_sources=OFFICIAL_SOURCE_REGISTRY,
        reason=f"User-requested official XU100 {year} monthly history; original cells and source-scale evidence reviewed")
    return {**result, "row_count": len(observations), "workbook_sha256": parsed["workbook_sha256"],
            "observations": observations, "cross_checks": parsed["cross_checks"]}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--year", type=int, default=2010)
    parser.add_argument("--publish", action="store_true", help="Explicitly publish (default: inspect only)")
    parser.add_argument("--store", type=Path, help="Runtime store, not the base DuckDB")
    parser.add_argument("--workspace-id", help="Existing operator import workspace")
    args = parser.parse_args()
    if args.publish and (args.store is None or not args.workspace_id):
        parser.error("--publish requires --store and --workspace-id")
    raw = download()
    result = (publish(LakehouseStore(args.store), args.workspace_id, raw, args.year,
                      supporting_raw=download_supporting_documents())
              if args.publish else parse_archive(raw, args.year))
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
