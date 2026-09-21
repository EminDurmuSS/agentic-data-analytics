"""Operator-only import of official XU100, XBANK and XUMAL base values.

Default execution only downloads and validates. Publication requires --publish,
an existing operator workspace and an explicit runtime store. The original XLSX
is retained: its date=value text cells, not hardcoded answers, supply the data.
These are cross-sectional index base levels at different dates, not returns or
a common-date time series. This importer is deliberately not an agent tool.
"""
from __future__ import annotations

import argparse
import csv
from datetime import datetime
from decimal import Decimal
import hashlib
import io
import json
from pathlib import Path
import re
import tempfile
import zipfile

import httpx
from openpyxl import load_workbook

from agentic_analytics.agent.tools.documents import DocumentTools, OFFICIAL_SOURCE_REGISTRY
from agentic_analytics.lakehouse.shared import SharedLakehouse
from agentic_analytics.lakehouse.store import LakehouseStore, _write_json

SOURCE_URL = "https://www.borsaistanbul.com/files/bist-endeks-kodlari-ve-baslangic-degerleri.xlsx"
SHEET = "Sheet1"
TITLE = "Endeks Başlangıç Değerleri (Base Values of Indices)"
HEADERS = (
    "Endeks Kodu / Index Code",
    "Endeksler / Index Names In Turkish",
    "Endekslerin İngilizce İsimleri / Index Names In English",
    "Endeksin Başlangıç Değeri / Base Value of Index",
)
IDENTITIES = {
    "XU100": ("BIST 100", "BIST 100"),
    "XBANK": ("BIST BANKA", "BIST BANKS"),
    "XUMAL": ("BIST MALI", "BIST FINANCIALS"),
}
PARSER_VERSION = "bist-index-baselines-v1"
MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_DOWNLOAD = 8 * 1024**2
MAX_EXPANDED = 64 * 1024**2
SOURCE_COLUMNS = ("index_code", "index_name", "index_name_en", "base_text")
SEMANTICS = (
    "Source-reported index base level in index points at that index's own base date. "
    "Başlangıç değeri; başlangıç tarihleri farklıdır. Not a return, performance "
    "ranking, current closing level, money amount or common-date observation. "
    "No rebasing or historical-scale correction is applied."
)


def download() -> bytes:
    """Read one fixed anonymous official URL; no redirects or credentials."""
    with httpx.Client(timeout=60, follow_redirects=False) as client:
        with client.stream("GET", SOURCE_URL) as response:
            response.raise_for_status()
            raw = bytearray()
            for chunk in response.iter_bytes():
                raw.extend(chunk)
                if len(raw) > MAX_DOWNLOAD:
                    raise ValueError("Official workbook exceeds download limit")
    return bytes(raw)


def parse_workbook(raw: bytes) -> dict:
    """Validate source layout/identity and split original literal D-column cells.

    Row numbers are discovered, not assumed. Formula cells (including cached
    formulas), duplicate target codes and ambiguous numeric formats fail closed.
    """
    if not raw or len(raw) > MAX_DOWNLOAD:
        raise ValueError("Invalid workbook size")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            parts = archive.infolist()
            if (len(parts) > 1000 or sum(part.file_size for part in parts) > MAX_EXPANDED
                    or len({part.filename for part in parts}) != len(parts)
                    or any(part.flag_bits & 1 for part in parts)
                    or any("vbaProject" in part.filename or "externalLinks/" in part.filename for part in parts)):
                raise ValueError("Unsafe or oversized workbook")
    except zipfile.BadZipFile as exc:
        raise ValueError("Expected an original XLSX workbook, not HTML or another response") from exc
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    observations = []
    try:
        if workbook.sheetnames != [SHEET]:
            raise ValueError("Official workbook sheets have changed")
        sheet = workbook[SHEET]
        if not 4 <= sheet.max_row <= 10000 or not 4 <= sheet.max_column <= 5:
            raise ValueError("Official table shape has changed")
        rows = sheet.iter_rows()
        if next(rows)[0].value != TITLE:
            raise ValueError("Official workbook title has changed")
        next(rows)
        headers = next(rows)
        if tuple(str(cell.value).strip() for cell in headers[:4]) != HEADERS:
            raise ValueError("Official headers have changed")
        for physical_row, cells in enumerate(rows, 4):
            code = cells[0].value
            if code not in IDENTITIES:
                continue
            if any(cell.data_type == "f" for cell in cells[:4]):
                raise ValueError("Selected source cells must be literal text, not formulas")
            code, name, english, base_text = (cell.value for cell in cells[:4])
            if (name, english) != IDENTITIES[code]:
                raise ValueError("Requested index identity does not match official labels")
            if any(cell.value is not None for cell in cells[4:]):
                raise ValueError("Unexpected extra information on a selected source row")
            match = re.fullmatch(r"(\d{2}\.\d{2}\.\d{4})\s*=\s*((?:0|[1-9]\d*)(?:,\d+)?)", base_text or "") if isinstance(base_text, str) else None
            if match is None:
                raise ValueError("Base cell must contain an unambiguous DD.MM.YYYY=decimal-comma literal")
            base_date = datetime.strptime(match[1], "%d.%m.%Y").date()
            amount = Decimal(match[2].replace(",", "."))
            if not amount.is_finite() or amount <= 0:
                raise ValueError("Base value must be finite and positive")
            observations.append({
                "index_code": code, "index_name": name, "index_name_en": english,
                "base_text": base_text, "base_date": base_date.isoformat(),
                "base_value": str(amount), "source_row": physical_row,
            })
    finally:
        workbook.close()
    if sorted(row["index_code"] for row in observations) != sorted(IDENTITIES):
        raise ValueError("Expected exactly one original source row for each of XU100, XBANK and XUMAL")
    # Source order remains intact for inspection and cell origins. The store may
    # sort keys; its source_row_indices preserve the resulting permutation.
    return {
        "parser": PARSER_VERSION, "source_url": SOURCE_URL,
        "raw_sha256": hashlib.sha256(raw).hexdigest(), "sheet": SHEET,
        "source_title": TITLE, "source_headers": list(HEADERS),
        "frequency": "static", "semantics": SEMANTICS, "observations": observations,
    }


def publish(store: LakehouseStore, workspace_id: str, raw: bytes) -> dict:
    """Validate, ingest and promote without changing any existing workspace pin."""
    parsed = parse_workbook(raw)  # Complete source validation before any writes.
    workspace = store.workspace(workspace_id)
    observations = parsed["observations"]
    import_key = f"{PARSER_VERSION}:{parsed['raw_sha256']}"
    for dataset_id in workspace["datasets"]:
        provenance = store.dataset_manifest(dataset_id)["contract"].get("document_provenance", {})
        if provenance.get("import_key") == import_key:
            result = SharedLakehouse(store).promote(
                workspace_id, dataset_id, official_sources=OFFICIAL_SOURCE_REGISTRY,
                reason="Repeat of the same validated official index base-values import",
            )
            return {**result, "row_count": len(observations), "raw_sha256": parsed["raw_sha256"],
                    "observations": observations}

    documents = DocumentTools(store, workspace_id, searxng_url=False)
    source = documents._register(raw, Path(SOURCE_URL).name, MIME_TYPE, SOURCE_URL)
    origins = []
    for number, row in enumerate(observations, 1):
        cells = {column: {"sheet": SHEET, "cell": f"{letter}{row['source_row']}",
                          "candidate_row": number, "candidate_column": column,
                          "original_value": row[column]}
                 for column, letter in zip(SOURCE_COLUMNS, "ABCD")}
        cells["base_date"] = {**cells["base_text"], "candidate_column": "base_date",
                              "transform": "split_equals_left_dd_mm_yyyy_to_iso_date"}
        cells["base_value"] = {**cells["base_text"], "candidate_column": "base_value",
                               "transform": "split_equals_right_decimal_comma_to_decimal_point"}
        origins.append(cells)
    inspection = {
        "source_id": source["source_id"], "parser": PARSER_VERSION, "pages": [],
        "processed_pages": [], "text_truncated": False,
        "text": TITLE + "\n" + " | ".join(HEADERS) + "\n" + "\n".join(
            " | ".join(row[column] for column in SOURCE_COLUMNS) for row in observations),
        "warnings": ["Operator selected only XU100, XBANK and XUMAL; other workbook rows were not imported.", SEMANTICS],
        "tables": [{
            "table_id": "table_001", "origin": "parsed", "sheet": SHEET, "page": None,
            "columns": list(SOURCE_COLUMNS), "original_columns": dict(zip(SOURCE_COLUMNS, HEADERS)),
            "rows": [[row[column] for column in SOURCE_COLUMNS] for row in observations],
            "row_count": len(observations), "selection": {"index_codes": list(IDENTITIES)},
            "row_origins": [{"sheet": SHEET, "source_row": row["source_row"]} for row in observations],
            "cell_origins": [{column: row[column] for column in SOURCE_COLUMNS} for row in origins],
        }],
    }
    _write_json(documents._directory(source["source_id"]) / "inspection.json", inspection)
    dimension = {"dtype": "string", "unit": "label", "kind": "dimension", "nullable": False}
    contract = {
        "name": "Borsa İstanbul endeks başlangıç değerleri XU100 BIST 100 XBANK BIST Banka XUMAL BIST Mali (base values)",
        "frequency": "static", "date_column": None, "key": ["index_code"], "grain": ["index_code"],
        "expected_rows": len(IDENTITIES), "source_namespace": source["source_namespace"],
        "columns": {
            "index_code": dimension, "index_name": dimension, "index_name_en": dimension,
            "base_text": dimension,
            "base_date": {**dimension, "dtype": "date", "unit": "calendar"},
            "base_value": {"dtype": "float", "unit": "index", "kind": "index", "scale": 1,
                           "aggregation": "none", "additive_over_time": False, "nullable": False,
                           "source_semantics": SEMANTICS},
        },
        "document_provenance": {
            "source_id": source["source_id"], "table_id": "table_001", "source_url": SOURCE_URL,
            "raw_sha256": parsed["raw_sha256"], "sheet": SHEET, "page": None,
            "import_key": import_key, "parser": PARSER_VERSION, "cell_origins": origins,
            "numeric_verification": "strict_split_of_original_literal_date_equals_decimal_comma_cells",
            "unit_evidence": {"base_value": HEADERS[3]}, "source_title": TITLE,
            "date_normalization": {"operation": "parse_each_original_base_date_dd_mm_yyyy_to_iso_date",
                                   "aggregation_performed": False, "common_date_asserted": False,
                                   "rows": [{"index_code": row["index_code"], "original": row["base_text"],
                                             "normalized": row["base_date"]} for row in observations]},
            "semantic_warning": SEMANTICS,
        },
    }
    with tempfile.TemporaryDirectory(prefix="bist-index-baselines-") as temporary:
        csv_path = Path(temporary) / "baselines.csv"
        with csv_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(contract["columns"]), extrasaction="ignore")
            writer.writeheader()
            writer.writerows(observations)
        updated = store.ingest_csv(workspace_id, csv_path, contract, expected_version=workspace["version"])
    added = set(updated["datasets"]) - set(workspace["datasets"])
    if len(added) != 1:
        raise ValueError("Expected exactly one new verified dataset")
    result = SharedLakehouse(store).promote(
        workspace_id, added.pop(), official_sources=OFFICIAL_SOURCE_REGISTRY,
        reason="User-requested official index base-values import; three original text cells verified",
    )
    return {**result, "row_count": len(observations), "raw_sha256": parsed["raw_sha256"],
            "observations": observations}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--publish", action="store_true", help="Publish to shared release (default: inspect only)")
    parser.add_argument("--store", type=Path, help="Lakehouse runtime store, not the base DuckDB path")
    parser.add_argument("--workspace-id", help="Existing operator import workspace")
    args = parser.parse_args()
    if args.publish and (args.store is None or not args.workspace_id):
        parser.error("--publish requires --store and --workspace-id")
    raw = download()
    result = publish(LakehouseStore(args.store), args.workspace_id, raw) if args.publish else parse_workbook(raw)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
