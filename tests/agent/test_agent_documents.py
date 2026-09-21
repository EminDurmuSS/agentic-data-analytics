"""Real document parsing/publication plus mocked public-network boundary checks."""

import io
import copy
import json
from pathlib import Path
import socket
import tempfile
import threading
import unittest
from unittest.mock import patch, MagicMock
from urllib import error

import duckdb

from agentic_analytics.agent.tools.documents import (
    DocumentError,
    DocumentTools,
    _ascii_safe_url,
    _article_metadata,
    _direct_document_url,
    _document_type,
    _official_registry,
    _public_destination,
    fetch_public_url,
)
from agentic_analytics.lakehouse.store import LakehouseStore


def _text_pdf():
    stream = b"BT /F1 16 Tf 50 740 Td (REPORT 12345 million TL) Tj ET"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
               b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"]
    data, offsets = b"%PDF-1.4\n", []
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += str(index).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    xref = len(data)
    data += b"xref\n0 6\n0000000000 65535 f \n"
    data += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    return data + b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n" + str(xref).encode() + b"\n%%EOF\n"


class AgentDocumentTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.directory = Path(self.temporary.name)
        database = self.directory / "source.duckdb"
        with duckdb.connect(str(database)) as connection:
            connection.execute("CREATE TABLE seed(value INTEGER)")
        self.store = LakehouseStore(self.directory / "store")
        snapshot = self.store.publish_snapshot(database)
        self.store.create_workspace(snapshot["snapshot_id"], "workspace_docs")
        self.docs = DocumentTools(self.store, "workspace_docs")
        self.contract = {"name": "imported_profit", "frequency": "monthly", "date_column": "month",
                         "key": ["month"], "grain": ["month"], "expected_rows": 2,
                         "columns": {
                             "month": {"dtype": "date", "unit": "calendar", "kind": "dimension", "nullable": False},
                             "value": {"dtype": "integer", "unit": "TRY_million", "kind": "flow", "nullable": False}}}

    def upload(self, name, data):
        path = self.docs.upload_root / name
        path.write_bytes(data if isinstance(data, bytes) else data.encode())
        source = self.docs.register_upload(path)
        path.unlink()
        return self.docs.inspect_source(source_id=source["source_id"])

    def test_csv_preserves_raw_publishes_and_recovers_without_duplicate_write(self):
        content = "month,value (million TL)\n2026-01,100\n2026-02,120\n"
        inspected = self.upload("profit.csv", content)
        self.assertEqual(inspected["tables"][0]["preview"][0]["value_million_TL"], "100")
        source_id = inspected["source_id"]
        self.assertEqual((self.docs._directory(source_id) / "raw.bin").read_text(), content)
        arguments = {"source_id": source_id, "table_id": "table_001", "contract": self.contract,
                     "expected_version": 0, "column_mapping": {"month": "month", "value_million_TL": "value"},
                     "unit_evidence": {"value": "million TL"}}
        result = self.docs.publish_selected_table(**arguments)
        self.assertEqual(result["status"], "ok")
        self.assertTrue(result["source_namespace"].startswith("external:upload:"))
        self.assertEqual(result["provenance"]["raw_sha256"], inspected["raw_sha256"])
        recovered = self.docs.extra_tools()["publish_selected_table"]["recover"](arguments, {})
        self.assertEqual(recovered["dataset_id"], result["dataset_id"])
        self.assertEqual(self.store.workspace("workspace_docs")["version"], 1)
        self.assertEqual(result["row_count"], 2)

    def test_ragged_csv_downgrades_to_a_warning_instead_of_losing_the_source(self):
        # A trailing footnote line (common in official exports) makes the CSV
        # inconsistent-width. That single ragged table must not discard the
        # whole source; the text stays readable and the issue is a warning.
        content = "month,value\n2026-01,100\n2026-02,120\nNot: kaynak dipnotu\n"
        inspected = self.upload("ragged.csv", content)
        self.assertEqual(inspected["status"], "ok")
        self.assertEqual(inspected["tables"], [])
        self.assertIn("AMBIGUOUS_TABLE", [w["code"] for w in inspected["warnings"]])
        self.assertIn("2026-01", inspected["text"])

    def test_unit_quote_and_numeric_format_are_gates(self):
        inspected = self.upload("profit.csv", 'month,value (million TL)\n2026-01,"1,25"\n2026-02,120\n')
        arguments = {"source_id": inspected["source_id"], "table_id": "table_001", "contract": self.contract,
                     "expected_version": 0, "column_mapping": {"month": "month", "value_million_TL": "value"}}
        result = self.docs.extra_tools()["publish_selected_table"]["handler"](arguments)
        self.assertEqual(result["code"], "UNITS_REVIEW_REQUIRED")
        result = self.docs.extra_tools()["publish_selected_table"]["handler"]({**arguments, "unit_evidence": {"value": "million TL"}})
        self.assertEqual(result["code"], "INVALID_NUMERIC_CELL")
        self.assertEqual(self.store.workspace("workspace_docs")["version"], 0)

    def test_publication_schema_exposes_contract_and_blocks_invalid_contract_before_write(self):
        source = self.upload("profit.csv", "month,value (million TL)\n2026-01,100\n2026-02,120\n")
        tool = self.docs.extra_tools()["publish_selected_table"]
        schema = tool["schema"]["function"]["parameters"]["properties"]["contract"]
        self.assertEqual(set(schema["properties"]["columns"]["additionalProperties"]["required"]), {"dtype", "unit", "kind", "nullable"})
        self.assertEqual(schema["properties"]["number_format"]["default"], "decimal_dot")
        args = {"source_id": source["source_id"], "table_id": "table_001", "contract": self.contract, "expected_version": 0,
                "column_mapping": {"month": "month", "value_million_TL": "value"}, "unit_evidence": {"value": "million TL"}}
        invalid = []
        for field in ["dtype", "unit", "kind", "nullable"]:
            contract = copy.deepcopy(self.contract)
            del contract["columns"]["value"][field]
            invalid.append(contract)
        no_date = copy.deepcopy(self.contract)
        del no_date["date_column"]
        invalid.extend([no_date, {**self.contract, "frequency": "made_up"}])
        with patch.object(self.store, "ingest_csv", wraps=self.store.ingest_csv) as ingest:
            for contract in invalid:
                result = tool["handler"]({**args, "contract": contract})
                self.assertEqual(result["code"], "INVALID_ARGUMENTS")
            ingest.assert_not_called()
        self.assertEqual(self.store.workspace("workspace_docs")["version"], 0)

    def test_monthly_iso_dates_require_declared_label_format_and_preserve_cells_and_units(self):
        import pandas as pd
        raw = "date,visits (count),staff (persons)\n2024-01-01,10,2\n2024-02-01,20,3\n"
        source = self.upload("monthly-clinic.csv", raw)
        self.assertEqual(source["tables"][0]["source_header_quotes"]["staff_persons"], "staff (persons)")
        contract = {"name": "monthly_clinic", "frequency": "monthly", "date_column": "date", "key": ["date"], "grain": ["date"],
                    "expected_rows": 2, "expected_periods": ["2024-01", "2024-02"], "columns": {
                        "date": {"dtype": "date", "kind": "dimension", "unit": "calendar", "nullable": False},
                        "visits_count": {"dtype": "integer", "kind": "count_flow", "unit": "visits", "nullable": False},
                        "staff_persons": {"dtype": "integer", "kind": "count_stock", "unit": "persons", "nullable": False}}}
        args = {"source_id": source["source_id"], "table_id": "table_001", "contract": contract, "expected_version": 0,
                "unit_evidence": {"visits_count": "visits (count)", "staff_persons": "staff (persons)"}}
        handler = self.docs.extra_tools()["publish_selected_table"]["handler"]
        explained_quote = {**args, "unit_evidence": {"visits_count": "visits (count)", "staff_persons": "Header: staff (persons), monthly staff count"}}
        self.assertEqual(handler(explained_quote)["code"], "UNITS_REVIEW_REQUIRED")
        self.assertEqual(handler(args)["code"], "SOURCE_DATE_FORMAT_REQUIRED")
        self.assertEqual(self.store.workspace("workspace_docs")["version"], 0)
        contract["source_date_format"] = "iso_period_start"
        result = handler(args)
        self.assertEqual(result["status"], "ok", result)
        frame = pd.read_parquet(self.store.overlay_path(result["dataset_id"]))
        self.assertEqual(frame["date"].tolist(), ["2024-01", "2024-02"])
        self.assertEqual(frame["visits_count"].tolist(), [10, 20])
        self.assertEqual(frame["staff_persons"].tolist(), [2, 3])
        self.assertEqual(self.docs.raw_source_bytes(source["source_id"]).decode(), raw)
        proof = result["provenance"]["date_normalization"]
        self.assertFalse(proof["aggregation_performed"])
        self.assertEqual(proof["rows"][0], {"candidate_row": 1, "original": "2024-01-01", "normalized": "2024-01"})

    def test_period_end_normalization_handles_leap_month_quarters_and_years(self):
        import pandas as pd
        for frequency, dates, expected in [
            ("monthly", ["2024-02-29", "2024-03-31"], ["2024-02", "2024-03"]),
            ("quarterly", ["2024-03-31", "2024-06-30"], ["2024-Q1", "2024-Q2"]),
            ("annual", ["2023-12-31", "2024-12-31"], ["2023", "2024"]),
        ]:
            with self.subTest(frequency=frequency):
                source = self.upload(frequency + ".csv", f"month,value (million TL)\n{dates[0]},100\n{dates[1]},120\n")
                contract = {**self.contract, "frequency": frequency, "source_date_format": "iso_period_end", "expected_periods": expected}
                result = self.docs.publish_selected_table(source["source_id"], "table_001", contract, self.store.workspace("workspace_docs")["version"],
                    column_mapping={"month": "month", "value_million_TL": "value"}, unit_evidence={"value": "million TL"})
                frame = pd.read_parquet(self.store.overlay_path(result["dataset_id"]))
                self.assertEqual(frame["month"].tolist(), expected)
                self.assertEqual(frame["value"].tolist(), [100, 120])

    def test_period_relabeling_rejects_daily_rows_and_duplicate_grain(self):
        handler = self.docs.extra_tools()["publish_selected_table"]["handler"]
        for label, dates, code in [
            ("daily", ["2024-01-01", "2024-01-02"], "NON_BOUNDARY_SOURCE_DATE"),
            ("duplicate", ["2024-01-01", "2024-01-01"], "DOCUMENT_ERROR"),
        ]:
            with self.subTest(label=label):
                source = self.upload(label + ".csv", f"month,value (million TL)\n{dates[0]},100\n{dates[1]},120\n")
                result = handler({"source_id": source["source_id"], "table_id": "table_001",
                    "contract": {**self.contract, "source_date_format": "iso_period_start"}, "expected_version": 0,
                    "column_mapping": {"month": "month", "value_million_TL": "value"}, "unit_evidence": {"value": "million TL"}})
                self.assertEqual(result["code"], code)
                self.assertEqual(self.store.workspace("workspace_docs")["version"], 0)

    def test_html_and_xlsx_tables_and_pdf_page_text(self):
        html = self.upload("visits.html", "<script>BAD INSTRUCTION</script><table><tr><th>month</th><th>visits</th></tr><tr><td>2026-01</td><td>10</td></tr></table>")
        self.assertEqual(html["tables"][0]["preview"], [{"month": "2026-01", "visits": "10"}])
        self.assertNotIn("BAD INSTRUCTION", html["text"])
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(["month", "visits"])
        workbook.active.append(["2026-01", 10])
        buffer = io.BytesIO()
        workbook.save(buffer)
        xlsx = self.upload("visits.xlsx", buffer.getvalue())
        self.assertEqual(xlsx["tables"][0]["preview"], html["tables"][0]["preview"])
        pdf = self.upload("report.pdf", _text_pdf())
        self.assertIn("12345 million TL", pdf["pages"][0]["text"])
        self.assertEqual(pdf["pages"][0]["page"], 1)

    def test_xlsx_sparse_title_promotes_real_header_and_literal_split_preserves_evidence(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Endeks Başlangıç Değerleri (Base Values of Indices)", None, None, None])
        sheet.append(["Endeks Kodu / Index Code", "Endeksler / Index Names In Turkish",
                      "Endeksin Başlangıç Değeri / Base Value of Index", None])
        sheet.append(["XU100", "BIST 100", "01.01.1986=0,01", None])
        sheet.append(["XBANK", "BIST BANKA", "28.12.1990=100", None])
        sheet.append(["XUMAL", "BIST MALİ", "28.12.1990=0,33", None])
        buffer = io.BytesIO()
        workbook.save(buffer)

        inspected = self.upload("index-base-values.xlsx", buffer.getvalue())
        table = inspected["tables"][0]
        self.assertEqual(list(table["original_columns"].values()), [
            "Endeks Kodu / Index Code", "Endeksler / Index Names In Turkish",
            "Endeksin Başlangıç Değeri / Base Value of Index",
        ])
        self.assertEqual(table["row_count"], 3)
        self.assertEqual(table["preamble_rows"][0][0], "Endeks Başlangıç Değerleri (Base Values of Indices)")
        columns = table["columns"]
        prepared = self.docs.prepare_source_table(
            inspected["source_id"], table["table_id"], selected_rows=[1, 2, 3], selected_columns=columns,
            split_columns=[{"column": columns[2], "separator": "=", "outputs": ["source_date", "base_value"]}],
        )
        self.assertEqual([row["source_date"] for row in prepared["preview"]],
                         ["01.01.1986", "28.12.1990", "28.12.1990"])
        self.assertEqual([row["base_value"] for row in prepared["preview"]], ["0,01", "100", "0,33"])
        self.assertEqual(prepared["source_header_quotes"]["base_value"],
                         "Endeksin Başlangıç Değeri / Base Value of Index")

        contract = {"name": "index_base_values", "frequency": "static", "date_column": None,
                    "key": ["code"], "grain": ["code"], "expected_rows": 3,
                    "number_format": "decimal_comma", "columns": {
                        "code": {"dtype": "string", "unit": "label", "kind": "dimension", "nullable": False},
                        "name": {"dtype": "string", "unit": "label", "kind": "dimension", "nullable": False},
                        "source_date": {"dtype": "string", "unit": "calendar", "kind": "dimension", "nullable": False},
                        "base_value": {"dtype": "float", "unit": "index", "kind": "index", "nullable": False},
                    }}
        published = self.docs.publish_selected_table(
            inspected["source_id"], prepared["table_id"], contract, expected_version=0,
            column_mapping={columns[0]: "code", columns[1]: "name", "source_date": "source_date",
                            "base_value": "base_value"},
            unit_evidence={"base_value": "Endeksin Başlangıç Değeri / Base Value of Index"},
        )
        self.assertEqual(published["status"], "ok", published)
        self.assertEqual(published["row_count"], 3)

    def test_utf8_text_and_full_review_candidate_preserve_source_content(self):
        text = self.upload("notes.txt", "Kaynak notu: İstanbul, milyon TL.\n")
        self.assertEqual(text["tables"], [])
        self.assertIn("İstanbul", text["text"])
        self.assertEqual(self.docs.raw_source_bytes(text["source_id"]).decode(), "Kaynak notu: İstanbul, milyon TL.\n")
        source = self.upload("long.csv", "month,visits\n" + "".join(f"2026-{month:02d},{month}\n" for month in range(1, 13)))
        self.assertEqual(len(source["tables"][0]["preview"]), 8)
        candidate = self.docs.review_candidate(source["source_id"], "table_001")
        self.assertEqual(len(candidate["rows"]), 12)
        self.assertFalse(candidate["preview_truncated"])
        self.assertEqual(candidate["original_columns"], {"month": "month", "visits": "visits"})
        self.assertNotIn("review_candidate", self.docs.extra_tools())

    def test_scanned_pdf_uses_bounded_image_callback_and_preserves_pages(self):
        from PIL import Image
        buffer = io.BytesIO()
        images = [Image.new("RGB", (120, 120), "white") for _ in range(4)]
        images[0].save(buffer, format="PDF", save_all=True, append_images=images[1:])
        calls = []
        def extractor(data, mime):
            calls.append((len(data), mime))
            return {"text": "month visits 2026-01 10", "tables": [{"columns": ["month", "visits"], "rows": [["2026-01", 10]]}]}
        self.docs.ocr_callback = extractor
        source = self.upload("scanned.pdf", buffer.getvalue())
        self.assertEqual(len(calls), 3)
        self.assertTrue(all(mime == "image/png" for _, mime in calls))
        self.assertEqual([table["page"] for table in source["tables"]], [1, 2, 3])
        self.assertTrue(all(table["origin"] == "ocr" for table in source["tables"]))
        self.assertEqual(next(warning for warning in source["warnings"] if warning["code"] == "OCR_PAGE_LIMIT")["pages"], [4])
        self.assertEqual(source["pages"][0]["extraction_method"], "ocr")

    def test_image_callback_is_bounded_and_incomplete_ocr_is_not_a_table(self):
        from PIL import Image
        buffer = io.BytesIO()
        Image.new("RGB", (30, 30), "white").save(buffer, format="PNG")
        seen = []
        def ocr(data, mime):
            seen.append((len(data), mime))
            return {"status": "ok", "text": "month visits 2026-01 10", "tables": [{"columns": ["month", "visits"], "rows": [["2026-01", 10]]}]}
        self.docs.ocr_callback = ocr
        inspected = self.upload("image.png", buffer.getvalue())
        self.assertEqual(seen[0][1], "image/png")
        self.assertEqual(inspected["tables"][0]["origin"], "ocr")
        self.assertEqual(inspected["tables"][0]["preview"][0]["visits"], "10")
        self.docs.ocr_callback = lambda *args: {"status": "ok", "text": "nonsense", "finish_reason": "length"}
        with self.assertRaisesRegex(DocumentError, "incomplete"):
            self.upload("other.png", buffer.getvalue())

    def test_paths_cross_workspace_and_raw_tampering_are_rejected(self):
        outside = self.directory / "outside.csv"
        outside.write_text("x\n1\n")
        with self.assertRaises(DocumentError):
            self.docs.register_upload(outside)
        with self.assertRaises(DocumentError):
            self.docs.source("../../outside")
        inspected = self.upload("ok.csv", "x\n1\n")
        (self.docs._directory(inspected["source_id"]) / "raw.bin").write_text("x\n2\n")
        with self.assertRaisesRegex(DocumentError, "hash mismatch"):
            self.docs.source(inspected["source_id"])

    def test_unconfigured_search_and_unknown_write_outcome_are_explicit(self):
        self.docs.searxng_url = False
        self.assertEqual(self.docs.web_search("test")["code"], "WEB_SEARCH_UNCONFIGURED")
        self.assertEqual(self.docs.recover_publication({"source_id": "absent", "table_id": "table_001", "contract": {}, "expected_version": 0}, {})["code"], "WRITE_OUTCOME_UNKNOWN")
        for tool in self.docs.extra_tools().values():
            self.assertNotIn("path", tool["schema"]["function"]["parameters"]["properties"])

    def test_private_dns_and_redirect_destinations_are_blocked(self):
        def resolve(host, port, **kwargs):
            ip = "127.0.0.1" if host == "local.test" else "93.184.216.34"
            return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", (ip, port))]
        with patch("agentic_analytics.agent.tools.documents.socket.getaddrinfo", side_effect=resolve):
            for url in ["http://local.test/x", "file:///etc/passwd", "https://u:p@public.test", "https://public.test:8080"]:
                with self.assertRaises(DocumentError):
                    _public_destination(url)
            redirect = error.HTTPError("https://public.test", 302, "Found", {"Location": "http://local.test/secret"}, io.BytesIO())
            with patch("agentic_analytics.agent.tools.documents.request.build_opener") as builder:
                builder.return_value.open.side_effect = redirect
                with self.assertRaisesRegex(DocumentError, "Private"):
                    fetch_public_url("https://public.test")
                self.assertEqual(builder.return_value.open.call_count, 1)

    def test_dns_resolution_has_a_time_budget(self):
        release = threading.Event()
        def delayed(*args, **kwargs):
            release.wait(1)
            return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 443))]
        with patch("agentic_analytics.agent.tools.documents.socket.getaddrinfo", side_effect=delayed):
            try:
                with self.assertRaises(DocumentError) as caught:
                    _public_destination("https://public.test", timeout=0.001)
                self.assertEqual(caught.exception.code, "FETCH_TIMEOUT")
            finally:
                release.set()

    def test_default_bing_search_is_bounded_and_does_not_claim_source_verification(self):
        rss = b'<rss><channel><item><title>Report</title><link>https://example.org/report</link><description>123</description></item></channel></rss>'
        with patch("agentic_analytics.agent.tools.documents.fetch_public_url", return_value=(rss, "application/rss+xml", "https://www.bing.com/search")) as fetch:
            result = self.docs.web_search("public report", limit=1)
            self.assertEqual(result["source_backend"], "Bing RSS")
            self.assertFalse(result["sources_verified"])
            self.assertEqual(len(result["results"]), 1)
            self.assertIn("format=rss", fetch.call_args.args[0])
        with patch("agentic_analytics.agent.tools.documents.fetch_public_url", return_value=(b"<html>challenge</html>", "text/html", "https://www.bing.com/search")):
            self.assertEqual(self.docs.web_search("public report")["code"], "SEARCH_INVALID_RESPONSE")

    def test_drifted_search_uses_one_alternate_index_and_retains_full_query(self):
        from urllib.parse import parse_qs, urlsplit
        rss = b'<rss><channel><item><title>Example Bank</title><link>https://example.org/</link></item><item><title>Job listings</title><link>https://jobs.test/</link></item></channel></rss>'
        html = b'''<table><tr><td><a class="result-link" href="//duckduckgo.com/l/?uddg=https%3A%2F%2Fexample.org%2F2026%2Ffinancial-report.pdf">Example Bank March 2026 consolidated report</a></td></tr>
        <tr><td class="result-snippet">Consolidated <b>financial statements</b> for March 2026.</td></tr></table>'''
        query = 'Example Bank March 2026 consolidated financial report PDF'
        with patch('agentic_analytics.agent.tools.search_backend.kap_financial_search', return_value=[]), \
                patch("agentic_analytics.agent.tools.documents.fetch_public_url", side_effect=[
                (rss, 'application/rss+xml', 'https://www.bing.com/search'),
                (html, 'text/html', 'https://lite.duckduckgo.com/lite/')]) as fetch:
            result = self.docs.web_search(query)
        self.assertEqual(fetch.call_count, 2)
        self.assertEqual(parse_qs(urlsplit(fetch.call_args_list[1].args[0]).query)['q'], [query])
        self.assertEqual(result['results'][0]['url'], 'https://example.org/2026/financial-report.pdf')
        self.assertNotIn('jobs.test', json.dumps(result['results']))
        self.assertIn('DuckDuckGo Lite', result['source_backend'])
        self.assertFalse(result['sources_verified'])

    def test_search_domain_filter_applies_to_every_backend_and_returns_actionable_recovery(self):
        self.docs.searxng_url = 'http://localhost:8080'
        entries = [
            {'title': '31 meaning', 'url': 'https://unrelated.test/31', 'content': '31 March 2026'},
            {'title': 'Financial report', 'url': 'https://example.org.evil.test/report.pdf'},
            {'title': 'Example Bank', 'url': 'https://example.org/'}]
        with patch('agentic_analytics.agent.tools.search_backend.configured_search', return_value=entries), \
                patch('agentic_analytics.agent.tools.documents.fetch_public_url', side_effect=OSError('Index unavailable')) as public:
            result = self.docs.web_search('site:example.org 31 March 2026 consolidated financial report')
        self.assertEqual(public.call_count, 2)
        self.assertEqual([item['url'] for item in result['results']], ['https://example.org/'])
        self.assertEqual(result['code'], 'SEARCH_DISCOVERY_ONLY')
        self.assertEqual(result['recovery']['arguments']['domains'], ['example.org'])
        self.assertEqual(next(item['count'] for item in result['warnings'] if item['code'] == 'IRRELEVANT_SEARCH_RESULTS_SKIPPED'), 2)

    def test_search_challenge_preserves_navigation_lead_without_claiming_report_found(self):
        rss = b'<rss><channel><item><title>Example Bank</title><link>https://example.org/</link></item></channel></rss>'
        with patch('agentic_analytics.agent.tools.search_backend.kap_financial_search', return_value=[]), \
                patch('agentic_analytics.agent.tools.documents.fetch_public_url', side_effect=[
                (rss, 'text/xml', 'https://www.bing.com/search'),
                (b'<html>Challenge required</html>', 'text/html', 'https://lite.duckduckgo.com/lite/')]) as fetch:
            result = self.docs.web_search('Example Bank 2026 financial report')
        self.assertEqual(fetch.call_count, 2)
        self.assertEqual(result['code'], 'SEARCH_DISCOVERY_ONLY')
        self.assertEqual(result['warnings'][0]['code'], 'SEARCH_FALLBACK_UNAVAILABLE')
        self.assertTrue(result['results'][0]['discovery_only'])

    def test_research_reports_login_paywall_and_captcha_without_registering_them_as_sources(self):
        cases = [
            ("Sign in to continue. Login required.", "AUTHENTICATION_REQUIRED"),
            ("Subscribe to continue reading. Subscription required.", "PAYWALL_REQUIRED"),
            ("Verify you are human. CAPTCHA security challenge.", "BOT_CHALLENGE"),
        ]
        for body, code in cases:
            with self.subTest(code=code):
                url = "https://example.org/restricted"
                self.docs.web_search = lambda *args, **kwargs: {"status": "ok", "results": [{
                    "title": "Official 2026 report", "url": url, "discovery_only": False,
                    "entity_verification_required": False,
                }]}
                html = f"<html><head><title>Official 2026 report</title></head><body><main>{body}</main></body></html>".encode()
                with patch("agentic_analytics.agent.tools.documents.fetch_public_url",
                           return_value=(html, "text/html", url)):
                    result = self.docs.research_web(
                        "Official 2026 report",
                        limit=1,
                        domains=["example.org"],
                    )
                self.assertEqual("unavailable", result["status"])
                self.assertEqual([], result["sources"])
                self.assertIn(code, {failure["code"] for failure in result["failures"]})
                self.assertTrue(any("no " in failure["message"].casefold()
                                    for failure in result["failures"] if failure["code"] == code))

    def test_research_verifies_issuer_after_dated_search_abstract_omits_it(self):
        from agentic_analytics.agent.tools.search_backend import rank_search_results
        query = 'Example Bank 31 March 2026 consolidated financial report'
        result = {'title': '31 March 2026 consolidated financial report', 'url': 'https://filings.test/report.pdf', 'snippet': ''}
        ranked, _, _ = rank_search_results(query, [result])
        self.assertTrue(ranked[0]['entity_verification_required'])
        self.docs.web_search = lambda *args, **kwargs: {'status': 'ok', 'results': ranked}
        self.docs.inspect_source = lambda **kwargs: {'source_id': 'report', 'source_url': result['url'],
            'text': 'Unrelated Company 31 March 2026 consolidated financial report'}
        self.assertEqual(self.docs.research_web(query, limit=1)['sources'], [])
        self.docs.inspect_source = lambda **kwargs: {'source_id': 'report', 'source_url': result['url'],
            'text': 'Example Bank 31 March 2026 consolidated financial report'}
        self.assertEqual(self.docs.research_web(query, limit=1)['sources'][0]['source_id'], 'report')

    def test_logo_names_and_late_document_links_survive_repeated_navigation(self):
        navigation = '<nav><a href="/home">Home</a></nav>' * 600
        html = (navigation + '<main><h1>Shareholders</h1><a href="https://example.org"><img alt="Example Bank" src="bank.png"></a>'
                '<a href="/report.pdf">2026 consolidated report</a></main>').encode()
        result = _article_metadata(html, 'text/html', 'https://issuer.test/owners')
        self.assertIn('Example Bank', result['readable_text'])
        self.assertEqual(result['source_links'][0]['title'], 'Example Bank')
        self.assertEqual(result['document_links'][0]['url'], 'https://issuer.test/report.pdf')

    def test_svg_chart_source_labels_stay_with_shareholder_section_without_inferred_names_or_amounts(self):
        html = '''<main><h4>Ortaklık Yapısı</h4><svg>
            <defs><symbol id="unused-bank"><title>Unused Bank</title></symbol></defs>
            <circle role="progressbar" data-slice-name="examplebank" aria-valuenow="18" />
            <circle role="progressbar" aria-label="Başka Banka" aria-valuenow="9" />
            <use href="#unlabelled-symbol" />
            </svg><h4>Üye Dağılımı</h4><p>Only Member Bank</p></main>'''.encode()
        text = _article_metadata(html, 'text/html', 'https://issuer.test/owners')['readable_text']
        owners, members = text.split('Üye Dağılımı')
        self.assertIn('Ortaklık Yapısı', owners)
        self.assertIn('Grafik kategorisi (kaynak etiketi): examplebank', owners)
        self.assertIn('Başka Banka', owners)
        self.assertNotIn('Only Member Bank', owners)
        self.assertIn('Only Member Bank', members)
        for absent in ('Unused Bank', 'unlabelled-symbol', '18', '9', 'Example Bank'):
            self.assertNotIn(absent, text)

    def test_ownership_research_follows_about_page_across_www_alias_before_unrelated_pdf(self):
        root = 'https://kkb.com.tr/'
        about = 'https://www.kkb.com.tr/hakkimizda'
        owners = 'https://www.kkb.com.tr/ortaklik-yapisi'
        pages = {
            root: {'source_id': 'root', 'source_url': root, 'text': 'KKB ana sayfa', 'article': {
                'title': 'Ana Sayfa', 'source_links': [
                    {'url': about, 'title': 'Hakkımızda', 'in_main_content': True},
                    {'url': about, 'title': 'Hakkımızda', 'in_navigation': True}],
                'document_links': [{'url': 'https://www.kkb.com.tr/other.pdf', 'title': 'Diğer rapor'}]}},
            about: {'source_id': 'about', 'source_url': about, 'text': 'KKB kurum bilgileri', 'article': {
                'title': 'Hakkımızda', 'link_count': 40, 'source_links': [
                    {'url': owners, 'title': 'Ortaklık Yapısı', 'in_main_content': True}]}},
            owners: {'source_id': 'owners', 'source_url': owners,
                     'text': 'KKB Ortaklık Yapısı\nÖrnek Bankası\nBaşka Bankası',
                     'article': {'title': 'Ortaklık Yapısı', 'link_count': 40}}}
        with patch.object(self.docs, 'web_search', return_value={'status': 'ok', 'results': [
                {'url': root, 'title': 'KKB', 'discovery_only': True}]}) as search, \
                patch.object(self.docs, 'inspect_source', side_effect=lambda **args: pages[args['url']]) as read:
            result = self.docs.research_web('KKB Bankası ortakları hissedarları', limit=1)
        self.assertTrue(search.call_args.args[0].startswith('site:kkb.com.tr '))
        self.assertEqual([call.kwargs['url'] for call in read.call_args_list], [root, about, owners])
        self.assertEqual(result['sources'][0]['url'], owners)
        self.assertIn('Örnek Bankası', result['sources'][0]['content'])

    def test_article_metadata_survives_valueless_html_attributes(self):
        # A bare attribute (<img alt>, <a href>, <link rel>, <script type>) is
        # valid HTML and html.parser hands it back as None, not "". Any
        # unguarded attrs.get(name, "").strip()/.split() crashes with an
        # AttributeError on real-world pages using this shorthand.
        html = (b'<html><head><title>T</title><link rel><script type>{}</script></head>'
                b'<body><a href><img src="x" alt></a><svg><path aria-label></path></svg></body></html>')
        article = _article_metadata(html, "text/html", "https://example.org/page")
        self.assertEqual(article["title"], "T")

    def test_research_web_reads_json_ld_article_content_and_skips_unreadable_results(self):
        html = b'''<html><head><title>Fallback title</title>
        <script type="application/ld+json">{"@type":"NewsArticle","headline":"Official report",
        "datePublished":"2026-09-11T10:00:00Z","articleBody":"The report states the latest result."}</script>
        </head><body>Visible article text</body></html>'''
        article = _article_metadata(html, "text/html", "https://example.org/report")
        self.assertEqual(article["title"], "Official report")
        self.assertEqual(article["date_published"], "2026-09-11T10:00:00Z")
        self.assertIn("latest result", article["article_body"])
        self.docs.web_search = lambda query, limit=5, **kwargs: {"status": "ok", "results": [
            {"title": "Unreadable", "url": "https://bad.example/no", "snippet": ""},
            {"title": "Readable", "url": "https://example.org/report", "snippet": ""},
        ]}
        self.docs.inspect_source = lambda url=None, source_id=None, **kwargs: (
            {"status": "ok", "source_id": "source_" + "a" * 64, "source_url": url,
             "text": "Visible article text", "article": article}
            if "example.org" in url else (_ for _ in ()).throw(DocumentError("blocked", "FETCH_FAILED")))
        result = self.docs.research_web("latest report", limit=1)
        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["read"], 1)
        self.assertEqual(result["sources"][0]["title"], "Official report")

    def test_research_web_does_not_fall_back_to_unrelated_domains_for_official_queries(self):
        self.docs.web_search = lambda query, limit=5, **kwargs: {"status": "ok", "results": [
            {"title": "Housing listings", "url": "https://emlakjet.com/listings", "snippet": ""},
        ]}
        with patch.object(self.docs, "inspect_source", side_effect=DocumentError("Official root unavailable", "FETCH_FAILED")) as inspect:
            result = self.docs.research_web("TCMB konut fiyat endeksi 2026 raporu", limit=3)
            self.assertTrue(all("tcmb.gov.tr" in call.kwargs["url"] for call in inspect.call_args_list))
        self.assertIn(result.get("code"), {None, "OFFICIAL_SOURCE_NOT_FOUND", "NO_READABLE_SOURCES"})
        self.assertEqual(result.get("sources"), [])
        self.assertEqual(result["sources"], [])

    def test_official_registry_resolves_subdomains_without_accepting_lookalikes(self):
        self.assertEqual(_official_registry("www.garantibbvainvestorrelations.com")["institution"],
                         "Garanti BBVA Yatırımcı İlişkileri")
        self.assertEqual(_official_registry("evds3.tcmb.gov.tr")["institution"], "TCMB EVDS")
        self.assertEqual(_official_registry("veriportali.tuik.gov.tr")["institution"], "TÜİK Veri Portalı")
        self.assertIsNone(_official_registry("borsaistanbul.com.evil.test"))

    def test_named_official_sources_are_scoped_before_search_and_use_safe_entrypoints(self):
        cases = [
            ("Garanti BBVA 31 Mart 2026 konsolide finansal raporu", "garantibbvainvestorrelations.com",
             "https://www.garantibbvainvestorrelations.com/en/library/brsa-consolidated-financials-pdf/PDF/1268/0/0"),
            ("KAP 2025 finansal raporu", "kap.org.tr", "https://www.kap.org.tr/tr/sirketler/ALL"),
            ("İMKB 100 BIST 100 endeks ad değişikliği", "borsaistanbul.com",
             "https://www.borsaistanbul.com/datum/duyuru_ekleri/GenelMektup_4030_Endeks_Adlari.pdf"),
            ("EVDS TP.MK.F.BILESIK aylık kapanış", "evds3.tcmb.gov.tr", "https://evds3.tcmb.gov.tr/anasayfa"),
            ("TÜİK Veri Portalı İstanbul konut satışları", "veriportali.tuik.gov.tr", "https://veriportali.tuik.gov.tr/"),
        ]
        for query, domain, entrypoint in cases:
            with self.subTest(query=query):
                searches, reads = [], []
                self.docs.web_search = lambda value, **kwargs: searches.append(value) or {
                    "status": "ok", "results": []}
                def unreadable(**kwargs):
                    reads.append(kwargs["url"])
                    raise DocumentError("unavailable", "FETCH_FAILED")
                self.docs.inspect_source = unreadable
                result = self.docs.research_web(query, limit=1)
                self.assertTrue(searches[0].startswith("site:" + domain + " "))
                self.assertEqual(reads[0], entrypoint)
                self.assertEqual(result["status"], "unavailable")

    def test_borsa_2010_research_starts_from_the_official_annual_report_document(self):
        searches, reads = [], []
        self.docs.web_search = lambda value, **kwargs: searches.append(value) or {
            "status": "ok", "results": []}

        def unreadable(**kwargs):
            reads.append(kwargs["url"])
            raise DocumentError("unavailable", "FETCH_FAILED")

        self.docs.inspect_source = unreadable
        result = self.docs.research_web(
            "2010 yılına ait İMKB 100 kapanış verisini arıyorum.",
            limit=1,
        )
        annual_report = "https://www.borsaistanbul.com/files/IMKB_FINAL.pdf"
        self.assertTrue(searches[0].startswith("site:borsaistanbul.com "))
        self.assertEqual(annual_report, reads[0])
        self.assertTrue(_direct_document_url(annual_report))
        self.assertFalse(_direct_document_url("https://www.borsaistanbul.com/endeks/xu100"))
        self.assertEqual("unavailable", result["status"])

    def test_research_web_starts_at_curated_entrypoints_when_search_finds_nothing(self):
        from agentic_analytics.agent.tools.documents import OFFICIAL_SOURCE_REGISTRY
        self.docs.web_search = MagicMock(return_value={"status": "ok", "results": []})
        with patch.object(self.docs, "inspect_source", side_effect=DocumentError("unavailable", "FETCH_FAILED")) as inspect:
            self.docs.research_web("2010 IMKB 100 kapanış verisi", limit=1)
        attempted = {call.kwargs["url"] for call in inspect.call_args_list}
        self.assertTrue(attempted & set(OFFICIAL_SOURCE_REGISTRY["borsaistanbul.com"]["entrypoints"]))
        self.assertNotIn("https://borsaistanbul.com/", attempted)

    def test_search_publication_date_never_becomes_fetched_document_date(self):
        from agentic_analytics.agent.context import _model_tool_result
        url = 'https://example.org/financial-report'
        self.docs.web_search = lambda *args, **kwargs: {'status': 'ok', 'results': [{
            'url': url, 'title': 'Example Bank financial report', 'published_at': '2026-04-30',
            'publication_date_basis': 'search_metadata_unverified'}]}
        for source_date in (None, '2026-05-02'):
            with self.subTest(source_date=source_date):
                metadata = {'@type': 'Article', 'headline': 'Example Bank financial report'}
                if source_date:
                    metadata['datePublished'] = source_date
                html = ('<html><head><script type="application/ld+json">' + json.dumps(metadata)
                        + '</script></head><body><main><h1>Example Bank financial report</h1>'
                        + '<p>Example Bank financial report presents annual performance.</p></main></body></html>').encode()
                with patch('agentic_analytics.agent.tools.documents.fetch_public_url', return_value=(html, 'text/html', url)):
                    result = self.docs.research_web('Example Bank financial report', limit=1)
                card = result['sources'][0]
                self.assertEqual(card['date_published'], source_date)
                self.assertEqual(card['search_published_at'], '2026-04-30')
                self.assertEqual(card['search_publication_date_basis'], 'search_metadata_unverified')
                self.assertEqual(card['verification'], 'direct_public_fetch')
                projected = _model_tool_result('research_web', result)['sources'][0]
                self.assertEqual(projected['date_published'], source_date)
                self.assertEqual(projected['search_published_at'], '2026-04-30')
                self.assertEqual(projected['search_publication_date_basis'], 'search_metadata_unverified')

    def test_human_review_enables_ocr_publication_and_preserves_review_provenance(self):
        from PIL import Image
        buffer = io.BytesIO()
        Image.new("RGB", (30, 30), "white").save(buffer, format="PNG")
        self.docs.ocr_callback = lambda *args: {"status": "ok", "text": "month value million TL",
            "tables": [{"columns": ["month", "value"], "rows": [["2026-01", 999], ["2026-02", 120]]}]}
        source = self.upload("review.png", buffer.getvalue())
        args = {"source_id": source["source_id"], "table_id": "table_001", "contract": self.contract, "expected_version": 0}
        self.assertEqual(self.docs.extra_tools()["publish_selected_table"]["handler"](args)["code"], "OCR_REVIEW_REQUIRED")
        review = self.docs.review_table(source["source_id"], "table_001", [{"month": "2026-01", "value": 100}, {"month": "2026-02", "value": 120}], {"value": "million TL"})
        result = self.docs.publish_selected_table(**args)
        self.assertEqual(result["provenance"]["human_review"]["review_sha256"], review["review_sha256"])
        reviewed_origin = result["provenance"]["cell_origins"][0]["value"]
        self.assertNotIn("candidate_row", reviewed_origin)
        self.assertEqual(reviewed_origin["review_row"], 1)
        self.assertEqual(reviewed_origin["review_sha256"], review["review_sha256"])
        self.assertEqual(self.store.raw_source_path(result["dataset_id"]).read_text(), "month,value\n2026-01,100\n2026-02,120\n")
        self.assertEqual(self.docs.recover_publication(args, {})["dataset_id"], result["dataset_id"])
        self.assertNotIn("review_table", self.docs.extra_tools())

    def test_actual_http_connection_uses_pinned_public_ip_and_download_size_cap(self):
        class FakeSocket:
            def __init__(self, body, content_length=None):
                self.body = body
                self.content_length = content_length or str(len(body)).encode()
            def makefile(self, *args):
                return io.BytesIO(b"HTTP/1.1 200 OK\r\nContent-Type: text/csv\r\nContent-Length: " + self.content_length + b"\r\n\r\n" + self.body)
            def sendall(self, data):
                pass
            def close(self):
                pass
        addresses = [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 80))]
        with patch("agentic_analytics.agent.tools.documents.socket.getaddrinfo", return_value=addresses), patch("agentic_analytics.agent.tools.documents.socket.create_connection", return_value=FakeSocket(b"x\n1\n")) as connect:
            body, mime, _ = fetch_public_url("http://public.test/data.csv")
            self.assertEqual(body, b"x\n1\n")
            self.assertEqual(mime, "text/csv")
            self.assertEqual(connect.call_args.args[0], ("93.184.216.34", 80))
        with patch("agentic_analytics.agent.tools.documents.socket.getaddrinfo", return_value=addresses), patch(
                "agentic_analytics.agent.tools.documents.socket.create_connection",
                return_value=FakeSocket(b"x\n1\n", b"4        ")):
            body, mime, _ = fetch_public_url("http://public.test/data.csv")
            self.assertEqual(body, b"x\n1\n")
            self.assertEqual(mime, "text/csv")
        with patch("agentic_analytics.agent.tools.documents.socket.getaddrinfo", return_value=addresses), patch("agentic_analytics.agent.tools.documents.socket.create_connection", return_value=FakeSocket(b"123456")):
            with self.assertRaisesRegex(DocumentError, "size limit"):
                fetch_public_url("http://public.test/data.csv", max_bytes=3)

    def test_raw_non_ascii_url_is_percent_encoded_before_the_request_line(self):
        # Search providers and discovered links can hand back an unescaped
        # Turkish path (e.g. a Wikipedia URL). http.client would otherwise
        # raise UnicodeEncodeError trying to send the request line as ASCII.
        self.assertEqual(_ascii_safe_url("https://tr.wikipedia.org/wiki/İhlas_Finans"),
                          "https://tr.wikipedia.org/wiki/%C4%B0hlas_Finans")
        self.assertEqual(_ascii_safe_url("https://example.com/already%20encoded?x=1&y=2"),
                          "https://example.com/already%20encoded?x=1&y=2")

        class FakeSocket:
            def __init__(self):
                self.sent = b""
            def makefile(self, *args):
                return io.BytesIO(b"HTTP/1.1 200 OK\r\nContent-Type: text/html\r\nContent-Length: 2\r\n\r\nok")
            def sendall(self, data):
                self.sent += data
            def close(self):
                pass
        addresses = [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 80))]
        sock = FakeSocket()
        with patch("agentic_analytics.agent.tools.documents.socket.getaddrinfo", return_value=addresses), \
                patch("agentic_analytics.agent.tools.documents.socket.create_connection", return_value=sock):
            body, _, final_url = fetch_public_url("http://public.test/wiki/İhlas_Finans")
        self.assertEqual(body, b"ok")
        self.assertIn(b"GET /wiki/%C4%B0hlas_Finans HTTP", sock.sent)
        self.assertEqual(final_url, "http://public.test/wiki/%C4%B0hlas_Finans")

    def test_monetary_alias_publication_canonicalizes_metadata_without_changing_values(self):
        import pandas as pd
        source = self.upload("unit.csv", "month,value (million TL)\n2026-01,100\n2026-02,120\n")
        contract = copy.deepcopy(self.contract)
        contract["columns"]["value"].update(kind="stock", aggregation="period_end_stock")
        result = self.docs.publish_selected_table(source["source_id"], "table_001", contract, 0,
            column_mapping={"month": "month", "value_million_TL": "value"}, unit_evidence={"value": "million TL"})
        spec = self.store.dataset_manifest(result["dataset_id"])["contract"]["columns"]["value"]
        self.assertEqual((spec["unit"], spec["scale"], spec["currency"], spec["aggregation"]), ("TRY", 1000000, "TRY", "last"))
        self.assertEqual(pd.read_parquet(self.store.overlay_path(result["dataset_id"]))["value"].tolist(), [100, 120])
        self.assertFalse(result["provenance"]["unit_normalization"]["value"]["values_changed"])

    def test_conflicting_units_scales_and_semantics_block_before_publication(self):
        from agentic_analytics.lakehouse.units import normalize_column, unit_quote_matches
        base = self.contract["columns"]["value"]
        for update in [{"scale": 1e6}, {"currency": "USD"}, {"scale": 0}, {"scale": -1},
                       {"scale": True}, {"scale": float("inf")}, {"kind": "made_up"},
                       {"kind": "stock", "aggregation": "sum"}]:
            with self.subTest(update=update), self.assertRaises(ValueError):
                normalize_column({**base, **update})
        for alias, canonical, scale in [("milyar TL", "TRY", 1e9), ("USD_million", "USD", 1e6), ("bin EUR", "EUR", 1000)]:
            value, _ = normalize_column({**base, "unit": alias})
            self.assertEqual((value["unit"], value["scale"], value["currency"]), (canonical, scale, canonical))
        self.assertFalse(unit_quote_matches("TRY", 1, "million TL"))
        self.assertFalse(unit_quote_matches("TRY", 1e6, "billion TL"))
        self.assertTrue(unit_quote_matches("USD", 1e6, "million USD"))

    def test_long_pdf_selects_pages_and_retains_stable_candidate_ids(self):
        def page(number):
            item = MagicMock()
            item.extract_text.return_value = f"Page {number} million TL"
            item.extract_tables.return_value = [[['month', 'value'], [f'2026-{number:02d}', str(number)]]]
            return item
        pdf = MagicMock()
        pdf.pages = [page(number) for number in range(1, 32)]
        with patch("pdfplumber.open") as opened:
            opened.return_value.__enter__.return_value = pdf
            source = self.upload("long.pdf", _text_pdf())
            self.assertEqual(source["total_pages"], 31)
            self.assertFalse(source["inspection_complete"])
            self.assertEqual(source["processed_pages"], list(range(1, 31)))
            old_id = source["tables"][0]["table_id"]
            later = self.docs.inspect_source(source_id=source["source_id"], page_numbers=[31])
            self.assertEqual(later["processed_pages"], [31])
            self.assertEqual(later["tables"][0]["table_id"], "table_p000031_001")
            self.assertEqual(self.docs.review_candidate(source["source_id"], old_id)["rows"][0][1], "1")
            for pages in [[0], [1, 1], [32], list(range(1, 32))]:
                with self.subTest(pages=pages), self.assertRaises(DocumentError):
                    self.docs.inspect_source(source_id=source["source_id"], page_numbers=pages)

    def test_explicit_pdf_continuation_preserves_page_rows_and_rejects_mismatched_headers(self):
        pdf = MagicMock()
        pdf.pages = []
        for number in range(1, 4):
            page = MagicMock()
            page.extract_text.return_value = "million TL"
            page.extract_tables.return_value = [[['month', 'value' if number < 3 else 'other'], [f'2026-0{number}', str(number)]]]
            pdf.pages.append(page)
        with patch("pdfplumber.open") as opened:
            opened.return_value.__enter__.return_value = pdf
            source = self.upload("continuation.pdf", _text_pdf())
        ids = [table["table_id"] for table in source["tables"]]
        combined = self.docs.combine_source_tables(source["source_id"], ids[:2], "Same period report and identical unit scope.")
        self.assertEqual(combined["row_count"], 2)
        published = self.docs.publish_selected_table(source["source_id"], combined["table_id"], self.contract, 0,
                                                     unit_evidence={"value": "million TL"})
        self.assertEqual(published["provenance"]["source_pages"], [1, 2])
        self.assertEqual(published["provenance"]["row_origins"][1]["page"], 2)
        with self.assertRaises(DocumentError):
            self.docs.combine_source_tables(source["source_id"], ids[1:], "Different metric header must fail.")

    def test_xlsx_preserves_cached_formula_cells_and_blocks_missing_cache(self):
        import zipfile
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(["month", "value (million TL)"])
        workbook.active.append(["2026-01", "=50+50"])
        workbook.active.append(["2026-02", 120])
        data = io.BytesIO()
        workbook.save(data)
        missing = self.upload("formula-missing.xlsx", data.getvalue())
        self.assertEqual(missing["tables"][0]["missing_formula_cache"], ["B2"])
        args = {"source_id": missing["source_id"], "table_id": "table_001", "contract": self.contract, "expected_version": 0,
                "column_mapping": {"month": "month", "value_million_TL": "value"}, "unit_evidence": {"value": "million TL"}}
        self.assertEqual(self.docs.extra_tools()["publish_selected_table"]["handler"](args)["code"], "FORMULA_VALUES_REVIEW_REQUIRED")
        output = io.BytesIO()
        with zipfile.ZipFile(data) as source_zip, zipfile.ZipFile(output, "w") as target:
            for item in source_zip.infolist():
                body = source_zip.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    body = body.replace(b"<f>50+50</f><v></v>", b"<f>50+50</f><v>100</v>")
                    body = body.replace(b"<f>50+50</f><v/>", b"<f>50+50</f><v>100</v>")
                    body = body.replace(b"<f>50+50</f><v />", b"<f>50+50</f><v>100</v>")
                target.writestr(item, body)
        source = self.upload("formula-cached.xlsx", output.getvalue())
        self.assertEqual(source["tables"][0]["preview"][0]["value_million_TL"], "100")
        self.assertEqual(source["tables"][0]["formula_cells"][0]["verification"], "saved_source_value_not_recalculated")
        self.assertEqual(self.docs.publish_selected_table(**{**args, "source_id": source["source_id"]})["status"], "ok")

    def test_structural_html_multiline_headers_and_merged_numeric_cells(self):
        source = self.upload("multi.html", '''<table><thead><tr><th rowspan="2">metric</th><th colspan="2">million TL</th></tr>
          <tr><th>2024</th><th>2025</th></tr></thead><tbody><tr><td>loans</td><td>100</td><td>120</td></tr></tbody></table>''')
        table = source["tables"][0]
        self.assertEqual(list(table["original_columns"].values()), ["metric", "million TL / 2024", "million TL / 2025"])
        self.assertEqual(list(table["preview"][0].values()), ["loans", "100", "120"])
        self.assertFalse(table["layout_review_required"])
        merged = self.upload("merged.html", '<table><tr><th>month</th><th>value</th><th>other</th></tr><tr><td>2026-01</td><td colspan="2">100</td></tr></table>')
        self.assertTrue(merged["tables"][0]["layout_review_required"])
        self.assertEqual(list(merged["tables"][0]["preview"][0].values()), ["2026-01", "100", None])

    def test_financial_report_unpivot_copies_exact_source_cells_and_publishes(self):
        import pandas as pd
        source = self.upload("wide.csv", "metric (million TL),2024,2025\nloans,100,120\nassets,150,180\n")
        prepared = self.docs.prepare_source_table(source["source_id"], "table_001", selected_rows=[1],
            unpivot={"columns": ["c_2024", "c_2025"], "period_column": "year", "value_column": "value"})
        self.assertEqual([row["value"] for row in prepared["preview"]], ["100", "120"])
        contract = {"name": "loans", "frequency": "annual", "date_column": "year", "key": ["year"], "grain": ["year"],
                    "columns": {"metric_million_TL": {"dtype": "string", "kind": "dimension", "unit": "label", "nullable": False},
                                "year": {"dtype": "date", "kind": "dimension", "unit": "calendar", "nullable": False},
                                "value": {"dtype": "integer", "kind": "stock", "unit": "TRY_million", "nullable": False}}}
        result = self.docs.publish_selected_table(source["source_id"], prepared["table_id"], contract, 0, unit_evidence={"value": "million TL"})
        self.assertEqual(pd.read_parquet(self.store.overlay_path(result["dataset_id"]))["value"].tolist(), [100, 120])
        self.assertEqual(result["provenance"]["cell_origins"][1]["value"], {"candidate_row": 1, "candidate_column": "c_2025"})
        self.assertEqual(self.docs.review_candidate(source["source_id"], "table_001")["row_count"], 2)

    def test_research_follows_report_downloads_and_rejects_official_redirects(self):
        self.docs.web_search = lambda query, limit=5, **kwargs: {"status": "ok", "results": [{"title": "Housing report", "url": "https://tcmb.gov.tr/reports"}]}
        self.docs.inspect_source = lambda url=None, **kwargs: {"source_id": "source_" + "a" * 64, "source_url": url,
            "text": "Housing report contains monthly housing prices.", "article": {"document_links": [{"url": "https://tcmb.gov.tr/data.pdf", "title": "Housing report data"}]} if url.endswith("reports") else {},
            "tables": [] if url.endswith("reports") else [{"table_id": "table_p000001_001", "preview": [{"price": "100"}]}]}
        result = self.docs.research_web("housing report", limit=1, domains=["tcmb.gov.tr"])
        self.assertEqual(result["sources"][0]["url"], "https://tcmb.gov.tr/data.pdf")
        self.assertEqual(result["sources"][0]["discovered_from"], "https://tcmb.gov.tr/reports")
        self.docs.inspect_source = lambda url=None, **kwargs: {"source_id": "source_" + "a" * 64, "source_url": "https://elsewhere.org/report",
            "text": "Housing report contains monthly housing prices.", "article": {}}
        failed = self.docs.research_web("housing report", domains=["tcmb.gov.tr"])
        self.assertEqual(failed["failures"][0]["code"], "OFFICIAL_SOURCE_REDIRECT")

    def test_report_title_header_repair_accounting_values_and_label_lineage(self):
        import pandas as pd
        source = self.upload("financial.csv", "fragment,rest,current (million TL),prior (million TL)\n,,(30/06/2026),(30/06/2025)\nCURRENT PERIOD PRO,FIT/LOSS,34.333,24.850\nOTHER COMPREHENSI,VE INCOME,(8.207),19\n")
        window = self.docs.read_source_table(source["source_id"], "table_001", row_start=2, row_limit=1)
        self.assertEqual(window["rows"][0]["candidate_row"], 2)
        self.assertTrue(window["preview_truncated"])
        prepared = self.docs.prepare_source_table(source["source_id"], "table_001", selected_rows=[2, 3],
            unpivot={"columns": ["current_million_TL", "prior_million_TL"], "period_column": "period_end", "value_column": "amount",
                     "header_row": 1, "period_format": "parenthesized_dmy"},
            join_columns={"columns": ["fragment", "rest"], "output": "metric", "separator": ""})
        self.assertEqual(prepared["preview"][0]["period_end"], "2026-06-30")
        self.assertEqual(prepared["preview"][0]["metric"], "CURRENT PERIOD PROFIT/LOSS")
        contract = {"name": "reported_income", "frequency": "event", "date_column": "period_end", "key": ["period_end", "metric"], "grain": ["period_end", "metric"],
                    "number_format": "decimal_comma", "negative_format": "accounting_parentheses", "columns": {
                        "period_end": {"dtype": "date", "kind": "dimension", "unit": "calendar", "nullable": False},
                        "metric": {"dtype": "string", "kind": "dimension", "unit": "label", "nullable": False},
                        "amount": {"dtype": "integer", "kind": "unknown", "unit": "TRY_million", "nullable": False}}}
        invalid = copy.deepcopy(contract)
        invalid["columns"]["metric"].update(dtype="integer", kind="flow")
        with self.assertRaisesRegex(DocumentError, "must remain string dimensions"):
            self.docs.publish_selected_table(source["source_id"], prepared["table_id"], invalid, 0, unit_evidence={"amount": "million TL"})
        result = self.docs.publish_selected_table(source["source_id"], prepared["table_id"], contract, 0, unit_evidence={"amount": "million TL"})
        frame = pd.read_parquet(self.store.overlay_path(result["dataset_id"]))
        self.assertEqual(sorted(frame["amount"].tolist()), [-8207, 19, 24850, 34333])
        self.assertEqual(result["provenance"]["cell_origins"][0]["period_end"], {"candidate_row": 1, "candidate_column": "prior_million_TL"})
        self.assertEqual(len(result["provenance"]["cell_origins"][0]["metric"]["parts"]), 2)
        self.assertEqual(result["provenance"]["row_order"]["stored_row_to_source_csv_row"], [2, 4, 1, 3])
        raw = self.docs.review_candidate(source["source_id"], "table_001")
        for record, origin in zip(frame.to_dict("records"), result["provenance"]["cell_origins"]):
            date_cell = origin["period_end"]
            source_date = raw["rows"][date_cell["candidate_row"] - 1][raw["columns"].index(date_cell["candidate_column"])]
            self.assertEqual(self.docs._source_period_label(source_date, "parenthesized_dmy"), record["period_end"])
            amount_cell = origin["amount"]
            value = raw["rows"][amount_cell["candidate_row"] - 1][raw["columns"].index(amount_cell["candidate_column"])]
            parsed = -int(value[1:-1].replace(".", "")) if value.startswith("(") else int(value.replace(".", ""))
            self.assertEqual(parsed, record["amount"])

    def test_pdf_alternate_table_strategy_does_not_replace_original_candidates(self):
        pdf, page = MagicMock(), MagicMock()
        page.extract_text.return_value = "monthly value million TL"
        page.extract_tables.side_effect = lambda *args: [[['month', 'value'], ['2026-01', '100' if not args else '120']]]
        pdf.pages = [page]
        with patch("pdfplumber.open") as opened:
            opened.return_value.__enter__.return_value = pdf
            source = self.upload("strategies.pdf", _text_pdf())
            alternate = self.docs.inspect_source(source_id=source["source_id"], page_numbers=[1], table_strategy="text")
        self.assertNotEqual(source["tables"][0]["table_id"], alternate["tables"][0]["table_id"])
        self.assertEqual(self.docs.review_candidate(source["source_id"], source["tables"][0]["table_id"])["rows"][0][1], "100")
        self.assertEqual(alternate["tables"][0]["preview"][0]["value"], "120")

    def test_research_traverses_archive_and_never_uses_language_root_as_answer(self):
        self.docs.web_search = lambda query, limit=5, **kwargs: {"status": "ok", "results": [{"title": "Example Bank", "url": "https://example.org/en"}]}
        def inspect(url=None, **kwargs):
            body = {"source_id": "source_" + "a" * 64, "source_url": url, "text": "Annual financial results 2025", "tables": []}
            if url.endswith("/en"):
                body["article"] = {"title": "Example Bank", "source_links": [{"url": "https://example.org/reports/2025", "title": "Annual financial results 2025"}], "link_count": 50}
            else:
                body["article"] = {"title": "Annual financial results 2025", "article_body": "Annual financial results 2025: reported profit 100 million USD."}
            return body
        self.docs.inspect_source = inspect
        result = self.docs.research_web("financial results 2025", limit=1, domains=["example.org"])
        self.assertEqual(result["sources"][0]["url"], "https://example.org/reports/2025")

    def test_unit_evidence_cannot_be_borrowed_from_another_currency_column(self):
        source = self.upload("mixed-units.csv", "month,asset (TL),debt (USD)\n2026-01,10,20\n2026-02,30,40\n")
        contract = {"name": "mixed", "frequency": "monthly", "date_column": "month", "key": ["month"], "grain": ["month"],
                    "columns": {"month": self.contract["columns"]["month"],
                                "asset_TL": {"dtype": "integer", "kind": "stock", "unit": "TRY", "nullable": False},
                                "debt_USD": {"dtype": "integer", "kind": "stock", "unit": "TRY", "nullable": False}}}
        result = self.docs.extra_tools()["publish_selected_table"]["handler"]({"source_id": source["source_id"], "table_id": "table_001", "contract": contract,
            "expected_version": 0, "unit_evidence": {"asset_TL": "TL", "debt_USD": "TL"}})
        self.assertEqual(result["code"], "COLUMN_UNIT_CONFLICT")
        self.assertEqual(self.store.workspace("workspace_docs")["version"], 0)

    def test_nonmonetary_multiplier_requires_source_evidence(self):
        from agentic_analytics.lakehouse.units import unit_quote_matches
        self.assertFalse(unit_quote_matches("count", 1000, "count"))
        self.assertFalse(unit_quote_matches("persons", 1000, "persons"))
        self.assertTrue(unit_quote_matches("persons", 1000, "thousand persons"))

    def test_joined_label_cannot_be_laundered_to_a_number_by_preparing_twice(self):
        source = self.upload("labels.csv", "left,right\n10,20\n")
        first = self.docs.prepare_source_table(source["source_id"], "table_001", join_columns={"columns": ["left", "right"], "output": "label", "separator": ""})
        second = self.docs.prepare_source_table(source["source_id"], first["table_id"])
        self.assertEqual(second["dimension_only_columns"], ["label"])
        contract = {"name": "invalid", "frequency": "static", "key": ["label"], "grain": ["label"],
                    "columns": {"label": {"dtype": "integer", "kind": "count", "unit": "count", "nullable": False}}}
        with self.assertRaisesRegex(DocumentError, "must remain string dimensions"):
            self.docs.publish_selected_table(source["source_id"], second["table_id"], contract, 0)

    def test_explicit_cumulative_column_header_is_preserved_even_when_model_omits_flag(self):
        source = self.upload("ytd.csv", "month,YTD profit (million TL)\n2026-01,100\n2026-02,120\n")
        result = self.docs.publish_selected_table(source["source_id"], "table_001", self.contract, 0,
            column_mapping={"month": "month", "YTD_profit_million_TL": "value"}, unit_evidence={"value": "million TL"})
        spec = self.store.dataset_manifest(result["dataset_id"])["contract"]["columns"]["value"]
        self.assertEqual(spec["temporal_semantics"], "year_to_date_flow")
        self.assertIn("YTD profit", spec["source_semantics"])
        self.assertEqual(spec["aggregation"], "none")

    def test_ocr_missing_item_code_header_is_reviewable_and_preserves_every_cell(self):
        from PIL import Image
        data = io.BytesIO()
        Image.new("RGB", (30, 30), "white").save(data, format="PNG")
        raw = {"text": "Amounts are expressed in thousands of TRY.", "tables": [{"columns": ["Assets", "Current", "Prior"],
            "rows": [["I.", "Cash", "100", "90"], ["1.1", "Deposits", "200", "180"]]}]}
        self.docs.ocr_callback = lambda *args: raw
        source = self.upload("dense.png", data.getvalue())
        table = source["tables"][0]
        self.assertEqual(table["columns"][0], "source_item_code")
        self.assertTrue(table["header_hypothesis"]["requires_independent_review"])
        self.assertEqual(table["preview"][0]["Current"], "100")
        full = self.docs.review_candidate(source["source_id"], table["table_id"])
        self.assertEqual(full["raw_machine_rows"], raw["tables"][0]["rows"])
        evidence = json.loads((self.docs.root / "extractions" / (table["extraction_artifact_ref"] + ".json")).read_text())
        self.assertEqual(evidence["raw_machine_output"], raw)
        self.assertFalse(evidence["verified"])
        with self.assertRaisesRegex(DocumentError, "Review uncertain extraction"):
            self.docs.prepare_source_table(source["source_id"], table["table_id"])

    def test_irregular_ocr_rows_are_retained_without_padding_or_dropping_numbers(self):
        from PIL import Image
        data = io.BytesIO()
        Image.new("RGB", (30, 30), "white").save(data, format="PNG")
        raw = [["one", "100", "999"], ["two", "200"]]
        self.docs.ocr_callback = lambda *args: {"text": "Amounts in TRY.", "tables": [{"columns": ["item", "amount"], "rows": raw}]}
        source = self.upload("ragged.png", data.getvalue())
        self.assertEqual(source["tables"][0]["preview"], [])
        self.assertEqual(source["tables"][0]["raw_preview"], raw)
        window = self.docs.read_source_table(source["source_id"], "table_001")
        self.assertIsNone(window["rows"][0]["values"])
        self.assertEqual(window["rows"][0]["raw_cells"], raw[0])

    def test_incomplete_ocr_retains_evidence_reference(self):
        from PIL import Image
        data = io.BytesIO()
        Image.new("RGB", (30, 30), "white").save(data, format="PNG")
        self.docs.ocr_callback = lambda *args: {"finish_reason": "length", "text": "incomplete", "tables": []}
        path = self.docs.upload_root / "length.png"
        path.write_bytes(data.getvalue())
        source = self.docs.register_upload(path)
        result = self.docs.extra_tools()["inspect_source"]["handler"]({"source_id": source["source_id"]})
        self.assertEqual(result["code"], "OCR_INVALID_OUTPUT")
        self.assertTrue((self.docs.root / "extractions" / (result["artifact_ref"] + ".json")).exists())

    def test_merged_period_header_cells_are_explicitly_mapped_without_numeric_changes(self):
        source = self.upload("merged.csv", "metric,current_date,current_total,prior_day,prior_month_year,prior_total\n,31 March 2026,,31,December 2025,\ncash,,867799356,,,1005229845\nassets,,1279933463,,,1260568409\n")
        arguments = {"source_id": source["source_id"], "table_id": "table_001", "selected_rows": [2, 3],
            "selected_columns": ["metric", "current_total", "prior_total"], "unpivot": {
                "columns": ["current_total", "prior_total"], "period_column": "date", "value_column": "amount",
                "period_format": "english_dmy", "period_sources": {
                    "current_total": {"row": 1, "columns": ["current_date"], "separator": " "},
                    "prior_total": {"row": 1, "columns": ["prior_day", "prior_month_year"], "separator": " "}}}}
        prepared = self.docs.prepare_source_table(**arguments)
        self.assertEqual([row["date"] for row in prepared["preview"]], ["2026-03-31", "2025-12-31"] * 2)
        self.assertEqual([row["amount"] for row in prepared["preview"]], ["867799356", "1005229845", "1279933463", "1260568409"])
        full = self.docs.review_candidate(source["source_id"], prepared["table_id"])
        self.assertEqual(full["cell_origins"][1]["date"]["parts"], [{"candidate_row": 1, "candidate_column": "prior_day"}, {"candidate_row": 1, "candidate_column": "prior_month_year"}])
        bad = copy.deepcopy(arguments)
        bad["unpivot"]["period_sources"]["prior_total"]["columns"] = ["invented_header"]
        with self.assertRaises(DocumentError):
            self.docs.prepare_source_table(**bad)
        bad = copy.deepcopy(arguments)
        bad["unpivot"]["period_sources"]["prior_total"]["row"] = 2
        with self.assertRaises(DocumentError):
            self.docs.prepare_source_table(**bad)

    def test_inflation_adjusted_caption_requires_preserved_price_basis(self):
        source = self.upload("adjusted.html", '<p>Amounts expressed in thousands of Turkish Lira in terms of purchasing power at 30 June 2025.</p><table><tr><th>item</th><th>amount</th></tr><tr><td>cash</td><td>100</td></tr></table>')
        contract = {"name": "adjusted", "frequency": "static", "key": ["item"], "grain": ["item"], "columns": {
            "item": {"dtype": "string", "unit": "label", "kind": "dimension", "nullable": False},
            "amount": {"dtype": "integer", "unit": "TRY", "scale": 1000, "kind": "stock", "nullable": False}}}
        arguments = {"source_id": source["source_id"], "table_id": "table_001", "contract": contract,
                     "expected_version": 0, "unit_evidence": {"amount": "thousands of Turkish Lira"}}
        result = self.docs.extra_tools()["publish_selected_table"]["handler"](arguments)
        self.assertEqual(result["code"], "PRICE_BASIS_REVIEW_REQUIRED")
        contract["columns"]["amount"]["price_basis"] = "2025-06-30 purchasing power"
        result = self.docs.publish_selected_table(**arguments)
        self.assertEqual(result["status"], "ok")
        spec = self.store.dataset_manifest(result["dataset_id"])["contract"]["columns"]["amount"]
        self.assertIn("30 June 2025", spec["price_basis_evidence"])

    def test_dated_research_prioritizes_exact_body_link_and_skips_wrong_year_archives(self):
        root, archive, target = "https://example.org/", "https://example.org/archive2025", "https://example.org/press/2025/decision15"
        self.docs.web_search = lambda *args, **kwargs: {"status": "ok", "results": [{"title": "Central bank", "url": root}]}
        visited = []
        def inspect(url=None, **kwargs):
            visited.append(url)
            if url == root:
                links = [{"url": "https://example.org/press/2026/decision", "title": "2026 monetary policy interest rate decision"},
                         {"url": archive, "title": "Monetary policy meetings 2025", "in_main_content": True}]
                return {"source_id": "root", "source_url": url, "text": "central bank", "article": {"source_links": links}}
            if url == archive:
                links = [{"url": "https://example.org/interest-rates/history", "title": "Policy rates monetary policy interest repo auctions"},
                         {"url": target, "title": "6 Mart 2025", "in_main_content": True}]
                return {"source_id": "archive", "source_url": url, "text": "2025 monetary policy meeting calendar", "article": {"title": "2025 monetary policy meetings", "source_links": links, "link_count": 150}}
            if url != target:
                raise AssertionError("Navigation or a wrong-year source was fetched before the exact decision.")
            return {"source_id": "decision", "source_url": url, "text": "6 March 2025 monetary policy committee lowered the repo interest rate from 45 to 42.5.", "article": {"title": "Monetary policy decision 6 March 2025"}}
        self.docs.inspect_source = inspect
        result = self.docs.research_web("Monetary policy committee 6 March 2025 interest rate decision weekly repo auctions", limit=1, domains=["example.org"])
        self.assertEqual(result["sources"][0]["url"], target)
        self.assertEqual(visited, [root, archive, target])

    def test_official_financial_report_archive_is_seeded_even_when_search_returns_a_presentation(self):
        presentation = "https://www.garantibbvainvestorrelations.com/en/images/pdf/1Q26_Financial_Results.pdf"
        archive = "https://www.garantibbvainvestorrelations.com/en/library/brsa-consolidated-financials-pdf/PDF/1268/0/0"
        target = "https://www.garantibbvainvestorrelations.com/en/images/pdf/31_March_2026_Consolidated_Financial_Report.pdf"
        self.docs.web_search = lambda *args, **kwargs: {"status": "ok", "results": [{
            "title": "31 March 2026 financial results presentation", "url": presentation}]}
        visited = []

        def inspect(url=None, **kwargs):
            visited.append(url)
            if url == presentation:
                return {"status": "ok", "source_id": "presentation", "source_url": url,
                    "raw_sha256": "1" * 64, "text": "Garanti BBVA 31 March 2026 financial results presentation",
                    "article": {"title": "31 March 2026 Financial Results Presentation"}}
            if url == archive:
                decoys = [{"url": f"https://www.garantibbvainvestorrelations.com/en/images/pdf/report_{index}_2026.pdf",
                            "title": "2026 consolidated financial report", "in_main_content": True}
                           for index in range(12)]
                return {"status": "ok", "source_id": "archive", "source_url": url,
                    "raw_sha256": "2" * 64, "text": "Garanti BBVA consolidated financial reports archive",
                    "article": {"title": "Consolidated Financial Reports", "source_links": [*decoys, {
                        "url": target, "title": "2026/1Q", "in_main_content": True}], "link_count": 375}}
            if url == target:
                return {"status": "ok", "source_id": "target", "source_url": url,
                    "raw_sha256": "3" * 64,
                    "text": "Garanti BBVA consolidated financial report. TOTAL ASSETS 4,783,750,292. Amounts in thousands of Turkish Lira.",
                    "article": {"title": "31 March 2026 Consolidated Financial Report"}}
            raise DocumentError("Decoy should not be fetched before the exact dated report", "WRONG_CANDIDATE")

        self.docs.inspect_source = inspect
        result = self.docs.research_web("Garanti BBVA 31 Mart 2026 konsolide finansal raporu", limit=1)
        self.assertEqual(result["status"], "ok", result)
        self.assertEqual(result["sources"][0]["url"], target)
        self.assertEqual(result["sources"][0]["document_type"], "financial_report")
        self.assertEqual(result["sources"][0]["reporting_period"], "2026-03-31")
        self.assertEqual(result["sources"][0]["consolidation_scope"], "consolidated")
        self.assertEqual(visited[:3], [presentation, archive, target])

    def test_tcmb_policy_decision_archive_is_seeded_and_meeting_summary_is_rejected(self):
        summary = "https://www.tcmb.gov.tr/meeting-summary-2025-15"
        archive = "https://www.tcmb.gov.tr/wps/wcm/connect/TR/TCMB+TR/Main+Menu/Duyurular/Basin/2025"
        decision = archive + "/DUY2025-15"
        self.docs.web_search = lambda *args, **kwargs: {"status": "ok", "results": [{
            "title": "6 Mart 2025 Para Politikası Kurulu Toplantı Özeti", "url": summary}]}
        visited = []

        def inspect(url=None, **kwargs):
            visited.append(url)
            if url == summary:
                return {"status": "ok", "source_id": "summary", "source_url": url, "raw_sha256": "4" * 64,
                    "text": "6 Mart 2025 Para Politikası Kurulu toplantı özeti ve politika değerlendirmesi",
                    "article": {"title": "Para Politikası Kurulu Toplantı Özeti"}}
            if url == archive:
                return {"status": "ok", "source_id": "archive", "source_url": url, "raw_sha256": "5" * 64,
                    "text": "TCMB 2025 basın duyuruları", "article": {"title": "2025 Basın Duyuruları",
                    "source_links": [{"url": decision, "title": "Faiz Oranlarına İlişkin Basın Duyurusu (2025-15)",
                                      "in_main_content": True}], "link_count": 150}}
            if url == decision:
                return {"status": "ok", "source_id": "decision", "source_url": url, "raw_sha256": "6" * 64,
                    "text": "6 Mart 2025. Para Politikası Kurulu politika faizi olan bir hafta vadeli repo ihale faiz oranının yüzde 45'ten yüzde 42,5'e indirilmesine karar vermiştir.",
                    "article": {"title": "Faiz Oranlarına İlişkin Basın Duyurusu (2025-15)"}}
            raise DocumentError("Unexpected source", "WRONG_CANDIDATE")

        self.docs.inspect_source = inspect
        result = self.docs.research_web("TCMB 6 Mart 2025 Para Politikası Kurulu faiz kararı", limit=1)
        self.assertEqual(result["status"], "ok", result)
        self.assertEqual(result["sources"][0]["url"], decision)
        self.assertEqual(result["sources"][0]["document_type"], "policy_decision")
        self.assertEqual(visited[:3], [summary, archive, decision])
        self.assertTrue(any(failure["code"] == "DOCUMENT_TYPE_MISMATCH" for failure in result["failures"]))

    def test_requested_policy_decision_rejects_unclassified_same_date_document(self):
        wrong = "https://example.org/2025-policy-program.pdf"
        self.docs.web_search = lambda *args, **kwargs: {"status": "ok", "results": [
            {"title": "2025 monetary policy programme", "url": wrong},
        ]}
        self.docs.inspect_source = lambda url=None, **kwargs: (
            {"source_id": "wrong", "source_url": url,
             "text": "2025 monetary policy programme. Calendar item: 6 March 2025.",
             "article": {"title": "2025 Monetary Policy Programme"}}
        )
        result = self.docs.research_web(
            "Monetary policy decision 6 March 2025 interest rate",
            limit=1,
            domains=["example.org"],
        )
        self.assertEqual("unavailable", result["status"])
        self.assertEqual([], result["sources"])
        self.assertTrue(any(
            failure["url"] == wrong and failure["code"] == "DOCUMENT_TYPE_MISMATCH"
            for failure in result["failures"]
        ))

    def test_document_type_uses_leading_official_title_before_navigation_links(self):
        decision = (
            "TCMB - Faiz Oranlarına İlişkin Basın Duyurusu (2025-15) "
            "Para Politikası Kurulu Toplantı Özeti bağlantısı"
        )
        summary = (
            "TCMB - Para Politikası Kurulu Toplantı Özeti "
            "Faiz Oranlarına İlişkin Basın Duyurusu bağlantısı"
        )
        self.assertEqual("policy_decision", _document_type(decision))
        self.assertEqual("meeting_summary", _document_type(summary))
        self.assertEqual("financial_report", _document_type("Example Bank financial report"))

    def test_url_underscore_date_establishes_exact_financial_reporting_period(self):
        url = "https://example.org/31_March_2026_Consolidated_Financial_Report.pdf"
        self.docs.web_search = lambda *args, **kwargs: {"status": "ok", "results": [{"title": "2026/1Q", "url": url}]}
        self.docs.inspect_source = lambda **kwargs: {"status": "ok", "source_id": "report", "source_url": url,
            "raw_sha256": "7" * 64, "text": "Example Bank consolidated financial report. TOTAL ASSETS.",
            "article": {"title": "2026/1Q Consolidated Financial Report"}}
        result = self.docs.research_web("Example Bank 31 March 2026 consolidated financial report", limit=1,
                                        domains=["example.org"])
        self.assertEqual(result["status"], "ok", result)
        self.assertEqual(result["sources"][0]["reporting_period"], "2026-03-31")

    def test_research_deduplicates_identical_document_bytes_across_urls(self):
        urls = ["https://example.org/report", "https://example.org/report-copy"]
        self.docs.web_search = lambda *args, **kwargs: {"status": "ok", "results": [
            {"title": "Official report 2026", "url": url} for url in urls]}
        self.docs.inspect_source = lambda url=None, **kwargs: {"status": "ok", "source_id": url.rsplit("/", 1)[-1],
            "source_url": url, "raw_sha256": "8" * 64, "text": "Official report 2026 verified source text",
            "article": {"title": "Official report 2026"}}
        result = self.docs.research_web("Official report 2026", limit=2, domains=["example.org"])
        self.assertEqual(len(result["sources"]), 1)
        self.assertTrue(any(failure["code"] == "DUPLICATE_SOURCE_CONTENT" for failure in result["failures"]))

    def test_find_source_table_rows_locates_a_row_beyond_the_preview(self):
        rows = ["line,value", *[f"Other line {index},{index}" for index in range(1, 12)],
                "TOTAL ASSETS,4783750292", "Other line 13,13"]
        source = self.upload("long-financial-table.csv", "\n".join(rows) + "\n")
        table_id = source["tables"][0]["table_id"]
        self.assertTrue(source["tables"][0]["preview_truncated"])
        result = self.docs.extra_tools()["find_source_table_rows"]["handler"]({
            "source_id": source["source_id"], "table_id": table_id, "query": "total assets"})
        self.assertEqual(result["status"], "ok", result)
        self.assertEqual(result["total_matches"], 1)
        self.assertEqual(result["rows"][0]["candidate_row"], 12)
        self.assertEqual(result["rows"][0]["values"]["value"], "4783750292")
        self.assertEqual(result["raw_sha256"], source["raw_sha256"])


    def test_year_only_navigation_link_inherits_the_parent_archive_topic(self):
        urls = ["https://example.org/", "https://example.org/meetings", "https://example.org/meetings/2025", "https://example.org/decisions/march"]
        self.docs.web_search = lambda *args, **kwargs: {"status": "ok", "results": [{"title": "Central Bank", "url": urls[0]}]}
        visited = []
        def inspect(url=None, **kwargs):
            visited.append(url)
            index = urls.index(url)
            result = {"source_id": str(index), "source_url": url, "text": "Monetary policy interest rate decisions", "article": {"title": "Monetary policy decisions", "source_links": []}}
            if index == 0:
                result["article"]["source_links"] = [{"url": urls[1], "title": "Monetary policy decisions", "in_navigation": True}]
            elif index == 1:
                result["article"]["source_links"] = [{"url": urls[2], "title": "2025", "in_navigation": True}]
            elif index == 2:
                result["article"].update(title="2025 decisions calendar", link_count=150, source_links=[{"url": urls[3], "title": "6 March 2025", "in_main_content": True}])
            else:
                result["article"].update(title="Interest rate decision 6 March 2025", article_body="Monetary policy interest rate decision on 6 March 2025: 42.5 percent.")
            return result
        self.docs.inspect_source = inspect
        result = self.docs.research_web("Monetary policy 6 March 2025 interest rate decision", limit=1, domains=["example.org"])
        self.assertEqual(result["sources"][0]["url"], urls[-1])
        self.assertEqual(visited, urls)

    def test_html_visual_line_breaks_never_concatenate_financial_amounts(self):
        for index, cells in enumerate(["100<br>200", "<div>100</div><div>200</div>", "<p>100</p><p>200</p>"]):
            with self.subTest(cells=cells):
                source = self.upload(f"broken-amount-{index}.html", '<table><tr><th>month</th><th>amount (TRY)</th></tr><tr><td>2026-01</td><td>' + cells + '</td></tr></table>')
                candidate = source["tables"][0]
                self.assertEqual(candidate["preview"][0]["amount_TRY"], "100\n200")
                self.assertTrue(candidate["layout_review_required"])
                contract = copy.deepcopy(self.contract)
                contract["columns"]["value"]["unit"] = "TRY"
                result = self.docs.extra_tools()["publish_selected_table"]["handler"]({"source_id": source["source_id"], "table_id": "table_001", "contract": contract,
                    "expected_version": 0, "column_mapping": {"month": "month", "amount_TRY": "value"}, "unit_evidence": {"value": "TRY"}})
                self.assertEqual(result["code"], "TABLE_LAYOUT_REVIEW_REQUIRED")
        source = self.upload("header-break.html", '<table><tr><th>metric</th><th>31 December<br>2025</th></tr><tr><td>cash</td><td>100</td></tr></table>')
        self.assertEqual(source["tables"][0]["original_columns"]["c_31_December_2025"], "31 December\n2025")

    def test_search_result_shapes_are_checked_before_reading_urls(self):
        self.docs.searxng_url = "http://localhost:8080"
        with patch("agentic_analytics.agent.tools.search_backend.configured_search", return_value=[None]), \
                patch("agentic_analytics.agent.tools.documents.fetch_public_url", side_effect=OSError('Index unavailable')):
            result = self.docs.web_search("financial report")
        self.assertEqual(result["provider_attempts"][0]["code"], "SEARCH_INVALID_RESPONSE")
        self.assertEqual(result["results"], [])
        with patch("agentic_analytics.agent.tools.search_backend.configured_search", return_value=[None, {"url": []}, {"url": "https://example.org/report", "title": "Report"}]):
            result = self.docs.web_search("financial report", limit=1)
        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["results"][0]["url"], "https://example.org/report")
        self.assertEqual(result["warnings"][0]["count"], 2)

    def test_research_accepts_earnings_wording_for_financial_results_without_year_only_match(self):
        self.docs.web_search = lambda *args, **kwargs: {"status": "ok", "results": [{"title": "Earnings presentation", "url": "https://example.org/earnings.pdf"}]}
        self.docs.inspect_source = lambda **kwargs: {"source_id": "report", "source_url": kwargs["url"], "text": "Earnings Presentation\n2026 H1\nMacro Financial\nOutlook Performance"}
        result = self.docs.research_web("2026 financial results", limit=1, domains=["example.org"])
        self.assertEqual(result["sources"][0]["url"], "https://example.org/earnings.pdf")
        self.docs.inspect_source = lambda **kwargs: {"source_id": "unrelated", "source_url": kwargs["url"], "text": "2026 Financial calendar dates"}
        result = self.docs.research_web("2026 financial results", limit=1, domains=["example.org"])
        self.assertEqual(result["status"], "unavailable")

    def test_complete_inverse_column_mapping_is_normalized_with_source_proof(self):
        source = self.upload("inverse.csv", "raw_month,raw_amount (million TL)\n2026-01,100\n2026-02,120\n")
        arguments = {"source_id": source["source_id"], "table_id": "table_001", "contract": self.contract,
                     "expected_version": 0, "column_mapping": {"month": "raw_month", "value": "raw_amount_million_TL"},
                     "unit_evidence": {"value": "million TL"}}
        published = self.docs.publish_selected_table(**arguments)
        proof = published["provenance"]["column_mapping"]
        self.assertEqual(proof["direction"], "unambiguous_output_to_source_normalized")
        self.assertEqual(proof["declared"], arguments["column_mapping"])
        self.assertEqual(proof["source_to_output"], {"raw_month": "month", "raw_amount_million_TL": "value"})
        self.assertFalse(proof["source_values_changed"])
        recovered = self.docs.recover_publication(arguments, {})
        self.assertEqual(recovered["dataset_id"], published["dataset_id"])
        self.assertEqual(self.store.workspace("workspace_docs")["version"], 1)

    def test_incomplete_or_many_to_one_mappings_explain_actual_columns_without_writing(self):
        source = self.upload("mapping.csv", "raw_month,raw_amount (million TL),suffix\n2026-01,100,profit\n2026-02,120,profit\n")
        for mapping in [{"month": "raw_month", "value": "raw_amount_million_TL"}, {"month": "raw_month", "value": "raw_month"}, {}]:
            with self.subTest(mapping=mapping):
                result = self.docs.extra_tools()["publish_selected_table"]["handler"]({"source_id": source["source_id"], "table_id": "table_001",
                    "contract": self.contract, "expected_version": 0, "column_mapping": mapping, "unit_evidence": {"value": "million TL"}})
                self.assertEqual(result["code"], "INVALID_COLUMN_MAPPING")
                self.assertIn('"source_columns": ["raw_month", "raw_amount_million_TL", "suffix"]', result["message"])
                self.assertIn("prepare_source_table", result["message"])
        self.assertEqual(self.store.workspace("workspace_docs")["version"], 0)

    def test_multiple_dates_split_across_header_cells_use_source_date_occurrence(self):
        source = self.upload("split-dates.csv", "metric,left,current,prior\n,3,0 June 2025 31 De,cember 2024\ncash,,90351730,85795568\nassets,,545630257,529848729\n")
        arguments = {"source_id": source["source_id"], "table_id": "table_001", "selected_rows": [2, 3],
                     "selected_columns": ["metric", "current", "prior"], "unpivot": {
                         "columns": ["current", "prior"], "period_column": "as_of", "value_column": "amount", "period_format": "english_dmy",
                         "period_sources": {column: {"row": 1, "columns": ["left", "current", "prior"], "separator": "", "date_index": index}
                                            for index, column in enumerate(["current", "prior"])}}}
        result = self.docs.prepare_source_table(**arguments)
        self.assertEqual([(row["as_of"], row["amount"]) for row in result["preview"]],
                         [("2025-06-30", "90351730"), ("2024-12-31", "85795568"), ("2025-06-30", "545630257"), ("2024-12-31", "529848729")])
        full = self.docs.review_candidate(source["source_id"], result["table_id"])
        proof = full["cell_origins"][1]["as_of"]
        self.assertEqual(proof["joined_source_text"], "30 June 2025 31 December 2024")
        self.assertEqual(proof["matched_source_date"], "31 December 2024")
        self.assertEqual(proof["date_index"], 1)
        self.assertEqual(len(proof["parts"]), 3)
        bad = copy.deepcopy(arguments)
        bad["unpivot"]["period_sources"]["prior"]["date_index"] = 2
        with self.assertRaisesRegex(DocumentError, "actual complete date"):
            self.docs.prepare_source_table(**bad)
        bad["unpivot"]["period_format"] = "source_header"
        with self.assertRaisesRegex(DocumentError, "explicit period_format"):
            self.docs.prepare_source_table(**bad)
        self.assertEqual(result["publication_guidance"]["candidate_columns"], ["metric", "as_of", "amount"])

    def test_explicit_low_source_limit_still_bounds_upload_and_url_reader(self):
        limited = DocumentTools(self.store, "workspace_docs", max_source_bytes=12)
        path = limited.upload_root / "small.txt"
        path.write_bytes(b"x" * 13)
        with self.assertRaises(DocumentError) as caught:
            limited.register_upload(path)
        self.assertEqual(caught.exception.code, "SOURCE_TOO_LARGE")
        with patch("agentic_analytics.agent.tools.documents.fetch_public_url", return_value=(b"short", "text/plain", "https://example.org/source.txt")) as fetch:
            limited.inspect_source(url="https://example.org/source.txt")
        self.assertEqual(fetch.call_args.kwargs["max_bytes"], 12)

    def test_bad_header_selection_returns_source_owned_retry_with_complete_dates(self):
        source = self.upload("recover-header.csv", "metric,prefix,current,prior\n,3,0 September 2026 30 No,vember 2025\ncash,,100,90\n")
        arguments = {"source_id": source["source_id"], "table_id": "table_001", "selected_rows": [2],
                     "selected_columns": ["metric", "current", "prior"], "unpivot": {"columns": ["current", "prior"],
                         "period_column": "as_of", "value_column": "amount", "header_row": 1, "period_format": "english_dmy"}}
        result = self.docs.extra_tools()["prepare_source_table"]["handler"](arguments)
        self.assertEqual(result["code"], "INVALID_SOURCE_DATE_FORMAT")
        recovery = result["recovery"]
        self.assertEqual(recovery["header_rows"][0]["cells"][1]["text"], "3")
        update = recovery["suggested_unpivot_update"]
        self.assertEqual(update["period_sources"]["current"]["columns"], ["prefix", "current", "prior"])
        self.assertEqual(update["period_sources"]["prior"]["date_index"], 1)
        self.assertIn('"date_index": 1', result["message"])
        retry = copy.deepcopy(arguments)
        retry["unpivot"].update(update)
        prepared = self.docs.extra_tools()["prepare_source_table"]["handler"](retry)
        self.assertEqual(prepared["status"], "ok")
        self.assertEqual([(row["as_of"], row["amount"]) for row in prepared["preview"]], [("2026-09-30", "100"), ("2025-11-30", "90")])

    def test_header_recovery_exposes_ambiguous_day_prefix_without_assigning_a_date(self):
        source = self.upload("ambiguous-day.csv", "metric,prefix,current,prior\n,2,8 February 2026 30 No,vember 2025\ncash,,100,90\n")
        result = self.docs.extra_tools()["prepare_source_table"]["handler"]({"source_id": source["source_id"], "table_id": "table_001", "selected_rows": [2],
            "selected_columns": ["metric", "current", "prior"], "unpivot": {"columns": ["current", "prior"], "period_column": "as_of", "value_column": "amount", "header_row": 1, "period_format": "english_dmy"}})
        self.assertTrue(result["recovery"]["ambiguous_date_interpretations"])
        self.assertIsNone(result["recovery"]["suggested_unpivot_update"])

    def test_header_diagnostics_do_not_assign_dates_to_nonoverlapping_value_columns(self):
        source = self.upload("ambiguous-header.csv", "date_left,date_right,current,prior\n31 March 2026,31 December 2025,,\n,,100,90\n")
        arguments = {"source_id": source["source_id"], "table_id": "table_001", "selected_rows": [2],
                     "selected_columns": ["current", "prior"], "unpivot": {"columns": ["current", "prior"],
                         "period_column": "as_of", "value_column": "amount", "header_row": 1, "period_format": "english_dmy"}}
        result = self.docs.extra_tools()["prepare_source_table"]["handler"](arguments)
        self.assertEqual(result["code"], "INVALID_UNPIVOT")
        self.assertTrue(result["recovery"]["date_candidates"])
        self.assertIsNone(result["recovery"]["suggested_unpivot_update"])

    def test_top_level_numeric_parsing_aliases_preserve_values_proof_and_recovery(self):
        import pandas as pd
        source = self.upload("numeric-alias.csv", 'month,value (million TL)\n2026-01,"(1,234)"\n2026-02,"2,345"\n')
        arguments = {"source_id": source["source_id"], "table_id": "table_001", "contract": self.contract,
                     "expected_version": 0, "column_mapping": {"month": "month", "value_million_TL": "value"},
                     "unit_evidence": {"value": "million TL"}, "number_format": "decimal_dot_grouped", "negative_format": "accounting_parentheses"}
        result = self.docs.extra_tools()["publish_selected_table"]["handler"](arguments)
        self.assertEqual(result["status"], "ok")
        values = pd.read_parquet(self.store.overlay_path(result["dataset_id"]))["value"].tolist()
        self.assertEqual(values, [-1234, 2345])
        self.assertEqual(result["provenance"]["numeric_parsing"]["declared_at_top_level"], {"number_format": "decimal_dot_grouped", "negative_format": "accounting_parentheses"})
        recovered = self.docs.recover_publication(arguments, {})
        self.assertEqual(recovered["dataset_id"], result["dataset_id"])
        conflicting = copy.deepcopy(arguments)
        conflicting["contract"]["number_format"] = "decimal_comma"
        conflict = self.docs.extra_tools()["publish_selected_table"]["handler"](conflicting)
        self.assertEqual(conflict["code"], "CONFLICTING_NUMERIC_FORMAT")
        self.assertEqual(self.store.workspace("workspace_docs")["version"], 1)

    def test_unprepared_source_origins_follow_typed_date_integer_and_label_sort(self):
        import pandas as pd
        source = self.upload("typed-sort.csv", "day,entity (count),label,amount (TRY)\n2026-02-01,2,Z,202\n2026-01-01,10,A,110\n2026-01-01,2,Z,102\n2026-01-01,2,A,101\n")
        contract = {"name": "typed_sort", "frequency": "event", "date_column": "day", "key": ["day", "entity_id", "label"], "grain": ["day", "entity_id", "label"], "columns": {
            "day": {"dtype": "date", "kind": "dimension", "unit": "calendar", "nullable": False},
            "entity_id": {"dtype": "integer", "kind": "dimension", "unit": "count", "nullable": False},
            "label": {"dtype": "string", "kind": "dimension", "unit": "label", "nullable": False},
            "amount": {"dtype": "integer", "kind": "stock", "unit": "TRY", "nullable": False}}}
        result = self.docs.publish_selected_table(source["source_id"], "table_001", contract, 0,
            column_mapping={"day": "day", "entity_count": "entity_id", "label": "label", "amount_TRY": "amount"},
            unit_evidence={"entity_id": "count", "amount": "TRY"})
        frame = pd.read_parquet(self.store.overlay_path(result["dataset_id"]))
        self.assertEqual(frame["amount"].tolist(), [101, 102, 110, 202])
        proof = result["provenance"]
        self.assertEqual(proof["row_order"]["stored_row_to_source_csv_row"], [4, 3, 2, 1])
        raw = self.docs.review_candidate(source["source_id"], "table_001")
        for record, origin in zip(frame.to_dict("records"), proof["cell_origins"]):
            self.assertEqual(set(origin), set(contract["columns"]))
            for column, cell in origin.items():
                value = raw["rows"][cell["candidate_row"] - 1][raw["columns"].index(cell["candidate_column"])]
                self.assertEqual(int(value) if column in {"entity_id", "amount"} else value, record[column])


if __name__ == "__main__":
    unittest.main()
