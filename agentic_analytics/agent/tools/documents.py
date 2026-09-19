"""Bounded source inspection and explicit publication of selected document tables.

Uploads enter through trusted application code. Model tools accept source IDs or
public URLs, never local paths. Extracted text is untrusted source data.
"""

from __future__ import annotations

import csv
from concurrent.futures import ThreadPoolExecutor, TimeoutError as FutureTimeout
import calendar
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import hashlib
from html.parser import HTMLParser
import http.client
import io
import ipaddress
import json
import mimetypes
import os
from pathlib import Path
import re
import socket
import ssl
import tempfile
import threading
import time
import unicodedata
from urllib import error, parse, request
import xml.etree.ElementTree as ET
import zipfile

from agentic_analytics.lakehouse.store import StoreError
from agentic_analytics.lakehouse.units import KINDS, normalize_column, unit_quote_matches, unit_header_conflict
from agentic_analytics.lakehouse.financial_semantics import cumulative_evidence
from agentic_analytics.agent.tools.document_tables import HTMLTableExtractor


_DNS_WORKERS = ThreadPoolExecutor(max_workers=4, thread_name_prefix="document-dns")
_DNS_SLOTS = threading.BoundedSemaphore(4)


# Curated provider hints stay separate from generic search. They improve query
# formulation without allowing a non-official result to enter an official run.
OFFICIAL_SOURCE_REGISTRY = {
    "kkb.com.tr": {
        "institution": "Kredi Kayıt Bürosu",
        "search_variants": ("{query}",),
    },
    "tcmb.gov.tr": {
        "institution": "TCMB",
        "search_variants": ("{query}", "{query} bülten", "{query} raporu", "{query} gelişmeleri", "{query} yayın"),
    },
    "bddk.org.tr": {
        "institution": "BDDK",
        "search_variants": ("{query}", "{query} bülten", "{query} raporu", "{query} gelişmeleri", "{query} duyuru"),
    },
    "tuik.gov.tr": {
        "institution": "TÜİK",
        "search_variants": ("{query}", "{query} bülten", "{query} raporu", "{query} gelişmeleri", "{query} haber bülteni"),
    },
    "borsaistanbul.com": {
        "institution": "Borsa İstanbul",
        "search_variants": ("{query}", "{query} endeks", "{query} endeks verisi", "{query} metodoloji"),
    },
    "kap.org.tr": {
        "institution": "Kamuyu Aydınlatma Platformu",
        "search_variants": ("{query}", "{query} finansal rapor", "{query} finansal tablo", "{query} duyuru"),
    },
}


class DocumentError(ValueError):
    def __init__(self, message, code="DOCUMENT_ERROR", artifact_ref=None, recovery=None):
        super().__init__(message)
        self.code = code
        self.artifact_ref = artifact_ref
        self.recovery = recovery


def _canonical(value):
    return json.dumps(value, sort_keys=True, ensure_ascii=False, allow_nan=False, separators=(",", ":")).encode()


def _search_text(value):
    return "".join(character for character in unicodedata.normalize("NFKD", str(value).casefold().replace("ı", "i"))
                   if not unicodedata.combining(character))


def _safe_truncate(text, max_len):
    if len(text) <= max_len:
        return text
    truncated = text[:max_len]
    cut = max(truncated.rfind("."), truncated.rfind("\n"))
    if cut <= max_len // 2:
        cut = truncated.rfind(" ")
    if cut <= 0:
        return truncated
    return truncated[:cut + 1]


def _unit_caption(text):
    """Retain explicit common-unit declarations, not other metric headers."""
    lines = []
    for line in text.splitlines():
        folded = _search_text(line).strip()
        declaration = any(term in folded for term in ("amounts are", "amounts expressed", "all amounts", "all figures", "figures in", "amounts in", "tutarlar", "aksi belirtilmedikce", "birim:", "unit:"))
        standalone = bool(re.fullmatch(r"[()\s]*(?:(?:tl|try|usd|eur|gbp)\s+(?:million|millions|thousand|thousands|billion|mn|bn|000)|(?:million|millions|thousand|thousands|billion|milyon|milyar|bin)\s+(?:tl|try|usd|eur|gbp))[()\s]*", folded))
        standalone = standalone or bool(re.fullmatch(r"[()\s]*(?:thousands?|millions?|billions?)(?: of)?\s+(?:turkish lira(?:\s*\(tl\))?|us dollars?(?:\s*\(usd\))?|euros?(?:\s*\(eur\))?)[()\s]*", folded))
        standalone = standalone or bool(re.fullmatch(r"[()\s]*(?:bin|milyon|milyar)\s+(?:turk lirasi(?:\s*\((?:tl|try)\))?|abd dolari(?:\s*\(usd\))?|avro(?:\s*\(eur\))?)[()\s]*", folded))
        if declaration or standalone:
            lines.append(line.strip())
    return "\n".join(lines)[:4000]


def _write_json(path, value):
    fd, name = tempfile.mkstemp(dir=path.parent, suffix=".tmp")
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(_canonical(value))
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(name, path)
    finally:
        Path(name).unlink(missing_ok=True)


def _public_destination(url, *, timeout=5):
    if not isinstance(url, str) or len(url) > 4096 or any(ord(char) < 32 for char in url):
        raise DocumentError("Invalid public URL.", "UNSAFE_URL")
    parsed = parse.urlsplit(url)
    if parsed.scheme not in {"https", "http"} or not parsed.hostname or parsed.username or parsed.password:
        raise DocumentError("Only public HTTP(S) URLs without credentials are allowed.", "UNSAFE_URL")
    try:
        port = parsed.port or (443 if parsed.scheme == "https" else 80)
        if port not in {80, 443}:
            raise DocumentError("Only ports 80 and 443 are allowed.", "UNSAFE_URL")
        if not _DNS_SLOTS.acquire(blocking=False):
            raise DocumentError("Source resolver is busy; retry within the request budget.", "FETCH_BUSY")
        try:
            pending = _DNS_WORKERS.submit(socket.getaddrinfo, parsed.hostname, port, type=socket.SOCK_STREAM)
        except Exception:
            _DNS_SLOTS.release()
            raise
        pending.add_done_callback(lambda future: _DNS_SLOTS.release())
        try:
            addresses = pending.result(timeout=timeout)
        except FutureTimeout as exc:
            raise DocumentError("Source DNS resolution exceeded its time budget.", "FETCH_TIMEOUT") from exc
    except DocumentError:
        raise
    except (ValueError, OSError) as exc:
        raise DocumentError("URL host could not be safely resolved.", "UNSAFE_URL") from exc
    resolved = []
    for entry in addresses:
        address = ipaddress.ip_address(entry[4][0])
        if not address.is_global or isinstance(address, ipaddress.IPv6Address) and address.ipv4_mapped and not address.ipv4_mapped.is_global:
            raise DocumentError("Private, loopback, reserved and local destinations are blocked.", "UNSAFE_URL")
        resolved.append(str(address))
    if not resolved:
        raise DocumentError("URL host has no public address.", "UNSAFE_URL")
    return parsed, resolved[0], port


def _fetch_trusted_internal(url: str, *, max_bytes: int = 2 * 1024 ** 2, timeout: int = 15) -> bytes:
    parsed = parse.urlsplit(url)
    if parsed.scheme not in {"http", "https"} or parsed.username or parsed.password:
        raise DocumentError("Invalid trusted URL.", "UNSAFE_URL")
    req = request.Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AgenticMinds-SourceReader/1",
            "Accept": "application/json",
            "Accept-Encoding": "identity",
        },
    )
    opener = request.build_opener(request.ProxyHandler({}))
    try:
        with opener.open(req, timeout=timeout) as response:
            blocks, total = [], 0
            while True:
                chunk = response.read(min(65536, max_bytes - total + 1))
                if not chunk:
                    break
                total += len(chunk)
                if total > max_bytes:
                    raise DocumentError("Response exceeds limit.", "SOURCE_TOO_LARGE")
                blocks.append(chunk)
            return b"".join(blocks)
    except (OSError, error.URLError) as exc:
        raise DocumentError(f"Internal fetch failed: {exc}", "FETCH_FAILED") from exc


class _NoRedirect(request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        return None


def _ascii_safe_url(url):
    """Percent-encode a raw path/query so http.client can send the request line.

    Search providers and discovered links can carry unescaped non-ASCII
    characters (e.g. Turkish letters in a Wikipedia path). http.client encodes
    the request line as ASCII and raises UnicodeEncodeError on those bytes, so
    encode defensively rather than letting a legitimate URL fail to fetch.
    """
    parsed = parse.urlsplit(url)
    path = parse.quote(parsed.path, safe="/%")
    query = parse.quote(parsed.query, safe="=&%")
    return parse.urlunsplit((parsed.scheme, parsed.netloc, path, query, ""))


def fetch_public_url(url, *, max_bytes=16 * 1024**2, timeout=20, max_redirects=3):
    """Revalidate each hop and pin the validated IP for the actual connection."""
    deadline = time.monotonic() + timeout
    current = _ascii_safe_url(url)
    for hop in range(max_redirects + 1):
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            raise DocumentError("Download time budget exceeded.", "FETCH_TIMEOUT")
        parsed, address, port = _public_destination(current, timeout=min(remaining, 5))
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            raise DocumentError("Download time budget exceeded.", "FETCH_TIMEOUT")

        class PinnedHTTP(http.client.HTTPConnection):
            def connect(self):
                self.sock = socket.create_connection((address, port), timeout=min(remaining, 5))

        class PinnedHTTPS(http.client.HTTPSConnection):
            def connect(self):
                raw = socket.create_connection((address, port), timeout=min(remaining, 5))
                try:
                    self.sock = ssl.create_default_context().wrap_socket(raw, server_hostname=parsed.hostname)
                except Exception:
                    raw.close()
                    raise

        class HTTPHandler(request.HTTPHandler):
            def http_open(self, req):
                return self.do_open(PinnedHTTP, req)

        class HTTPSHandler(request.HTTPSHandler):
            def https_open(self, req):
                return self.do_open(PinnedHTTPS, req)

        opener = request.build_opener(request.ProxyHandler({}), _NoRedirect(), HTTPHandler(), HTTPSHandler())
        req = request.Request(current, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AgenticMinds-SourceReader/1", "Accept-Encoding": "identity"})
        try:
            response = opener.open(req, timeout=min(remaining, 5))
        except error.HTTPError as exc:
            try:
                if exc.code in {301, 302, 303, 307, 308} and exc.headers.get("Location") and hop < max_redirects:
                    current = _ascii_safe_url(parse.urljoin(current, exc.headers["Location"]))
                    continue
                raise DocumentError(f"Source returned HTTP {exc.code}.", "FETCH_FAILED") from exc
            finally:
                exc.close()
        except (OSError, error.URLError) as exc:
            raise DocumentError("Source download failed or timed out.", "FETCH_FAILED") from exc
        with response:
            declared = response.headers.get("Content-Length")
            if declared and (not declared.isdecimal() or int(declared) > max_bytes):
                raise DocumentError("Source exceeds the download size limit.", "SOURCE_TOO_LARGE")
            from agentic_analytics.agent.tools.search_backend import read_bounded_response
            try:
                data = read_bounded_response(response, max_bytes=max_bytes, deadline=deadline)
            except TimeoutError as exc:
                raise DocumentError("Download time budget exceeded.", "FETCH_TIMEOUT") from exc
            except ValueError as exc:
                raise DocumentError("Source exceeds the download size limit.", "SOURCE_TOO_LARGE") from exc
            return data, response.headers.get_content_type(), current
    raise DocumentError("Too many redirects.", "FETCH_FAILED")


class _HTMLArticle(HTMLParser):
    """Extract bounded article metadata without executing page scripts."""
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.title, self.description, self.canonical = "", "", ""
        self.meta, self.json_ld = {}, []
        self._title_text, self._script_text, self._script_type = [], [], None

    def handle_starttag(self, tag, attrs):
        values = dict(attrs)
        if tag == "title":
            self._title_text = []
        elif tag == "meta":
            key = values.get("property") or values.get("name") or values.get("itemprop")
            content = values.get("content")
            if key and content:
                self.meta[key.casefold()] = " ".join(content.split())[:2000]
        elif tag == "link" and values.get("rel", "").casefold() == "canonical":
            self.canonical = values.get("href", "")[:4096]
        elif tag == "script" and values.get("type", "").casefold() == "application/ld+json":
            self._script_text, self._script_type = [], "application/ld+json"

    def handle_data(self, data):
        if self._script_type:
            self._script_text.append(data)
        elif self._title_text is not None:
            self._title_text.append(data)

    def handle_endtag(self, tag):
        if tag == "title" and self._title_text is not None:
            self.title = " ".join("".join(self._title_text).split())[:500]
            self._title_text = None
        elif tag == "script" and self._script_type:
            try:
                value = json.loads("".join(self._script_text))
                self.json_ld.extend(value if isinstance(value, list) else [value])
            except (TypeError, ValueError, json.JSONDecodeError):
                pass
            self._script_text, self._script_type = [], None


class _HTMLReadable(HTMLParser):
    """Collect visible content and downloadable links without page execution."""
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.hidden, self.focus = [], 0
        self.text, self.focus_text, self.links = [], [], []
        self.link_positions = {}
        self.anchor = None
        self.svg_depth = 0

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag in {"script", "style", "noscript", "nav", "footer", "defs", "symbol"}:
            self.hidden.append(tag)
        if tag == "svg":
            self.svg_depth += 1
        if tag in {"main", "article"}:
            self.focus += 1
        if tag in {"p", "div", "h1", "h2", "h3", "li", "tr", "br"}:
            self.text.append("\n")
            self.focus_text.append("\n")
        if tag == "a":
            self.anchor = {"href": attrs.get("href", ""), "text": [],
                           "in_main_content": bool(self.focus) and not self.hidden,
                           "in_navigation": bool(self.hidden)}
        if tag == "img" and attrs.get("alt", "").strip():
            # Shareholder and partner lists often consist of linked logos.
            # Their alternative text is source-authored content, not OCR.
            self.handle_data(" " + attrs["alt"].strip() + " ")
        if self.svg_depth and tag in {"svg", "g", "circle", "path", "rect", "use"} and not self.hidden:
            # Preserve source-authored names in charts whose logos are paths.
            # Symbol IDs alone are implementation details, not readable names.
            label = " ".join(attrs.get("aria-label", "").split())[:300]
            if label:
                self.handle_data("\n" + label + "\n")
            elif attrs.get("role") == "progressbar" and attrs.get("data-slice-name", "").strip():
                label = " ".join(attrs["data-slice-name"].split())[:300]
                self.handle_data("\nGrafik kategorisi (kaynak etiketi): " + label + "\n")
            # aria-valuenow may be a rounded visual size. Do not promote it
            # to an exact amount, percentage, or shareholder interest.

    def handle_data(self, value):
        if self.anchor is not None:
            self.anchor["text"].append(value)
        if self.hidden:
            return
        self.text.append(value)
        if self.focus:
            self.focus_text.append(value)

    def handle_endtag(self, tag):
        if tag == "svg":
            self.svg_depth = max(0, self.svg_depth - 1)
        if self.hidden and tag == self.hidden[-1]:
            self.hidden.pop()
        if tag in {"main", "article"}:
            self.focus = max(0, self.focus - 1)
        if tag == "a" and self.anchor is not None:
            link = {"href": self.anchor["href"], "title": " ".join("".join(self.anchor["text"]).split())[:300],
                    "in_main_content": self.anchor["in_main_content"], "in_navigation": self.anchor["in_navigation"]}
            key = (link["href"], link["title"])
            position = self.link_positions.get(key)
            if position is not None:
                old = self.links[position]
                old["in_main_content"] |= link["in_main_content"]
                old["in_navigation"] &= link["in_navigation"]
            elif len(self.links) < 2000:
                self.link_positions[key] = len(self.links)
                self.links.append(link)
            self.anchor = None


def _article_metadata(data, mime_type, final_url):
    if mime_type != "text/html":
        return {}
    parser = _HTMLArticle()
    parser.feed(data.decode("utf-8-sig", errors="replace"))
    readable = _HTMLReadable()
    readable.feed(data.decode("utf-8-sig", errors="replace"))
    candidates = []
    for item in parser.json_ld:
        if not isinstance(item, dict):
            continue
        types = item.get("@type", [])
        types = types if isinstance(types, list) else [types]
        if any(value in {"Article", "NewsArticle", "Report"} for value in types):
            candidates.append(item)
    item = candidates[0] if candidates else {}
    image = item.get("image")
    return {
        "title": str(item.get("headline") or parser.meta.get("og:title") or parser.title)[:500],
        "description": str(item.get("description") or parser.meta.get("description") or parser.meta.get("og:description", ""))[:2000],
        "date_published": item.get("datePublished") or parser.meta.get("article:published_time"),
        "date_modified": item.get("dateModified") or parser.meta.get("article:modified_time"),
        "canonical_url": parse.urljoin(final_url, parser.canonical) if parser.canonical else final_url,
        "author": item.get("author", {}).get("name") if isinstance(item.get("author"), dict) else item.get("author"),
        "article_body": str(item.get("articleBody", ""))[:20000],
        "readable_text": "\n".join(line.strip() for line in "".join(readable.focus_text if "".join(readable.focus_text).strip() else readable.text).splitlines() if line.strip())[:40000],
        "document_links": [{"url": parse.urljoin(final_url, link["href"]), "title": link["title"]}
                           for link in readable.links
                           if re.search(r"\.(?:pdf|xlsx?|csv)(?:$|[?#])", link["href"], re.I)][:50],
        "source_links": [{"url": parse.urljoin(final_url, link["href"]), "title": link["title"],
                          "in_main_content": link["in_main_content"], "in_navigation": link["in_navigation"]}
                         for link in sorted(readable.links, key=lambda link: (not link["in_main_content"], link["in_navigation"]))
                         if link["href"] and not link["href"].startswith(("#", "javascript:", "mailto:", "tel:"))][:500],
        "link_count": len(readable.links),
        "image": image if isinstance(image, str) else None,
    }


def _column_name(value, index, used):
    ascii_value = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode()
    base = re.sub(r"[^A-Za-z0-9_]+", "_", ascii_value.strip()).strip("_")[:100] or f"column_{index + 1}"
    if base[0].isdigit():
        base = "c_" + base
    name, suffix = base, 2
    while name in used:
        name, suffix = f"{base}_{suffix}", suffix + 1
    used.add(name)
    return name


class DocumentTools:
    def __init__(self, store, workspace_id, *, upload_root=None, ocr_callback=None, searxng_url=None,
                 trusted_internal_urls=None,
                 max_source_bytes=32 * 1024**2, max_rows=5000, max_columns=64, max_pages=30):
        self.store, self.workspace_id = store, workspace_id
        self.store.workspace(workspace_id)
        self.root = store.root / "document_sources" / workspace_id
        self.root.mkdir(parents=True, exist_ok=True)
        self.upload_root = Path(upload_root or store.root / "uploads").resolve()
        self.upload_root.mkdir(parents=True, exist_ok=True)
        self.ocr_callback = ocr_callback
        self.max_source_bytes, self.max_rows = max_source_bytes, max_rows
        self.max_columns, self.max_pages = max_columns, max_pages

        # SearXNG ve güvenilir iç ağ yapılandırması
        if searxng_url is None:
            searxng_url = os.environ.get("SEARXNG_URL")
        self.searxng_url = searxng_url

        if trusted_internal_urls is None:
            raw_trusted = os.environ.get("TRUSTED_INTERNAL_URLS", "")
            trusted_internal_urls = [u.strip() for u in raw_trusted.split(",") if u.strip()]

        self._trusted_prefixes = tuple(
            u.rstrip("/") for u in (trusted_internal_urls or []) if isinstance(u, str) and u.strip()
        )
        if self.searxng_url:
            s_pref = self.searxng_url.rstrip("/")
            if s_pref not in self._trusted_prefixes:
                self._trusted_prefixes = self._trusted_prefixes + (s_pref,)

    def _is_trusted_url(self, url: str) -> bool:
        """Verilen URL güvenilir iç ağ prefixlerinden biriyle başlıyorsa True döner."""
        return any(url.startswith(p) for p in self._trusted_prefixes)

    def _directory(self, source_id):
        if not isinstance(source_id, str) or not re.fullmatch(r"source_[a-f0-9]{64}", source_id):
            raise DocumentError("Invalid source identifier.")
        directory = self.root / source_id
        if directory.is_symlink() or not directory.resolve().is_relative_to(self.root.resolve()):
            raise DocumentError("Source path is invalid.")
        return directory

    def source(self, source_id):
        directory = self._directory(source_id)
        manifest = json.loads((directory / "manifest.json").read_text())
        if hashlib.sha256((directory / "raw.bin").read_bytes()).hexdigest() != manifest["raw_sha256"]:
            raise DocumentError("Raw source hash mismatch.", "SOURCE_HASH_MISMATCH")
        return manifest

    def list_sources(self):
        return [self.source(path.name) for path in sorted(self.root.glob("source_*")) if path.is_dir()]

    def raw_source_bytes(self, source_id):
        """Trusted UI raw-download helper, with no caller-supplied filesystem path."""
        manifest = self.source(source_id)
        data = (self._directory(source_id) / "raw.bin").read_bytes()
        if hashlib.sha256(data).hexdigest() != manifest["raw_sha256"]:
            raise DocumentError("Raw source hash mismatch.", "SOURCE_HASH_MISMATCH")
        return data

    def review_candidate(self, source_id, table_id):
        """Full bounded candidate for trusted UI review, never a model tool."""
        manifest = self.source(source_id)
        self.inspect_source(source_id=source_id)
        inspection = json.loads((self._directory(source_id) / "inspection.json").read_text())
        table = next((item for item in inspection["tables"] if item["table_id"] == table_id), None)
        if table is None:
            raise DocumentError("Unknown candidate table.")
        return {"status": "ok", "source_id": source_id, "raw_sha256": manifest["raw_sha256"],
                **self._reviewed_table(source_id, table), "preview_truncated": False}

    def _register(self, data, filename, mime_type, source_url=None):
        if not data or len(data) > self.max_source_bytes:
            raise DocumentError("Source is empty or exceeds its byte limit.", "SOURCE_TOO_LARGE")
        filename = Path(filename).name[:200]
        metadata = {"filename": filename, "mime_type": mime_type, "source_url": source_url,
                    "raw_sha256": hashlib.sha256(data).hexdigest(), "size_bytes": len(data),
                    "workspace_id": self.workspace_id}
        source_id = "source_" + hashlib.sha256(_canonical(metadata)).hexdigest()
        directory = self._directory(source_id)
        directory.mkdir(exist_ok=True)
        manifest = {**metadata, "source_id": source_id, "artifact_ref": source_id,
                    "source_namespace": "external:" + (parse.urlsplit(source_url).hostname if source_url else "upload") + ":" + source_id,
                    "status": "registered", "content_is_untrusted_data": True}
        if not (directory / "manifest.json").exists():
            fd, name = tempfile.mkstemp(dir=directory)
            try:
                with os.fdopen(fd, "wb") as handle:
                    handle.write(data)
                os.replace(name, directory / "raw.bin")
            finally:
                Path(name).unlink(missing_ok=True)
            _write_json(directory / "manifest.json", manifest)
        return self.source(source_id)

    def register_upload(self, path, filename=None):
        """Trusted server entry point; deliberately absent from model tool schemas."""
        source = Path(path).resolve(strict=True)
        if not source.is_relative_to(self.upload_root) or not source.is_file():
            raise DocumentError("Upload is outside the server upload directory.", "UNSAFE_UPLOAD")
        if source.stat().st_size > self.max_source_bytes:
            raise DocumentError("Upload exceeds its byte limit.", "SOURCE_TOO_LARGE")
        with source.open("rb") as handle:
            data = handle.read(self.max_source_bytes + 1)
        name = filename or source.name
        return self._register(data, name, mimetypes.guess_type(name)[0] or "application/octet-stream")

    def _table(self, rows, *, page=None, sheet=None, units=None, origin="parsed", header_rows=1):
        if not rows or len(rows) < 2:
            return None
        header_evidence = None
        if header_rows > 1:
            if header_rows >= len(rows) or any(len(row) != len(rows[0]) for row in rows[:header_rows]):
                return None
            header_evidence = rows[:header_rows]
            flattened = []
            for index in range(len(rows[0])):
                values = list(dict.fromkeys(str(row[index]).strip() for row in header_evidence if row[index] is not None and str(row[index]).strip()))
                flattened.append(" / ".join(values))
            rows = [flattened, *rows[header_rows:]]
        if len(rows) - 1 > self.max_rows or not 1 <= len(rows[0]) <= self.max_columns:
            raise DocumentError("Candidate table exceeds row or column limits.", "TABLE_TOO_LARGE")
        width, used = len(rows[0]), set()
        if any(len(row) != width for row in rows):
            raise DocumentError("Candidate contains inconsistent row widths.", "AMBIGUOUS_TABLE")
        original = [str(value or "").strip() for value in rows[0]]
        columns = [_column_name(value, index, used) for index, value in enumerate(original)]
        values = [[None if cell is None or str(cell).strip() == "" else str(cell).strip() for cell in row] for row in rows[1:]]
        if any(len(cell) > 20000 for row in values for cell in row if cell is not None):
            raise DocumentError("A candidate cell exceeds its text limit.", "TABLE_TOO_LARGE")
        collapsed_numeric = any(cell is not None and "\n" in cell and all(
            re.fullmatch(r"[\d\s,.%()+\-]+", line) for line in cell.splitlines() if line.strip()) for row in values for cell in row)
        axis_caption = original[0] if len(original) > 1 and all(re.fullmatch(r"[\dQqHh\s()/.,-]+", header) for header in original[1:] if header) and all(original[1:]) else ""
        return {"columns": columns, "original_columns": dict(zip(columns, original)), "rows": values,
                "row_count": len(values), "page": page, "sheet": sheet, "units": units or {}, "origin": origin,
                "source_header_quotes": {column: (str(header_evidence[-1][index] or "") if header_evidence else original[index])
                                         for index, column in enumerate(columns)},
                "header_rows": header_evidence,
                "unit_contexts": {column: [original[index]] for index, column in enumerate(columns)}, "unit_caption": axis_caption,
                "layout_review_required": collapsed_numeric,
                "quality_notes": (["Multiple numeric lines share a single extracted cell; review or choose another extraction."] if collapsed_numeric else [])
                                 + (["Some column headers are blank; establish their meaning from source context before publication."] if not all(original) else [])}

    def _extract_image(self, data, mime_type, *, page=1):
        from PIL import Image
        if len(data) > 8 * 1024**2:
            raise DocumentError("Image exceeds the 8 MiB extraction limit.", "SOURCE_TOO_LARGE")
        with Image.open(io.BytesIO(data)) as image:
            if image.width * image.height > 20_000_000:
                raise DocumentError("Image exceeds its pixel limit.", "SOURCE_TOO_LARGE")
            image.verify()
        if self.ocr_callback is None:
            raise DocumentError("Image extraction is not configured.", "OCR_UNAVAILABLE")
        extracted = self.ocr_callback(data, mime_type)
        machine_record = {"image_sha256": hashlib.sha256(data).hexdigest(), "page": page,
                          "raw_machine_output": extracted, "verified": False}
        extraction_bytes = _canonical(machine_record)
        extraction_ref = "extraction_" + hashlib.sha256(extraction_bytes).hexdigest()
        extraction_root = self.root / "extractions"
        extraction_root.mkdir(exist_ok=True)
        if len(extraction_bytes) <= 1024**2:
            _write_json(extraction_root / (extraction_ref + ".json"), {"extraction_id": extraction_ref, **machine_record})
        else:
            raise DocumentError("Machine extraction exceeds its evidence size limit.", "OCR_INVALID_OUTPUT")
        if not isinstance(extracted, dict) or not isinstance(extracted.get("text", ""), str):
            raise DocumentError("Image extractor returned an invalid envelope.", "OCR_INVALID_OUTPUT", extraction_ref)
        if extracted.get("finish_reason") == "length" or extracted.get("status", "ok") != "ok":
            raise DocumentError("Image extraction was incomplete or failed. Raw machine output was retained for review.", "OCR_INVALID_OUTPUT", extraction_ref)
        if not isinstance(extracted.get("tables", []), list):
            raise DocumentError("Image extractor tables must be a list.", "OCR_INVALID_OUTPUT", extraction_ref)
        tables = []
        for table in extracted.get("tables", [])[:10]:
            if not isinstance(table, dict):
                raise DocumentError("OCR table must be an object; raw output was retained.", "OCR_INVALID_OUTPUT", extraction_ref)
            rows, columns = table.get("rows", []), table.get("columns", [])
            if not isinstance(rows, list) or not isinstance(columns, list):
                raise DocumentError("OCR rows and columns must be lists; raw output was retained.", "OCR_INVALID_OUTPUT", extraction_ref)
            raw_rows, raw_columns = json.loads(_canonical(rows)), list(columns)
            if rows and isinstance(rows[0], dict):
                if any(not isinstance(row, dict) or set(row) != set(columns) for row in rows):
                    raise DocumentError("OCR row keys disagree with the extracted headers. Raw output was retained.", "OCR_INVALID_OUTPUT", extraction_ref)
                rows = [[row.get(column) for column in columns] for row in rows]
            if (not isinstance(columns, list) or not 1 <= len(columns) <= self.max_columns or not isinstance(rows, list)
                    or len(rows) > self.max_rows or any(not isinstance(row, list) for row in rows)):
                raise DocumentError("OCR table shape exceeds its bounds; raw output was retained.", "OCR_INVALID_OUTPUT", extraction_ref)
            hypothesis = None
            if rows and all(len(row) == len(columns) + 1 for row in rows) and len(columns) < self.max_columns:
                identifier = re.compile(r"(?:[IVXLCDM]+[.)]?|\d+(?:\.\d+)*[.)]?)?")
                if all(identifier.fullmatch(str(row[0] or "").strip()) and isinstance(row[1], str)
                       and bool(re.search(r"[^\W\d_]", row[1])) for row in rows):
                    columns = ["source_item_code", *columns]
                    hypothesis = {"operation": "add_missing_item_identifier_header", "original_headers": raw_columns,
                                  "added_header": "source_item_code", "evidence": "Every row has exactly one extra leading item identifier followed by a text label.",
                                  "numeric_cells_changed": False, "requires_independent_review": True}
            try:
                candidate = self._table([columns, *rows], page=page, units=table.get("units"), origin="ocr")
            except DocumentError as exc:
                if exc.code != "AMBIGUOUS_TABLE":
                    raise DocumentError(str(exc), exc.code, extraction_ref) from exc
                used = set()
                names = [_column_name(value, index, used) for index, value in enumerate(columns)]
                # Preserve all cells, including excess cells. There is no aligned
                # preview or publishable grid until independent review supplies it.
                candidate = {"columns": names, "original_columns": dict(zip(names, map(str, columns))),
                             "source_header_quotes": dict(zip(names, map(str, columns))), "rows": rows, "row_count": len(rows),
                             "page": page, "sheet": None, "units": table.get("units", {}), "origin": "ocr",
                             "layout_review_required": True, "extraction_status": "unaligned_rows_require_review"}
            if candidate:
                candidate.update(extraction_artifact_ref=extraction_ref, raw_machine_rows=raw_rows, raw_machine_headers=raw_columns,
                                 context_text=extracted.get("text", ""), unit_caption=_unit_caption(extracted.get("text", "")))
                if hypothesis:
                    candidate["header_hypothesis"] = hypothesis
                tables.append(candidate)
        return extracted.get("text", ""), tables

    def _parse(self, manifest, data, page_numbers=None, table_strategy="lines"):
        suffix, mime = Path(manifest["filename"]).suffix.lower(), manifest["mime_type"]
        tables, pages, text, warnings = [], [], "", []
        total_pages, document_metadata = None, {}
        if (page_numbers is not None or table_strategy != "lines") and suffix != ".pdf" and mime != "application/pdf":
            raise DocumentError("Page selection is available only for PDFs.", "INVALID_PAGE_SELECTION")
        if suffix == ".csv" or mime == "text/csv":
            text = data.decode("utf-8-sig")
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(io.StringIO(text), dialect=dialect, strict=True)
            rows = []
            for row in reader:
                rows.append(row)
                if len(rows) > self.max_rows + 1:
                    raise DocumentError("CSV exceeds its row limit.", "TABLE_TOO_LARGE")
            tables.append(self._table(rows))
        elif suffix == ".xlsx" or mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist()) > 10000 or sum(item.file_size for item in archive.infolist()) > 64 * 1024**2:
                    raise DocumentError("Expanded workbook exceeds its size limit.", "SOURCE_TOO_LARGE")
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                if len(workbook.worksheets) > 10:
                    warnings.append({"code": "SHEET_LIMIT", "total_sheets": len(workbook.worksheets), "processed_sheets": 10})
                for sheet in workbook.worksheets[:10]:
                    if sheet.max_row > self.max_rows + 1 or sheet.max_column > self.max_columns:
                        raise DocumentError("Workbook sheet exceeds table limits.", "TABLE_TOO_LARGE")
                    rows, formulas, missing = [], [], []
                    cached_rows = cached[sheet.title].iter_rows()
                    for source_row, cached_row in zip(sheet.iter_rows(), cached_rows):
                        row = []
                        for cell, saved in zip(source_row, cached_row):
                            value = cell.value
                            if cell.data_type == "f":
                                value = saved.value if saved.data_type != "e" else None
                                formulas.append({"cell": cell.coordinate, "formula": str(cell.value),
                                                 "cached_value": str(value) if value is not None else None,
                                                 "verification": "saved_source_value_not_recalculated"})
                                if value is None:
                                    missing.append(cell.coordinate)
                            row.append(value)
                        rows.append(row)
                    table = self._table(rows, sheet=sheet.title)
                    if table and formulas:
                        table.update(formula_cells=formulas, missing_formula_cache=missing)
                        warnings.append({"code": "FORMULA_VALUES_REVIEW_REQUIRED" if missing else "CACHED_FORMULA_VALUES",
                                         "sheet": sheet.title, "formula_cells": len(formulas), "missing_cells": missing,
                                         "message": "Formula code was not executed. Cached values are the values saved in the source workbook; their freshness is not independently verified."})
                    tables.append(table)
                    text += "\n" + "\n".join(" | ".join(str(value or "") for value in row) for row in rows)
            finally:
                workbook.close()
                cached.close()
        elif suffix in {".html", ".htm"} or mime == "text/html":
            parser = HTMLTableExtractor(self.max_rows + 1, self.max_columns)
            decoded = data.decode("utf-8-sig", errors="replace")
            readable = _HTMLReadable()
            readable.feed(decoded)
            text = "".join(readable.text)
            try:
                parser.feed(decoded)
                warnings.extend(parser.warnings)
                for extracted in parser.tables[:10]:
                    try:
                        table = self._table(extracted["rows"], header_rows=extracted["header_rows"])
                        if table:
                            table.update(cell_spans=extracted["spans"], layout_review_required=table["layout_review_required"] or extracted["requires_review"])
                            table["unit_caption"] += "\n" + _unit_caption(text)
                        tables.append(table)
                    except DocumentError as exc:
                        warnings.append({"code": exc.code, "message": str(exc)})
            except (DocumentError, ValueError) as exc:
                warnings.append({"code": getattr(exc, "code", "AMBIGUOUS_TABLE"), "message": str(exc)})
        elif suffix == ".pdf" or mime == "application/pdf":
            import pdfplumber
            image_pages, deferred_pages = 0, []
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                document_metadata = {key.lower(): value[:500] for key, value in (pdf.metadata or {}).items()
                                     if key in {"Title", "Author", "Subject"} and isinstance(value, str)}
                total_pages = len(pdf.pages)
                selected = page_numbers or list(range(1, min(total_pages, self.max_pages) + 1))
                if any(number > total_pages for number in selected):
                    raise DocumentError("Selected PDF page does not exist.", "INVALID_PAGE_SELECTION")
                if len(selected) < total_pages:
                    warnings.append({"code": "PARTIAL_PDF_INSPECTION", "total_pages": total_pages,
                                     "processed_pages": selected,
                                     "message": "Only the listed pages were extracted. Select other 1-based page_numbers to continue; omitted pages are not evidence."})
                for number in selected:
                    page = pdf.pages[number - 1]
                    page_text = (page.extract_text() or "")[:30000]
                    page_method = "pdf_text"
                    if not page_text.strip() and self.ocr_callback is not None and image_pages < 3:
                        if (page.width * 110 / 72) * (page.height * 110 / 72) > 20_000_000:
                            raise DocumentError("Rendered PDF page exceeds its pixel limit.", "SOURCE_TOO_LARGE")
                        buffer = io.BytesIO()
                        page.to_image(resolution=110).original.convert("RGB").save(buffer, format="PNG")
                        page_text, extracted_tables = self._extract_image(buffer.getvalue(), "image/png", page=number)
                        tables.extend(extracted_tables)
                        image_pages += 1
                        page_method = "ocr"
                    elif not page_text.strip():
                        deferred_pages.append(number)
                    pages.append({"page": number, "text": page_text[:30000], "extraction_method": page_method})
                    text += "\n" + page_text
                    if page_method == "pdf_text":
                        extracted = page.extract_tables() if table_strategy == "lines" else page.extract_tables({
                            "vertical_strategy": "text", "horizontal_strategy": "text", "min_words_vertical": 3})
                        for rows in extracted[:10]:
                            try:
                                table = self._table(rows, page=number)
                                if table:
                                    table["context_text"] = page_text
                                    table["unit_caption"] += "\n" + _unit_caption(page_text)
                                    table["table_strategy"] = table_strategy
                                tables.append(table)
                            except DocumentError as exc:
                                warnings.append({"code": exc.code, "page": number, "message": str(exc)})
                    page.close()
            if image_pages:
                warnings.append({"code": "OCR_REVIEW_REQUIRED", "pages_processed": image_pages,
                                 "message": "Rendered PDF image cells require independent review before publication."})
            if deferred_pages:
                warnings.append({"code": "OCR_PAGE_LIMIT" if self.ocr_callback else "SCANNED_PDF_NEEDS_OCR",
                                 "pages": deferred_pages, "message": "Some image-only pages were not extracted; at most three OCR pages are processed per document."})
        elif suffix in {".png", ".jpg", ".jpeg"} or mime in {"image/png", "image/jpeg"}:
            text, tables = self._extract_image(data, "image/png" if data.startswith(b"\x89PNG") else "image/jpeg")
            pages = [{"page": 1, "text": text[:30000]}]
            warnings.append({"code": "OCR_REVIEW_REQUIRED", "message": "Image cells require independent numeric verification before publication."})
        elif suffix == ".txt" or mime == "text/plain":
            text = data.decode("utf-8-sig")
        else:
            raise DocumentError("Supported sources are CSV, XLSX, HTML, PDF, UTF-8 text, PNG and JPEG.", "UNSUPPORTED_FORMAT")
        tables = [table for table in tables if table is not None]
        if len(tables) > 30:
            warnings.append({"code": "TABLE_LIMIT", "tables_detected": len(tables), "tables_retained": 30,
                             "message": "Select fewer PDF pages to inspect the remaining tables."})
            tables = tables[:30]
        page_counts = {}
        for index, table in enumerate(tables, 1):
            if total_pages is not None:
                number = table["page"]
                page_counts[number] = page_counts.get(number, 0) + 1
                variant = "_text" if table_strategy == "text" else ""
                table["table_id"] = f"table_p{number:06d}{variant}_{page_counts[number]:03d}"
            else:
                table["table_id"] = f"table_{index:03d}"
        return {"text": text[:200000], "text_truncated": len(text) > 200000, "pages": pages,
                "tables": tables, "warnings": warnings, "total_pages": total_pages, "document_metadata": document_metadata,
                "processed_pages": [page["page"] for page in pages],
                "processed_extractions": [f"{page['page']}:{table_strategy}" for page in pages]}

    def inspect_source(self, source_id=None, url=None, page_numbers=None, table_strategy="lines", *, _deadline=None):
        if bool(source_id) == bool(url):
            raise DocumentError("Provide exactly one source_id or public URL.")
        if page_numbers is not None and (not isinstance(page_numbers, list) or not 1 <= len(page_numbers) <= self.max_pages
                or any(type(number) is not int or number < 1 for number in page_numbers) or len(set(page_numbers)) != len(page_numbers)):
            raise DocumentError(f"Select 1 to {self.max_pages} distinct 1-based PDF page numbers.", "INVALID_PAGE_SELECTION")
        if page_numbers:
            page_numbers = sorted(page_numbers)
        if table_strategy not in {"lines", "text"}:
            raise DocumentError("PDF table strategy must be lines or text.", "INVALID_TABLE_STRATEGY")
        article = {}
        if url:
            timeout = min(20, _deadline - time.monotonic()) if _deadline is not None else 20
            if timeout <= 0:
                raise DocumentError("Source research time budget exceeded.", "FETCH_TIMEOUT")
            data, mime, final_url = fetch_public_url(url, max_bytes=self.max_source_bytes, timeout=timeout)
            name = Path(parse.unquote(parse.urlsplit(final_url).path)).name or "source"
            manifest = self._register(data, name, mime, final_url)
            article = _article_metadata(data, mime, final_url)
            source_id = manifest["source_id"]
        manifest = self.source(source_id)
        if (page_numbers is not None or table_strategy != "lines") and Path(manifest["filename"]).suffix.lower() != ".pdf" and manifest["mime_type"] != "application/pdf":
            raise DocumentError("Page selection and table strategies are available only for PDFs.", "INVALID_PAGE_SELECTION")
        directory = self._directory(source_id)
        cache = directory / "inspection.json"
        if cache.exists():
            inspection = json.loads(cache.read_text())
            requested_pages = page_numbers or inspection.get("processed_pages", [])[:self.max_pages]
            done = set(inspection.get("processed_extractions", [f"{number}:lines" for number in inspection.get("processed_pages", [])]))
            if requested_pages and not all(f"{number}:{table_strategy}" in done for number in requested_pages):
                fresh = self._parse(manifest, (directory / "raw.bin").read_bytes(), requested_pages, table_strategy)
                # Existing candidate identities and human reviews never change
                # when another part of a source is inspected.
                by_id = {table["table_id"]: table for table in inspection["tables"]}
                for table in fresh["tables"]:
                    by_id.setdefault(table["table_id"], table)
                by_page = {page["page"]: page for page in inspection["pages"]}
                by_page.update({page["page"]: page for page in fresh["pages"]})
                if len(by_page) > 1000 or len(by_id) > 500:
                    raise DocumentError("Source inspection cache limit reached.", "SOURCE_TOO_LARGE")
                inspection = {**fresh, "tables": list(by_id.values()), "pages": [by_page[number] for number in sorted(by_page)],
                              "processed_pages": sorted(by_page), "processed_extractions": sorted(done.union(fresh["processed_extractions"]))}
                full_text = "\n".join(page["text"] for page in inspection["pages"])
                inspection.update(text=full_text[:200000], text_truncated=len(full_text) > 200000)
                _write_json(cache, inspection)
        else:
            inspection = self._parse(manifest, (directory / "raw.bin").read_bytes(), page_numbers, table_strategy)
            _write_json(cache, inspection)
        if not article and manifest["mime_type"] == "text/html":
            article = _article_metadata((directory / "raw.bin").read_bytes(), manifest["mime_type"], manifest.get("source_url") or "")
        inspection["tables"] = [self._reviewed_table(source_id, table) for table in inspection["tables"]]
        selected_pages = page_numbers or inspection.get("processed_pages", [])[:self.max_pages]
        selected_tables = [table for table in inspection["tables"] if (not selected_pages or table.get("page") in selected_pages)
                           and table.get("table_strategy", "lines") == table_strategy]
        selected_page_content = [page for page in inspection["pages"] if page["page"] in selected_pages]
        selected_text = "\n".join(page["text"] for page in selected_page_content) if selected_page_content else inspection["text"]
        # Focused page reads must carry the note body, not just repeating report
        # headers. Broad initial inspection remains a compact navigation view.
        page_text_limit = (12000 if len(selected_pages) == 1 else 6000) if page_numbers and len(selected_pages) <= 3 else 2000
        if page_numbers and len(selected_pages) > 3:
            page_text_limit = max(400, 12000 // len(selected_pages))
        previews = [{**{key: value for key, value in table.items() if key not in {"rows", "context_text", "cell_origins", "row_origins", "raw_machine_rows"}},
                     "context_text": table.get("context_text", "")[:3000],
                     "source_header_quotes": table.get("source_header_quotes", table["original_columns"]),
                     "preview": [dict(zip(table["columns"], row)) for row in table["rows"][:8]] if all(len(row) == len(table["columns"]) for row in table["rows"]) else [],
                     "raw_preview": table["rows"][:8] if any(len(row) != len(table["columns"]) for row in table["rows"]) else None,
                     "preview_truncated": table["row_count"] > 8} for table in selected_tables[:30]]
        return {**manifest, "status": "ok", "tables": previews, "text": selected_text[:12000],
                "text_truncated": inspection["text_truncated"] or len(selected_text) > 12000,
                "pages": [{**page, "text": page["text"][:page_text_limit],
                           "text_truncated": len(page["text"]) > page_text_limit} for page in selected_page_content],
                "total_pages": inspection.get("total_pages"), "processed_pages": selected_pages,
                **({"selected_pages": selected_pages} if page_numbers else {}),
                "cached_pages": inspection.get("processed_pages", []),
                "inspection_complete": inspection.get("total_pages") is None or len(selected_pages) == inspection["total_pages"],
                "warnings": inspection["warnings"], "article": article,
                "document_metadata": inspection.get("document_metadata", {}),
                "publication_requires_explicit_contract": True}

    def _research_pdf(self, inspected, topic_matches):
        """Read complete cached page text, not the inspection's short preview."""
        from agentic_analytics.agent.tools.pdf_research import pdf_passages, pdf_title
        manifest = self.source(inspected["source_id"])
        if manifest["raw_sha256"] != inspected.get("raw_sha256"):
            raise DocumentError("Research source identity changed.", "SOURCE_HASH_MISMATCH")
        inspection = json.loads((self._directory(manifest["source_id"]) / "inspection.json").read_text())
        pages = inspection.get("pages", [])
        passages = pdf_passages(pages, topic_matches)
        selected = [passage["page"] for passage in passages]
        cached = sorted({page["page"] for page in pages})
        title = pdf_title(inspection.get("document_metadata", {}), pages, manifest["filename"])
        navigation = {**title, "filename": manifest["filename"], "mime_type": manifest["mime_type"],
                      "total_pages": inspection.get("total_pages"), "cached_pages": cached[:60],
                      "matched_pages": selected, "passages": passages,
                      "page_number_basis": "physical_pdf_1_based",
                      "passage_search_scope": "cached_pdf_pages"}
        if selected:
            navigation["suggested_inspection"] = {"source_id": manifest["source_id"], "page_numbers": selected}
        navigation["next_step"] = ("These are separate, contiguous passages from the listed physical PDF pages. "
            "Do not merge different sections, periods or lists. If a passage is shortened, inspect its page with "
            "suggested_inspection before claiming a complete list. For missing context or pages outside cached_pages, "
            "use find_source_pages with this source_id and the specific heading; document-wide completeness is not asserted.")
        content = "\n".join(page.get("text", "") for page in pages)
        return content, navigation

    def research_web(self, query, limit=3, domains=None):
        """Search, read a few public URLs, and return readable source cards."""
        if (not isinstance(query, str) or not 1 <= len(query.strip()) <= 500 or type(limit) is not int
                or not 1 <= limit <= 3 or domains is not None and (not isinstance(domains, list)
                or len(domains) > 5 or any(not isinstance(domain, str) or not domain.strip() for domain in domains))):
            raise DocumentError("Research query and limit exceed their bounds.")
        read_deadline = time.monotonic() + 120
        search_deadline = min(read_deadline, time.monotonic() + 45)
        lowered = _search_text(query)
        inferred_domains = {
            "kkb": ["kkb.com.tr"],
            "kredi kayit burosu": ["kkb.com.tr"],
            "tcmb": ["tcmb.gov.tr"],
            "bddk": ["bddk.org.tr"],
            "tüik": ["tuik.gov.tr"],
            "tuik": ["tuik.gov.tr"],
            "borsa istanbul": ["borsaistanbul.com"],
            "bist": ["borsaistanbul.com"],
            "imkb": ["borsaistanbul.com"],
            "ise": ["borsaistanbul.com"],
            "kap": ["kap.org.tr"],
            "kamuyu aydinlatma platformu": ["kap.org.tr"],
        }
        from agentic_analytics.agent.tools.search_backend import search_domains
        preferred = domains or search_domains(query) or next((values for key, values in inferred_domains.items()
            if re.search(r"(?<!\w)" + re.escape(_search_text(key)) + r"(?!\w)", lowered)), None)
        if preferred:
            preferred = [domain.strip().casefold() for domain in preferred]
            if any(not re.fullmatch(r"[a-z0-9](?:[a-z0-9-]*[a-z0-9])?(?:\.[a-z0-9](?:[a-z0-9-]*[a-z0-9])?)+", domain) for domain in preferred):
                raise DocumentError("Domains must be hostname names without paths, ports or wildcards.")
        def allowed(url):
            hostname = (parse.urlsplit(url).hostname or "").casefold()
            return parse.urlsplit(url).scheme in {"https", "http"} and (not preferred or any(
                hostname == domain or hostname.endswith("." + domain) for domain in preferred))
        def same_source_host(left, right):
            # Public sites often serve the bare root while their own links use
            # www. Treat that alias as the same site, keeping all other hosts
            # subject to the original domain and public-address checks.
            host = lambda url: (parse.urlsplit(url).hostname or "").casefold().removeprefix("www.")
            return bool(host(left)) and host(left) == host(right)
        registry = OFFICIAL_SOURCE_REGISTRY.get(preferred[0]) if preferred else None
        topic_terms = [term for term in re.findall(r"[\wçğıöşü]+", lowered)
                   if len(term) >= 4 and term not in {"tcmb", "bddk", "tüik", "tuik", "yılında", "raporları", "için", "kaynak", "bağlantısı"}]
        if preferred:
            domain_names = " ".join(preferred)
            topic_terms = [term for term in topic_terms if term not in domain_names]
        month_pairs = dict(zip(
            ["january", "february", "march", "april", "may", "june", "july", "august", "september", "october", "november", "december"],
            ["ocak", "subat", "mart", "nisan", "mayis", "haziran", "temmuz", "agustos", "eylul", "ekim", "kasim", "aralik"]))
        topic_terms.extend(value for key, value in month_pairs.items() if key in lowered and value not in topic_terms)
        variants = registry["search_variants"] if registry else ("{query}",)
        variant_queries = [variant.format(query=query) for variant in variants]
        search_queries = ["site:" + preferred[0] + " " + variant for variant in variant_queries] if preferred and not search_domains(query) else variant_queries
        searches = []
        for search_query in search_queries[:3]:
            if time.monotonic() >= search_deadline:
                break
            search = self.web_search(search_query, limit=min(10, max(5, limit * 2)), _deadline=search_deadline)
            searches.append(search)
            if any(not item.get("discovery_only") and not item.get("entity_verification_required") for item in search.get("results", [])):
                break
        search_diagnostics = [{key: search.get(key) for key in ("query", "status", "code", "source_backend", "provider_attempts", "budget_exhausted", "warnings")}
                              for search in searches]
        results, seen_urls = [], set()
        for search in searches:
            if search.get("status") != "ok":
                continue
            for item in search.get("results", []):
                url = item.get("url", "").split("#", 1)[0]
                if url and url not in seen_urls and allowed(url):
                    seen_urls.add(url)
                    results.append(item)
        if not results:
            if preferred:
                # When keyword search cannot locate a page, start at the
                # explicitly requested institution and follow relevant archive
                # links. These roots remain discovery pages, never answer proof.
                results = [{"url": "https://" + domain + "/", "title": domain, "discovery_only": True} for domain in preferred]
            else:
                return {"status": "unavailable", "research_status": "unavailable", "code": "SEARCH_NO_RESULTS",
                        "message": "Search did not return usable result URLs.", "query": query, "sources": [],
                        "searches": search_diagnostics}
        if preferred:
            matching = [item for item in results if allowed(item.get("url", ""))]
            results = matching
            if not results:
                return {"status": "unavailable", "research_status": "unavailable",
                        "code": "OFFICIAL_SOURCE_NOT_FOUND",
                        "message": "İstenen resmi kurum alanında uygun kaynak bulunamadı.",
                        "query": query, "sources": [], "failures": []}
        requested_years = set(re.findall(r"\b(?:19|20)\d{2}\b", lowered))
        month_numbers = {name: index for index, pair in enumerate(month_pairs.items(), 1) for name in pair}
        month_pattern = "|".join(month_numbers)
        def source_dates(value):
            normalized = _search_text(value)
            observed = set()
            for day, month, year in re.findall(r"\b(0?[1-9]|[12]\d|3[01])\s+(" + month_pattern + r")\s+((?:19|20)\d{2})\b", normalized):
                observed.add((int(year), month_numbers[month], int(day)))
            for year, month, day in re.findall(r"\b((?:19|20)\d{2})-(\d{2})-(\d{2})\b", normalized):
                observed.add((int(year), int(month), int(day)))
            return observed
        requested_dates = source_dates(query)
        topical_terms = [term for term in topic_terms if not term.isdigit() and term not in month_numbers]
        topic_aliases = {"results": ("results", "earnings", "sonuclar", "financial performance"),
                         "earnings": ("earnings", "results", "financial performance"),
                         "financial": ("financial", "finansal"), "finansal": ("finansal", "financial")}
        ownership_terms = ("ortaklar", "ortaklari", "ortaklik", "hissedar", "hissedarlari", "hissedarlar",
                           "shareholder", "shareholders", "shareholding", "ownership")
        topic_aliases.update({term: ownership_terms for term in ownership_terms})
        founding_terms = ("kurulus", "kurulusu", "kurulusunda", "kurulan", "kuruldu", "kurulmustur", "kurucu",
                          "founding", "founded", "founders", "established")
        topic_aliases.update({term: founding_terms for term in founding_terms})
        def topic_matches(value):
            normalized = " ".join(_search_text(value).split())
            return sum(any(re.search(r"(?<!\w)" + re.escape(alias) + r"(?!\w)", normalized)
                           for alias in topic_aliases.get(term, (term,))) for term in topical_terms)
        # These are language/topic aliases, not institution-specific URL rules.
        archive_terms = ["archive", "arşiv", "press release", "basın duyuru", "announcements", "duyurular"]
        specialized_terms = []
        if any(term in lowered for term in ownership_terms):
            specialized_terms += ["ortaklar", "ortaklık yapısı", "hissedarlar", "shareholders", "ownership", "shareholding",
                                  "hakkımızda", "about us", "about-us", "kurumsal yönetim", "corporate governance"]
        if any(term in lowered for term in ("interest", "rate", "faiz", "policy", "politika")):
            specialized_terms += ["monetary policy", "para politikası", "interest rates", "faiz oranları"]
        if any(term in lowered for term in ("report", "rapor", "financial", "finansal")):
            specialized_terms += ["financial results", "finansal sonuç", "annual reports", "faaliyet rapor", "investor relations", "yatırımcı ilişkileri", "earnings", "finansal bilgiler", "publications and results", "financial statements", "mali tablolar"]
        financial_navigation = any(term in lowered for term in ("report", "rapor", "financial", "finansal"))
        specialized_terms = [_search_text(term) for term in specialized_terms]
        archive_terms = [_search_text(term) for term in archive_terms] + specialized_terms
        day_numbers = set(re.findall(r"\b(?:[1-9]|[12]\d|3[01])\b", lowered))
        def identity(item):
            # Ancestor URL segments repeat the entire navigation hierarchy on
            # some sites. Rank the actual link label and leaf, not that menu.
            path = parse.unquote(parse.urlsplit(item.get("url", "")).path).rstrip("/")
            return _search_text(item.get("title", "") + " " + path.rsplit("/", 1)[-1])
        def different_year(item):
            years = set(re.findall(r"(?<!\d)(?:19|20)\d{2}(?!\d)", identity(item)))
            return bool(requested_years and years and not requested_years.intersection(years))
        def relevance(item):
            label = identity(item)
            topical = topic_matches(label)
            return (4 * topical + 6 * sum(year in label for year in requested_years) * bool(topical or item.get("topic_context"))
                    + 24 * bool(item.get("topic_context") and str(item.get("title", "")).strip() in requested_years)
                    + 40 * bool(requested_dates.intersection(source_dates(label)))
                    + 2 * bool(item.get("in_main_content")) - 2 * bool(item.get("in_navigation"))
                    + 3 * sum(term in label for term in specialized_terms)
                    + sum(bool(re.search(r"\b" + number + r"\b", item.get("title", ""))) for number in day_numbers))
        results.sort(key=relevance, reverse=True)
        sources, deferred_sources, failures, attempts = [], [], [], 0
        while results:
            if len(sources) >= limit:
                break
            if attempts >= 18 or time.monotonic() >= read_deadline:
                break
            results.sort(key=lambda item: relevance(item) - 2 * item.get("discovery_depth", 0), reverse=True)
            result = results.pop(0)
            if different_year(result):
                failures.append({"url": result.get("url"), "code": "OUT_OF_DATE_LINK", "message": "Link explicitly names a different year; no fetch attempted."})
                continue
            attempts += 1
            try:
                inspected = self.inspect_source(url=result["url"], _deadline=read_deadline)
                text = inspected.get("text", "").strip()
                article = inspected.get("article", {})
                content = article.get("article_body") or article.get("readable_text") or text
                pdf_navigation = {}
                if inspected.get("mime_type") == "application/pdf" or inspected.get("filename", "").casefold().endswith(".pdf"):
                    content, pdf_navigation = self._research_pdf(inspected, topic_matches)
                if not content:
                    raise DocumentError("Source contained no readable text.", "EMPTY_SOURCE")
                source_url = inspected.get("source_url") or result["url"]
                if not allowed(source_url):
                    raise DocumentError("The fetched source redirected outside the requested source domains.", "OFFICIAL_SOURCE_REDIRECT")
                # A report landing page can be the path to the actual data.
                # Follow only bounded same-host links in addition to search hits.
                linked = [link for link in article.get("document_links", []) if allowed(link.get("url", ""))
                          and same_source_host(link["url"], source_url) and link["url"] not in seen_urls]
                depth = result.get("discovery_depth", 0)
                archive_links = []
                for link in article.get("source_links", []):
                    link = {**link, "topic_context": bool(topic_matches(identity(result)) or result.get("topic_context"))}
                    target = link.get("url", "").split("#", 1)[0]
                    label = identity(link)
                    precise_date = bool(requested_dates.intersection(source_dates(label)))
                    report_navigation = (not financial_navigation or precise_date or bool(re.search(
                        r"financ|finans|rapor|report|statement|earnings|investor|yatirimci|publications|mali.tablo", label))
                        or str(link.get("title", "")).strip() in requested_years and link.get("topic_context"))
                    if (depth < 4 or precise_date and depth < 6) and (target not in seen_urls and allowed(target)
                            and same_source_host(target, source_url) and not different_year(link)
                            and report_navigation
                            and (relevance(link) >= 2 or any(term in label for term in archive_terms))):
                        archive_links.append({**link, "url": target})
                candidates = {}
                for link in [*linked, *archive_links]:
                    if different_year(link):
                        continue
                    previous = candidates.get(link["url"])
                    if previous is None or relevance(link) > relevance(previous):
                        candidates[link["url"]] = link
                ranked_links = sorted(candidates.values(), key=lambda link: relevance(link) + sum(
                    term in identity(link) for term in archive_terms)
                    + 3 * sum(term in identity(link) for term in specialized_terms), reverse=True)
                for link in ranked_links[:3]:
                    seen_urls.add(link["url"])
                    results.append({**link, "discovered_from": source_url, "discovery_depth": depth + 1})
                path_text = parse.urlsplit(source_url).path.casefold()
                verified_title = pdf_navigation.get("title") or article.get("title") or ""
                source_title = verified_title or result.get("title", "")
                title_text = _search_text(source_title)
                searchable = _search_text(" ".join([verified_title, article.get("description", ""), content]))
                from agentic_analytics.agent.tools.search_backend import rank_search_results
                checked, _, _ = rank_search_results(query, [{"url": source_url,
                    "title": verified_title, "snippet": content[:2500]}])
                if result.get("entity_verification_required") and (not checked or checked[0]["entity_verification_required"]):
                    raise DocumentError("The requested issuer or named subject is not established in the fetched content.", "SOURCE_ENTITY_UNVERIFIED")
                if (result.get("discovery_only") or not path_text.strip("/") or re.fullmatch(r"/[a-z]{2}(?:-[a-z]{2})?/?", path_text)
                        or path_text.endswith(('/kurlar/kurlar_tr.html', '/main+page+site+area/bugun'))):
                    raise DocumentError("Source is a generic landing page.", "GENERIC_SOURCE")
                if (article.get("link_count", 0) > 20 and not article.get("article_body")
                        and not topic_matches(title_text)
                        and not any(year in title_text + path_text + str(article.get("date_published", "")) for year in requested_years)):
                    raise DocumentError("Navigation-heavy source has no topic-specific title or requested-period identity.", "GENERIC_SOURCE")
                if topical_terms and topic_matches(searchable) < min(2, len(topical_terms)):
                    raise DocumentError("Source content does not match the research topic.", "IRRELEVANT_SOURCE")
                source_years = set(re.findall(r"\b(?:19|20)\d{2}\b", title_text + " " + path_text))
                if requested_years and source_years and not requested_years.intersection(source_years) and not any(year in content for year in requested_years):
                    raise DocumentError("Source is outside the requested year.", "OUT_OF_DATE_SOURCE")
                identified_period = title_text + " " + path_text + " " + str(article.get("date_published", "")) + " " + content[:1500]
                if requested_years and not any(year in identified_period for year in requested_years):
                    raise DocumentError("The requested year is not established by the source title, address, publication date or opening content.", "SOURCE_PERIOD_UNVERIFIED")
                if requested_dates and not requested_dates.intersection(source_dates(identified_period + " " + content)):
                    raise DocumentError("Source does not establish the requested full calendar date.", "SOURCE_DATE_UNVERIFIED")
                paragraphs = [line.strip() for line in content.splitlines() if line.strip()]
                if pdf_navigation.get("passages"):
                    excerpt = pdf_navigation["passages"][0]["text"]
                elif len(content) > 3000 and len(paragraphs) > 1:
                    ranked = sorted(enumerate(paragraphs), key=lambda pair: topic_matches(pair[1]), reverse=True)
                    # Keep a contiguous passage around the strongest match.
                    # Joining isolated keyword hits can put members under an
                    # earlier shareholders heading or detach numbers from units.
                    start = max(0, ranked[0][0] - 1)
                    excerpt = _safe_truncate("\n".join(paragraphs[start:]), 3000)
                else:
                    excerpt = _safe_truncate(content, 3000)
                card = {
                    "title": source_title,
                    "url": source_url,
                    "domain": parse.urlsplit(source_url).hostname,
                    "snippet": article.get("description") or result.get("snippet", ""),
                    "date_published": article.get("date_published"),
                    **({"search_published_at": result["published_at"][:100],
                        "search_publication_date_basis": "search_metadata_unverified"}
                       if isinstance(result.get("published_at"), str) and result["published_at"].strip() else {}),
                    "date_modified": article.get("date_modified"),
                    "author": article.get("author"),
                    "content": excerpt,
                    "content_truncated": bool(inspected.get("text_truncated") or len(content) > len(excerpt)),
                    "tables": [table for table in inspected.get("tables", []) if table.get("preview")][:3],
                    "document_links": linked[:5],
                    "discovery_links": ranked_links[:5],
                    "discovered_from": result.get("discovered_from"),
                    "discovery_depth": depth,
                    "raw_sha256": inspected.get("raw_sha256"),
                    "inspection_complete": inspected.get("inspection_complete", True),
                    "warnings": inspected.get("warnings", []),
                    "source_id": inspected["source_id"],
                    "verification": "direct_public_fetch",
                    "content_is_untrusted_data": True,
                    **pdf_navigation,
                }
                discovery_index = (archive_links and not article.get("article_body") and article.get("link_count", 0) > 20
                                   and not topic_matches(title_text))
                discovery_index = discovery_index or bool(requested_dates and any(
                    requested_dates.intersection(source_dates(identity(link))) for link in archive_links)
                    and not requested_dates.intersection(source_dates(title_text)))
                if linked and not card["tables"] or discovery_index:
                    card["source_role"] = "discovery_index"
                    deferred_sources.append(card)
                else:
                    sources.append(card)
            except (DocumentError, OSError, ValueError, KeyError) as exc:
                failures.append({"url": result.get("url"), "code": getattr(exc, "code", "SOURCE_READ_FAILED"), "message": str(exc)})
        sources.extend(deferred_sources[:max(0, limit - len(sources))])
        if not sources:
            return {"status": "unavailable", "research_status": "unavailable", "code": "NO_READABLE_SOURCES",
                    "message": "Search results were found, but no result had readable public content.",
                    "query": query, "sources": [], "failures": failures, "searches": search_diagnostics,
                    "budget_exhausted": time.monotonic() >= read_deadline}
        return {"status": "ok", "research_status": "completed", "query": query,
                "sources": sources, "failures": failures, "searched": len(seen_urls),
                "searches": search_diagnostics,
                "read": len(sources), "attempted_reads": attempts, "content_is_untrusted_data": True,
                "next_step": "Sources are inspected and registered, not yet analytical datasets. For requested calculations: inspect returned source_id or document links, publish the selected table with an evidence-backed contract, then calculate and chart. Use only claims supported by returned evidence."}

    def combine_source_tables(self, source_id, table_ids, reason):
        """Explicitly concatenate compatible page tables, preserving row origins.

        Matching headers merely make tables eligible. The caller must select
        them and explain why they continue the same table; original candidates
        remain available and publication still validates the declared grain.
        """
        if (not isinstance(table_ids, list) or not 2 <= len(table_ids) <= 10 or len(set(table_ids)) != len(table_ids)
                or not isinstance(reason, str) or not 10 <= len(reason.strip()) <= 1000):
            raise DocumentError("Select 2 to 10 distinct continuation tables and explain their common scope.", "INVALID_TABLE_COMBINATION")
        manifest = self.source(source_id)
        self.inspect_source(source_id=source_id)
        cache = self._directory(source_id) / "inspection.json"
        inspection = json.loads(cache.read_text())
        by_id = {table["table_id"]: table for table in inspection["tables"]}
        if any(table_id not in by_id for table_id in table_ids):
            raise DocumentError("Every selected table must have been inspected in this source.", "INVALID_TABLE_COMBINATION")
        tables = [by_id[table_id] for table_id in table_ids]
        first = tables[0]
        if (any(table["origin"] != "parsed" or table.get("page") is None or table.get("combination") or table.get("preparation") for table in tables)
                or not all(first["original_columns"].values())
                or any(table["original_columns"] != first["original_columns"] for table in tables)):
            raise DocumentError("Only original parsed PDF tables with identical nonempty headers may be combined.", "INCOMPATIBLE_TABLE_HEADERS")
        page_numbers = [table["page"] for table in tables]
        if page_numbers != list(range(page_numbers[0], page_numbers[0] + len(tables))):
            raise DocumentError("Select tables from consecutive pages in page order.", "NONCONTIGUOUS_TABLE_PAGES")
        rows = [row for table in tables for row in table["rows"]]
        if len(rows) > self.max_rows:
            raise DocumentError("Combined table exceeds the row limit.", "TABLE_TOO_LARGE")
        combination = {"table_ids": table_ids, "reason": reason.strip(), "operation": "explicit_consecutive_page_concatenation",
                       "raw_sha256": manifest["raw_sha256"], "numeric_transformation": False}
        table_id = "table_c" + hashlib.sha256(_canonical(combination)).hexdigest()[:20]
        combined = {**first, "table_id": table_id, "rows": rows, "row_count": len(rows), "source_pages": page_numbers,
                    "row_origins": [{"page": table["page"], "table_id": table["table_id"], "candidate_row": index}
                                    for table in tables for index in range(1, len(table["rows"]) + 1)],
                    "context_text": "\n".join(table.get("context_text", "") for table in tables)[:200000], "combination": combination}
        if table_id not in by_id:
            if len(inspection["tables"]) >= 500:
                raise DocumentError("Source candidate cache limit reached.", "TABLE_LIMIT")
            inspection["tables"].append(combined)
            _write_json(cache, inspection)
        return {"status": "ok", "source_id": source_id,
                **{key: value for key, value in combined.items() if key not in {"rows", "context_text", "row_origins"}},
                "preview": [dict(zip(combined["columns"], row)) for row in rows[:8]],
                "preview_truncated": len(rows) > 8, "publication_requires_explicit_contract": True}

    def read_source_table(self, source_id, table_id, row_start=1, row_limit=30):
        """Read a bounded candidate row window with stable source row addresses."""
        if type(row_start) is not int or row_start < 1 or type(row_limit) is not int or not 1 <= row_limit <= 100:
            raise DocumentError("Read 1 to 100 candidate rows using a positive 1-based row_start.")
        table = self.review_candidate(source_id, table_id)
        rows = table.pop("rows")
        if row_start > len(rows):
            raise DocumentError("Requested row_start exceeds the candidate row count.")
        for key in ("context_text", "cell_origins", "row_origins", "raw_machine_rows"):
            table.pop(key, None)
        return {**table, "rows": [{"candidate_row": index + 1, "values": dict(zip(table["columns"], rows[index])) if len(rows[index]) == len(table["columns"]) else None,
                                   **({"raw_cells": rows[index]} if len(rows[index]) != len(table["columns"]) else {})}
                                  for index in range(row_start - 1, min(len(rows), row_start - 1 + row_limit))],
                "preview_truncated": row_start - 1 + row_limit < len(rows)}

    @staticmethod
    def _source_date_occurrences(value, source_format):
        """Find only complete literal dates in a bounded source header."""
        if not isinstance(value, str) or len(value) > 2000:
            raise DocumentError("A period header must be bounded source text.", "INVALID_UNPIVOT")
        patterns = {
            "english_dmy": r"(?<!\d)\d{1,2}\s+(?:January|February|March|April|May|June|July|August|September|October|November|December|Jan|Feb|Mar|Apr|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\s+\d{4}(?!\d)",
            "dmy": r"(?<!\d)\d{1,2}/\d{1,2}/\d{4}(?!\d)",
            "parenthesized_dmy": r"\(\d{1,2}/\d{1,2}/\d{4}\)",
            "year_month_slash": r"(?<!\d)\d{4}/\d{1,2}(?!\d)",
            "year_quarter_space": r"(?<!\d)\d{4} Q[1-4](?!\d)",
        }
        if source_format not in patterns:
            raise DocumentError("date_index requires an explicit period_format such as english_dmy; source_header leaves text unparsed.", "INVALID_UNPIVOT")
        return re.findall(patterns[source_format], value, re.I)

    def _period_header_recovery(self, table, unpivot):
        """Offer source-cell recipes, never dates supplied by the model."""
        header_rows = {source.get("row") for source in (unpivot.get("period_sources") or {}).values() if isinstance(source, dict)}
        if unpivot.get("header_row") is not None:
            header_rows.add(unpivot["header_row"])
        header_rows = sorted(row for row in header_rows if type(row) is int and 1 <= row <= len(table["rows"]))[:4]
        value_columns = sorted(unpivot["columns"], key=table["columns"].index)
        source_format = unpivot.get("period_format", "source_header")
        formats = [source_format] if source_format != "source_header" else ["english_dmy", "dmy", "parenthesized_dmy"]
        candidates, row_cells, signatures = [], [], set()
        for number in header_rows:
            row = table["rows"][number - 1]
            row_cells.append({"row": number, "cells": [{"column": column, "text": value[:200] if isinstance(value, str) else value,
                                                        "text_truncated": isinstance(value, str) and len(value) > 200}
                                                       for column, value in zip(table["columns"], row)]})
            for start in range(len(row)):
                for width in range(1, min(4, len(row) - start) + 1):
                    columns, parts = table["columns"][start:start + width], row[start:start + width]
                    if any(not isinstance(part, str) or not part.strip() for part in parts):
                        continue
                    parts = [part.strip() for part in parts]
                    for separator in (("",) if width == 1 else ("", " ")):
                        joined = separator.join(parts)
                        if len(joined) > 2000:
                            continue
                        positions, offset = [], 0
                        for column, part in zip(columns, parts):
                            positions.append((column, offset, offset + len(part)))
                            offset += len(part) + len(separator)
                        for fmt in formats:
                            try:
                                literal_dates = self._source_date_occurrences(joined, fmt)
                            except DocumentError:
                                continue
                            if len(literal_dates) > 4:
                                continue
                            dates, cursor = [], 0
                            for index, literal in enumerate(literal_dates):
                                begin = joined.find(literal, cursor)
                                cursor = begin + len(literal)
                                try:
                                    normalized = self._source_period_label(literal, fmt)
                                except DocumentError:
                                    continue
                                spans = [column for column, left, right in positions if left < cursor and right > begin]
                                dates.append({"date_index": index, "source_text": literal, "normalized_date": normalized, "source_columns": spans})
                            if not dates:
                                continue
                            signature = (number, fmt, tuple((date["source_text"], tuple(date["source_columns"])) for date in dates))
                            if signature in signatures:
                                continue
                            signatures.add(signature)
                            candidates.append({"row": number, "columns": columns, "separator": separator,
                                               "joined_source_text": joined, "period_format": fmt, "dates": dates})
        candidates.sort(key=lambda candidate: (len(candidate["columns"]), len(candidate["joined_source_text"])))
        suggested, selected, eligible = None, None, []
        for candidate in candidates:
            if len(candidate["dates"]) != len(value_columns):
                continue
            if not all(column in date["source_columns"] for column, date in zip(value_columns, candidate["dates"])):
                continue
            eligible.append(candidate)
        possible_dates = {tuple(date["normalized_date"] for date in candidate["dates"]) for candidate in eligible}
        if eligible and len(possible_dates) == 1:
            candidate = eligible[0]
            suggested = {"period_format": candidate["period_format"], "period_sources": {
                column: {"row": candidate["row"], "columns": candidate["columns"], "separator": candidate["separator"], "date_index": date["date_index"]}
                for column, date in zip(value_columns, candidate["dates"])}}
            selected = candidate
        shown = ([selected] if selected else []) + [candidate for candidate in candidates if candidate is not selected]
        return {"source_table_id": table["table_id"], "header_rows": row_cells, "value_columns_in_source_order": value_columns,
                "date_candidates": shown[:8], "suggested_unpivot_update": suggested,
                "ambiguous_date_interpretations": len(possible_dates) > 1,
                "mapping_basis": "Suggestion requires exactly one complete valid source date per value column in source order, each date overlapping that column's header-cell span. This is an explicit retry recipe, not an automatic transformation.",
                "next_step": "Keep selected data rows/columns and labels. Retry prepare_source_table with the suggested unpivot update if present. Otherwise review the actual header-cell candidates and choose period_sources/period_format; never retype a date. For one verified snapshot, select just its actual values and publish static with source-date scope."}

    @staticmethod
    def _source_period_label(value, source_format):
        if source_format == "source_header":
            return value
        try:
            if source_format in {"dmy", "parenthesized_dmy"}:
                if source_format == "parenthesized_dmy":
                    if not value.startswith("(") or not value.endswith(")"):
                        raise ValueError
                    value = value[1:-1]
                return datetime.strptime(value.strip(), "%d/%m/%Y").date().isoformat()
            if source_format == "year_month_slash":
                return datetime.strptime(value.strip(), "%Y/%m").strftime("%Y-%m")
            if source_format == "english_dmy":
                try:
                    return datetime.strptime(value.strip(), "%d %B %Y").date().isoformat()
                except ValueError:
                    return datetime.strptime(value.strip(), "%d %b %Y").date().isoformat()
            if source_format == "year_quarter_space" and re.fullmatch(r"\d{4} Q[1-4]", value.strip()):
                return value.strip().replace(" ", "-")
        except (ValueError, TypeError) as exc:
            raise DocumentError("Period header does not match the declared source format.", "INVALID_SOURCE_DATE_FORMAT") from exc
        raise DocumentError("Unsupported source period header format.", "INVALID_SOURCE_DATE_FORMAT")

    def prepare_source_table(self, source_id, table_id, selected_rows=None, selected_columns=None, unpivot=None, join_columns=None):
        """Select source cells and optionally turn period columns into rows.

        No model-supplied numerical values or expressions are accepted. Period
        labels come verbatim from original column headers, and every output
        cell retains its candidate-cell address.
        """
        manifest = self.source(source_id)
        self.inspect_source(source_id=source_id)
        cache = self._directory(source_id) / "inspection.json"
        inspection = json.loads(cache.read_text())
        candidate = next((table for table in inspection["tables"] if table["table_id"] == table_id), None)
        if candidate is None:
            raise DocumentError("Unknown candidate table.")
        table = self._reviewed_table(source_id, candidate)
        if ((table["origin"] == "ocr" or table.get("missing_formula_cache") or table.get("layout_review_required")) and "review" not in table):
            raise DocumentError("Review uncertain extraction cells before preparing this candidate.", "TABLE_REVIEW_REQUIRED")
        selected_rows = list(range(1, len(table["rows"]) + 1)) if selected_rows is None else selected_rows
        selected_columns = table["columns"] if selected_columns is None else selected_columns
        if (not isinstance(selected_rows, list) or not selected_rows or len(set(selected_rows)) != len(selected_rows)
                or any(type(number) is not int or not 1 <= number <= len(table["rows"]) for number in selected_rows)
                or not isinstance(selected_columns, list) or not selected_columns or len(set(selected_columns)) != len(selected_columns)
                or any(column not in table["columns"] for column in selected_columns)):
            raise DocumentError("Select distinct existing candidate row numbers and columns.", "INVALID_TABLE_SELECTION")
        indexes = {column: table["columns"].index(column) for column in selected_columns}
        columns = list(selected_columns)
        original_columns = {column: table["original_columns"][column] for column in columns}
        rows, origins = [], []
        if unpivot is not None:
            if (not isinstance(unpivot, dict) or not {"columns", "period_column", "value_column"}.issubset(unpivot)
                    or set(unpivot) - {"columns", "period_column", "value_column", "header_row", "period_format", "period_sources"}):
                raise DocumentError("Unpivot requires columns, period_column and value_column.", "INVALID_UNPIVOT")
            measures, period, value = unpivot["columns"], unpivot["period_column"], unpivot["value_column"]
            if (not isinstance(measures, list) or not measures or len(set(measures)) != len(measures)
                    or any(column not in columns for column in measures)
                    or any(not isinstance(name, str) or not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]{0,100}", name) for name in (period, value))
                    or period == value or period in columns or value in columns):
                raise DocumentError("Unpivot must name existing value columns and two new distinct output identifiers.", "INVALID_UNPIVOT")
            dimensions = [column for column in columns if column not in measures]
            header_row = unpivot.get("header_row")
            if header_row is not None and (type(header_row) is not int or not 1 <= header_row <= len(table["rows"]) or header_row in selected_rows):
                raise DocumentError("Use an actual candidate header row outside the selected data rows.", "INVALID_UNPIVOT")
            period_sources = unpivot.get("period_sources")
            if period_sources is not None:
                if not isinstance(period_sources, dict) or set(period_sources) != set(measures):
                    raise DocumentError("period_sources must map each unpivoted value column to its actual source date-header cells.", "INVALID_UNPIVOT")
                for source in period_sources.values():
                    if (not isinstance(source, dict) or not {"row", "columns", "separator"} <= set(source)
                            or set(source) - {"row", "columns", "separator", "date_index"}
                            or type(source["row"]) is not int or not 1 <= source["row"] <= len(table["rows"]) or source["row"] in selected_rows
                            or not isinstance(source["columns"], list) or not 1 <= len(source["columns"]) <= 4
                            or len(set(source["columns"])) != len(source["columns"])
                            or any(column not in table["columns"] for column in source["columns"])
                            or source["separator"] not in {"", " "}):
                        raise DocumentError("Period headers must use 1 to 4 actual cells from a source header row with an explicit text separator.", "INVALID_UNPIVOT")
                    if "date_index" in source and (type(source["date_index"]) is not int or not 0 <= source["date_index"] <= 3):
                        raise DocumentError("date_index selects source date occurrence 0 to 3; it cannot supply a new date.", "INVALID_UNPIVOT")
            columns = [*dimensions, period, value]
            original_columns = {column: table["original_columns"][column] for column in dimensions}
            original_columns.update({period: "source period header", value: "source cell value"})
            for number in selected_rows:
                source_row = table["rows"][number - 1]
                for measure in measures:
                    if period_sources is not None:
                        source = period_sources[measure]
                        parts = [table["rows"][source["row"] - 1][table["columns"].index(column)] for column in source["columns"]]
                        if any(not isinstance(part, str) or not part.strip() for part in parts):
                            raise DocumentError("Selected source date-header cells must contain text. Read this header row and map its populated cells with period_sources; header cells need not be in selected_columns. For a one-date snapshot, prepare only the actual metric and value columns and publish frequency=static with the source date in scope metadata.", "INVALID_UNPIVOT")
                        label = source["separator"].join(part.strip() for part in parts)
                        period_origin = {"parts": [{"candidate_row": source["row"], "candidate_column": column} for column in source["columns"]],
                                         "separator": source["separator"]}
                        if "date_index" in source:
                            source_format = unpivot.get("period_format", "source_header")
                            dates = self._source_date_occurrences(label, source_format)
                            if len(dates) > 4 or source["date_index"] >= len(dates):
                                raise DocumentError("date_index does not select an actual complete date in the source header. Parsed source dates: " + json.dumps(dates[:4]) + "; joined source text: " + repr(label[:500]) + ". Choose actual adjacent header cells, their separator and matching period_format.", "INVALID_UNPIVOT")
                            period_origin.update(joined_source_text=label, date_index=source["date_index"],
                                                 source_format=source_format, matched_source_date=dates[source["date_index"]])
                            label = dates[source["date_index"]]
                    else:
                        label = table["rows"][header_row - 1][indexes[measure]] if header_row is not None else (
                            table.get("source_header_quotes", {}).get(measure) or table["original_columns"][measure])
                        period_origin = {"candidate_row": header_row, "candidate_column": measure} if header_row else {"candidate_header": measure}
                    if not isinstance(label, str) or not label.strip():
                        raise DocumentError("Selected period header cell is empty. Read the header row and use period_sources to point to actual date cells, which may be above other columns. selected_columns should contain data values and labels only.", "INVALID_UNPIVOT")
                    try:
                        label = self._source_period_label(label, unpivot.get("period_format", "source_header"))
                    except DocumentError as exc:
                        recovery = self._period_header_recovery(table, unpivot)
                        hint = recovery["suggested_unpivot_update"]
                        message = str(exc) + " Source header cells may split a date across columns or contain multiple dates. "
                        message += ("Retry with this source-derived unpivot update: " + json.dumps(hint) if hint else
                                    "Read the supplied header_rows/date_candidates and select actual adjacent cells with period_sources and date_index.")
                        raise DocumentError(message, exc.code, recovery=recovery) from exc
                    rows.append([*[source_row[indexes[column]] for column in dimensions], label, source_row[indexes[measure]]])
                    origins.append({**{column: {"candidate_row": number, "candidate_column": column} for column in dimensions},
                                    period: period_origin,
                                    value: {"candidate_row": number, "candidate_column": measure}})
        else:
            rows = [[table["rows"][number - 1][indexes[column]] for column in columns] for number in selected_rows]
            origins = [{column: {"candidate_row": number, "candidate_column": column} for column in columns} for number in selected_rows]
        dimension_only = [column for column in table.get("dimension_only_columns", []) if column in columns]
        if unpivot and any(column in table.get("dimension_only_columns", []) for column in unpivot["columns"]):
            dimension_only.append(unpivot["value_column"])
        if join_columns is not None:
            if (not isinstance(join_columns, dict) or set(join_columns) != {"columns", "output", "separator"}
                    or not isinstance(join_columns["columns"], list) or not 2 <= len(join_columns["columns"]) <= 10
                    or len(set(join_columns["columns"])) != len(join_columns["columns"])
                    or any(column not in columns for column in join_columns["columns"])
                    or not isinstance(join_columns["output"], str) or not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]{0,100}", join_columns["output"])
                    or join_columns["output"] in columns or join_columns["separator"] not in {"", " ", " / "}):
                raise DocumentError("Join 2 to 10 existing label columns into one new dimension using an explicit separator.", "INVALID_LABEL_JOIN")
            joined, output = join_columns["columns"], join_columns["output"]
            indexes_now = {column: index for index, column in enumerate(columns)}
            retained = [column for column in columns if column not in joined]
            updated_rows, updated_origins = [], []
            for row, origin in zip(rows, origins):
                label = join_columns["separator"].join(str(row[indexes_now[column]] or "") for column in joined).strip()
                if len(label) > 20000:
                    raise DocumentError("Joined label exceeds its text bound.", "TABLE_TOO_LARGE")
                updated_rows.append([*[row[indexes_now[column]] for column in retained], label])
                updated_origins.append({**{column: origin[column] for column in retained}, output: {"parts": [origin[column] for column in joined]}})
            columns, rows, origins = [*retained, output], updated_rows, updated_origins
            original_columns = {column: original_columns[column] for column in retained}
            original_columns[output] = "joined source label"
            dimension_only = [column for column in dimension_only if column in retained] + [output]
        if len(rows) > self.max_rows or len(columns) > self.max_columns:
            raise DocumentError("Prepared table exceeds row or column bounds.", "TABLE_TOO_LARGE")
        recipe = {"source_table_id": table_id, "selected_rows": selected_rows, "selected_columns": selected_columns,
                  "unpivot": unpivot, "join_columns": join_columns, "numeric_values_changed": False, "raw_sha256": manifest["raw_sha256"]}
        if table.get("review"):
            def reviewed_origin(value):
                if isinstance(value, list):
                    return [reviewed_origin(item) for item in value]
                if not isinstance(value, dict):
                    return value
                annotated = {key: reviewed_origin(item) for key, item in value.items()}
                if "candidate_row" in annotated:
                    annotated["review_row"] = annotated.pop("candidate_row")
                    annotated["review_sha256"] = table["review"]["review_sha256"]
                return annotated
            origins = reviewed_origin(origins)
        prepared_id = "table_c" + hashlib.sha256(_canonical(recipe)).hexdigest()[:20]
        prepared = {**table, "table_id": prepared_id, "columns": columns, "original_columns": original_columns,
                    "source_header_quotes": {column: table.get("source_header_quotes", {}).get(column, original_columns[column]) for column in columns},
                    "rows": rows, "row_count": len(rows), "preparation": recipe,
                    "dimension_only_columns": dimension_only,
                    "cell_origins": origins, "context_text": (table.get("context_text", "") + "\n" + inspection["text"])[:200000]}
        prepared["unit_caption"] = table.get("unit_caption", "") + "\n" + _unit_caption(table.get("context_text", ""))
        contexts = table.get("unit_contexts", {column: [header] for column, header in table["original_columns"].items()})
        prepared["unit_contexts"] = {column: contexts.get(column, []) for column in columns}
        if unpivot:
            prepared["unit_contexts"][unpivot["value_column"]] = [context for column in unpivot["columns"] for context in contexts.get(column, [])]
        reviewed_units = {**table.get("verified_unit_evidence", {}), **table.get("review", {}).get("unit_evidence", {})}
        verified_units = {column: quote for column, quote in reviewed_units.items() if column in columns}
        if unpivot:
            quotes = [reviewed_units.get(column) for column in unpivot["columns"]]
            if quotes and quotes[0] and all(quote == quotes[0] for quote in quotes):
                verified_units[unpivot["value_column"]] = quotes[0]
        prepared["verified_unit_evidence"] = verified_units
        # Review provenance is retained as evidence, but its source row set must
        # never replace transformed rows during subsequent candidate reads.
        if prepared.get("review"):
            prepared["preparation"]["source_review"] = prepared.pop("review")
        prepared.update(origin="parsed", missing_formula_cache=[], layout_review_required=False)
        if not any(item["table_id"] == prepared_id for item in inspection["tables"]):
            if len(inspection["tables"]) >= 500:
                raise DocumentError("Source candidate cache limit reached.", "TABLE_LIMIT")
            inspection["tables"].append(prepared)
            _write_json(cache, inspection)
        empty_columns = [column for index, column in enumerate(columns) if all(row[index] is None or not str(row[index]).strip() for row in rows)]
        return {"status": "ok", "source_id": source_id,
                **{key: value for key, value in prepared.items() if key not in {"rows", "context_text", "cell_origins", "row_origins", "raw_machine_rows"}},
                "preview": [dict(zip(columns, row)) for row in rows[:8]], "preview_truncated": len(rows) > 8,
                "publication_requires_explicit_contract": True,
                "publication_guidance": {"candidate_columns": columns, "identity_column_mapping": {column: column for column in columns},
                    "dimension_only_columns": dimension_only, "entirely_empty_columns": empty_columns,
                    "period_column": unpivot["period_column"] if unpivot else None,
                    "period_format": unpivot.get("period_format", "source_header") if unpivot else None},
                "next_step": "Publish this new table_id. Contract columns must equal the returned candidate columns, or provide column_mapping={actual_source_column:output_contract_column} covering every column. Join fragmented labels with join_columns before publication. Remove unwanted all-empty/header-only columns through a new explicit preparation; source date-header cells belong in period_sources, not selected data columns. source_header keeps raw labels and may not be a valid date: use the matching period_format and date_index for multiple complete dates sharing header cells. For one source snapshot, static is valid; do not invent a time axis."}

    def _reviewed_table(self, source_id, table):
        review_path = self._directory(source_id) / (table["table_id"] + "_review.json")
        if not review_path.exists():
            return table
        review = json.loads(review_path.read_text())
        content = {key: value for key, value in review.items() if key != "review_sha256"}
        if hashlib.sha256(_canonical(content)).hexdigest() != review.get("review_sha256") or review["raw_sha256"] != self.source(source_id)["raw_sha256"]:
            raise DocumentError("Reviewed table hash does not match its source.", "SOURCE_HASH_MISMATCH")
        return {**table, "rows": review["rows"], "row_count": len(review["rows"]),
                "review": {key: value for key, value in review.items() if key != "rows"}}

    def review_table(self, source_id, table_id, reviewed_rows, unit_evidence):
        """Trusted UI-only human review; never registered as a model tool."""
        manifest = self.source(source_id)
        self.inspect_source(source_id=source_id)
        inspection = json.loads((self._directory(source_id) / "inspection.json").read_text())
        table = next((item for item in inspection["tables"] if item["table_id"] == table_id), None)
        if table is None or not isinstance(reviewed_rows, list) or not 1 <= len(reviewed_rows) <= self.max_rows:
            raise DocumentError("Select a known table and provide reviewed rows within the row limit.")
        if not isinstance(unit_evidence, dict) or any(column not in table["columns"] or not isinstance(quote, str) or not 1 <= len(quote) <= 1000 for column, quote in unit_evidence.items()):
            raise DocumentError("Reviewed unit declarations must name candidate columns and bounded source quotes or explicit user declarations.")
        rows = []
        for row in reviewed_rows:
            if isinstance(row, dict):
                if set(row) != set(table["columns"]):
                    raise DocumentError("Every reviewed row must contain exactly the candidate columns.")
                row = [row[column] for column in table["columns"]]
            if not isinstance(row, (list, tuple)) or len(row) != len(table["columns"]):
                raise DocumentError("Reviewed row width must equal the candidate column count.")
            if any(value is not None and (isinstance(value, (dict, list, bool)) or len(str(value)) > 20000) for value in row):
                raise DocumentError("Reviewed cells must be bounded scalar values.")
            rows.append([None if value is None or str(value).strip() == "" else str(value).strip() for value in row])
        review = {"source_id": source_id, "table_id": table_id, "raw_sha256": manifest["raw_sha256"],
                  "rows": rows, "unit_evidence": unit_evidence, "verification": "explicit_user_cell_and_unit_review"}
        review["review_sha256"] = hashlib.sha256(_canonical(review)).hexdigest()
        _write_json(self._directory(source_id) / (table_id + "_review.json"), review)
        return {"status": "ok", "source_id": source_id, "table_id": table_id, "row_count": len(rows),
                "review_sha256": review["review_sha256"], "verification": review["verification"]}

    def _publication_key(self, arguments):
        content = {"workspace_id": self.workspace_id, **arguments}
        source_id, table_id = arguments.get("source_id", ""), arguments.get("table_id", "")
        if re.fullmatch(r"source_[a-f0-9]{64}", source_id) and re.fullmatch(r"table_(?:\d{3}|p\d{6}(?:_text)?_\d{3}|c[a-f0-9]{20})", table_id):
            review = self._directory(source_id) / (table_id + "_review.json")
            if review.exists():
                content["review_sha256"] = json.loads(review.read_text())["review_sha256"]
        return hashlib.sha256(_canonical(content)).hexdigest()

    @staticmethod
    def _normalize_source_date(value, frequency, source_format):
        """Relabel a declared native period boundary without aggregating observations."""
        if source_format not in {"native", "iso_period_start", "iso_period_end"}:
            raise DocumentError("source_date_format must be native, iso_period_start or iso_period_end.", "INVALID_SOURCE_DATE_FORMAT")
        coarse = frequency in {"monthly", "quarterly", "annual"}
        if source_format != "native" and not coarse:
            raise DocumentError("ISO period-boundary normalization applies only to declared monthly, quarterly or annual data.", "INVALID_SOURCE_DATE_FORMAT")
        if value is None:
            return value
        is_iso = bool(re.fullmatch(r"\d{4}-\d{2}-\d{2}", value))
        if source_format == "native":
            if coarse and is_iso:
                raise DocumentError("Source dates use YYYY-MM-DD for declared " + frequency +
                    " data. Set source_date_format=iso_period_start or iso_period_end only if these dates label native period boundaries; this never aggregates daily observations.",
                    "SOURCE_DATE_FORMAT_REQUIRED")
            return value
        if not is_iso:
            raise DocumentError("The declared source_date_format requires every nonmissing date to use YYYY-MM-DD.", "INVALID_SOURCE_DATE_FORMAT")
        try:
            observed = date.fromisoformat(value)
        except ValueError as exc:
            raise DocumentError("Source contains an invalid calendar date.", "INVALID_SOURCE_DATE_FORMAT") from exc
        if frequency == "monthly":
            first = date(observed.year, observed.month, 1)
            last = date(observed.year, observed.month, calendar.monthrange(observed.year, observed.month)[1])
            label = f"{observed.year:04d}-{observed.month:02d}"
        elif frequency == "quarterly":
            quarter = (observed.month - 1) // 3 + 1
            first_month, last_month = (quarter - 1) * 3 + 1, quarter * 3
            first = date(observed.year, first_month, 1)
            last = date(observed.year, last_month, calendar.monthrange(observed.year, last_month)[1])
            label = f"{observed.year:04d}-Q{quarter}"
        else:
            first, last = date(observed.year, 1, 1), date(observed.year, 12, 31)
            label = f"{observed.year:04d}"
        expected = first if source_format == "iso_period_start" else last
        if observed != expected:
            raise DocumentError(f"Source date {value} is not the declared {source_format} boundary for {frequency} data; daily observations cannot be converted by relabeling.",
                                "NON_BOUNDARY_SOURCE_DATE")
        return label

    @staticmethod
    def _merge_numeric_arguments(contract, number_format=None, negative_format=None, null_values=None):
        resolved = dict(contract)
        aliases = {key: value for key, value in {"number_format": number_format, "negative_format": negative_format,
                                               "null_values": null_values}.items() if value is not None}
        for key, value in aliases.items():
            if key in resolved and resolved[key] != value:
                raise DocumentError(f"Top-level {key} conflicts with contract.{key}. Supply one consistent explicit source parsing rule.", "CONFLICTING_NUMERIC_FORMAT")
            resolved[key] = value
        return resolved, aliases

    def publish_selected_table(self, source_id, table_id, contract, expected_version, column_mapping=None, unit_evidence=None,
                               number_format=None, negative_format=None, null_values=None):
        contract, parsing_aliases = self._merge_numeric_arguments(contract, number_format, negative_format, null_values)
        contract = self.store._validate_contract(contract)
        manifest = self.source(source_id)
        self.inspect_source(source_id=source_id)
        inspection = json.loads((self._directory(source_id) / "inspection.json").read_text())
        table = next((item for item in inspection["tables"] if item["table_id"] == table_id), None)
        if table is None:
            raise DocumentError("Unknown candidate table.")
        table = self._reviewed_table(source_id, table)
        if table["origin"] == "ocr" and "review" not in table:
            raise DocumentError("OCR candidates need independent cell verification before publication.", "OCR_REVIEW_REQUIRED")
        if table.get("missing_formula_cache") and "review" not in table:
            raise DocumentError("Formula cells have no saved values. Review the source cells or upload a workbook with saved calculated values.", "FORMULA_VALUES_REVIEW_REQUIRED")
        if table.get("layout_review_required") and "review" not in table:
            raise DocumentError("Merged data cells require explicit review before publication; values were not copied across numeric observations.", "TABLE_LAYOUT_REVIEW_REQUIRED")
        arguments = {"source_id": source_id, "table_id": table_id, "contract": contract,
                     "expected_version": expected_version, "column_mapping": column_mapping, "unit_evidence": unit_evidence}
        key = self._publication_key(arguments)
        mapping = column_mapping if column_mapping is not None else {column: column for column in table["columns"]}
        source_columns, output_columns = set(table["columns"]), set(contract.get("columns", {}))
        valid_mapping = isinstance(mapping, dict) and all(isinstance(value, str) for value in mapping.values())
        forward = valid_mapping and set(mapping) == source_columns and len(set(mapping.values())) == len(mapping) and set(mapping.values()) == output_columns
        inverse = valid_mapping and set(mapping) == output_columns and len(set(mapping.values())) == len(mapping) and set(mapping.values()) == source_columns
        mapping_direction = "source_to_output"
        if not forward and inverse:
            mapping = {source: output for output, source in mapping.items()}
            mapping_direction = "unambiguous_output_to_source_normalized"
        elif not forward:
            mapped_sources = set(mapping) if valid_mapping else set()
            mapped_outputs = set(mapping.values()) if valid_mapping else set()
            diagnostic = {"source_columns": table["columns"], "contract_columns": list(contract.get("columns", {})),
                          "unmapped_source_columns": sorted(source_columns - mapped_sources),
                          "unmapped_output_columns": sorted(output_columns - mapped_outputs)}
            raise DocumentError("column_mapping must map every actual source column to exactly one output contract column: {source_column: output_column}. " + json.dumps(diagnostic, ensure_ascii=False) +
                ". Renaming cannot select rows, omit columns or turn metric rows into columns. First use prepare_source_table to select rows/columns, join fragmented labels or unpivot dates. Omit column_mapping only when contract names exactly equal candidate names.", "INVALID_COLUMN_MAPPING")
        contract = json.loads(_canonical(contract))
        if set(contract.get("columns", {})) != set(mapping.values()):
            raise DocumentError("Contract columns must exactly match the mapped candidate columns.")
        if any(contract["columns"][mapping[column]]["dtype"] != "string" or contract["columns"][mapping[column]]["kind"] != "dimension"
               for column in table.get("dimension_only_columns", [])):
            raise DocumentError("Joined source labels must remain string dimensions and cannot be published as computed numeric metrics.", "INVALID_LABEL_SEMANTICS")
        reviewed_evidence = {mapping[column]: quote for column, quote in {
            **table.get("verified_unit_evidence", {}), **table.get("review", {}).get("unit_evidence", {})}.items() if column in mapping}
        evidence = unit_evidence or reviewed_evidence
        normalization = {}
        for column, spec in contract["columns"].items():
            try:
                spec, proof = normalize_column(spec)
            except ValueError as exc:
                raise DocumentError(str(exc), "INVALID_UNIT_SEMANTICS") from exc
            contract["columns"][column] = spec
            if proof:
                normalization[column] = proof
            if spec.get("dtype") not in {"integer", "float"}:
                continue
            quote = evidence.get(column, "")
            unit = spec.get("unit", "")
            original_column = next(name for name, mapped in mapping.items() if mapped == column)
            local_header = table["original_columns"].get(original_column, "") + "\n" + table.get("source_header_quotes", {}).get(original_column, "")
            source_headers = [local_header, *table.get("unit_contexts", {}).get(original_column, [])]
            if any(unit_header_conflict(unit, spec.get("scale", 1), header) for header in source_headers):
                raise DocumentError("Declared unit or scale contradicts this numeric column's own source header: " + column,
                                    "COLUMN_UNIT_CONFLICT")
            explicitly_reviewed = quote == reviewed_evidence.get(column) and column in reviewed_evidence
            source_text = "\n".join(source_headers) + "\n" + table.get("unit_caption", "") + "\n" + _unit_caption(table.get("context_text", ""))
            semantic_evidence = cumulative_evidence({"source_semantics": source_text})
            if semantic_evidence and spec.get("kind") in {"flow", "count_flow", "unknown"}:
                spec["temporal_semantics"] = "year_to_date_flow" if re.search(r"ytd|year[ _-]to[ _-]date", source_text, re.I) else "cumulative_flow"
                spec["source_semantics"] = semantic_evidence["source_semantics"]
                spec["aggregation"] = "none"
            basis_text = _search_text(source_text)
            if unit in {"TRY", "USD", "EUR", "GBP"} and re.search(r"purchasing power|(?<!not )inflation[- ]adjusted|constant prices|satin alma gucu|sabit fiyat", basis_text):
                price_basis = spec.get("price_basis")
                if not isinstance(price_basis, str) or _search_text(price_basis).strip() in {"", "nominal", "current prices", "current_prices", "cari", "cari fiyatlar"}:
                    raise DocumentError("Source amounts already have a purchasing-power/constant-price basis. Preserve an explicit price_basis from the source caption: " + source_text[:1000], "PRICE_BASIS_REVIEW_REQUIRED")
                spec["price_basis_evidence"] = source_text
            if not isinstance(quote, str) or not quote.strip() or not (quote in source_text or explicitly_reviewed) or not unit_quote_matches(unit, spec.get("scale", 1), quote):
                original_column = next(name for name, mapped in mapping.items() if mapped == column)
                original_header = table["original_columns"].get(original_column, "")
                raise DocumentError("Numeric unit needs an exact source quote consistent with the declared unit: " + column +
                    ". Copy only a verbatim source substring, without explanatory additions. Candidate header: " + repr(original_header), "UNITS_REVIEW_REQUIRED")
        number_format = contract.get("number_format", "decimal_dot")
        if number_format not in {"decimal_dot", "decimal_dot_grouped", "decimal_comma"}:
            raise DocumentError("Choose decimal_dot, decimal_dot_grouped or decimal_comma numeric parsing explicitly.")
        negative_format = contract.get("negative_format", "leading_minus")
        if negative_format not in {"leading_minus", "accounting_parentheses"}:
            raise DocumentError("Choose leading_minus or accounting_parentheses negative parsing explicitly.")
        null_values = contract.get("null_values", [])
        if (not isinstance(null_values, list) or len(null_values) > 10
                or any(not isinstance(value, str) or not 1 <= len(value) <= 30 for value in null_values)):
            raise DocumentError("Null markers must be an explicit bounded list of source strings.")
        source_date_format = contract.get("source_date_format", "native")
        self._normalize_source_date(None, contract["frequency"], source_date_format)
        rows, date_proof = [], []
        for row_index, row in enumerate(table["rows"], 1):
            output = []
            for original, value in zip(table["columns"], row):
                spec = contract["columns"][mapping[original]]
                if value in null_values:
                    value = None
                if mapping[original] == contract.get("date_column"):
                    original_value = value
                    value = self._normalize_source_date(value, contract["frequency"], source_date_format)
                    if value != original_value:
                        date_proof.append({"candidate_row": row_index, "original": original_value, "normalized": value})
                if value is not None and spec.get("dtype") in {"integer", "float"}:
                    if negative_format == "accounting_parentheses" and value.startswith("(") and value.endswith(")"):
                        inner = value[1:-1].strip()
                        if inner.startswith(("+", "-")):
                            raise DocumentError("An accounting negative cannot contain a second sign.", "INVALID_NUMERIC_CELL")
                        value = "-" + inner
                    if number_format == "decimal_dot_grouped":
                        if not re.fullmatch(r"[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?", value):
                            raise DocumentError("Numeric cell does not match decimal_dot_grouped format.", "INVALID_NUMERIC_CELL")
                        value = value.replace(",", "")
                    if number_format == "decimal_comma":
                        if not re.fullmatch(r"[+-]?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?", value):
                            raise DocumentError("Numeric cell does not match decimal_comma format.", "INVALID_NUMERIC_CELL")
                        value = value.replace(".", "").replace(",", ".")
                    try:
                        if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", value) or not Decimal(value).is_finite():
                            raise InvalidOperation
                    except InvalidOperation as exc:
                        raise DocumentError("Numeric cell needs explicit format or correction.", "INVALID_NUMERIC_CELL") from exc
                output.append(value if value is not None else "")
            rows.append(output)
        contract["source_namespace"] = manifest["source_namespace"]
        contract["document_provenance"] = {"source_id": source_id, "table_id": table_id, "raw_sha256": manifest["raw_sha256"],
            "source_url": manifest["source_url"], "page": table["page"], "sheet": table["sheet"],
            "publication_key": key, "unit_evidence": evidence, "numeric_verification": "explicit_format_and_store_contract"}
        contract["document_provenance"]["column_mapping"] = {"declared": column_mapping, "source_to_output": mapping,
            "direction": mapping_direction, "source_values_changed": False}
        contract["document_provenance"]["numeric_parsing"] = {"number_format": number_format, "negative_format": negative_format, "null_values": null_values}
        if parsing_aliases:
            contract["document_provenance"]["numeric_parsing"]["declared_at_top_level"] = parsing_aliases
        if normalization:
            contract["document_provenance"]["unit_normalization"] = normalization
        for detail in ("formula_cells", "source_pages", "row_origins", "combination", "preparation", "cell_origins", "header_rows", "cell_spans", "extraction_artifact_ref", "header_hypothesis", "source_scope_evidence"):
            if detail == "row_origins" and (table.get("preparation") or table.get("review")):
                # A selected/unpivoted/reviewed candidate has its own row set.
                # Its cell addresses plus source recipe retain the parent row
                # identities; the parent's positional row array does not.
                continue
            if table.get(detail):
                contract["document_provenance"][detail] = table[detail]
        if table.get("review"):
            input_origins = [{column: {"review_row": row_index, "candidate_column": column, "review_sha256": table["review"]["review_sha256"]}
                             for column in table["columns"]} for row_index in range(1, len(table["rows"]) + 1)]
        else:
            input_origins = table.get("cell_origins") or [
                {column: {"candidate_row": row_index, "candidate_column": column} for column in table["columns"]}
                for row_index in range(1, len(table["rows"]) + 1)]
        contract["document_provenance"]["cell_origins"] = [
            {mapping[column]: origin[column] for column in table["columns"]} for origin in input_origins]
        if table.get("review"):
            contract["document_provenance"]["human_review"] = table["review"]
        if date_proof:
            contract["document_provenance"]["date_normalization"] = {
                "operation": "normalize_native_period_label", "source_date_format": source_date_format,
                "declared_frequency": contract["frequency"], "date_column": contract["date_column"],
                "aggregation_performed": False, "rows": date_proof}
        fd, name = tempfile.mkstemp(dir=self.root, suffix=".csv")
        try:
            with os.fdopen(fd, "w", encoding="utf-8", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow([mapping[column] for column in table["columns"]])
                writer.writerows(rows)
            workspace = self.store.ingest_csv(self.workspace_id, name, contract, expected_version=expected_version)
        finally:
            Path(name).unlink(missing_ok=True)
        result = self._find_publication(key, workspace)
        _write_json(self.root / ("publication_" + key + ".json"), result)
        return result

    def _find_publication(self, key, workspace=None):
        workspace = workspace or self.store.workspace(self.workspace_id)
        for dataset_id in workspace["datasets"]:
            manifest = self.store.dataset_manifest(dataset_id)
            provenance = manifest.get("contract", {}).get("document_provenance", {})
            if provenance.get("publication_key") == key:
                self.store.overlay_path(dataset_id)
                return {"status": "ok", "dataset_id": dataset_id, "artifact_ref": dataset_id,
                        "source_id": provenance["source_id"], "table_id": provenance["table_id"],
                        "workspace_version": workspace["version"], "row_count": manifest["row_count"],
                        "source_namespace": manifest["contract"]["source_namespace"], "provenance": provenance,
                        "frequency": manifest["contract"]["frequency"],
                        "published_columns": manifest["contract"]["columns"],
                        "metric_ids": [f"overlay:{dataset_id}:{column}" for column, definition in manifest["contract"]["columns"].items()
                                       if definition["dtype"] in {"integer", "float"}],
                        "next_step": ("Use aggregate_dataset on this dataset_id for static/event data with explicit grouping and time_bucket where needed."
                                      if manifest["contract"]["frequency"] in {"static", "event"} else
                                      "Use describe on the returned metric_ids, then execute to align with existing series. discover can search by dataset name. aggregate_dataset is for within-dataset grouping, not joining this source to sector data.")}
        return {"status": "blocked", "code": "WRITE_OUTCOME_UNKNOWN", "message": "No committed publication was found; the write was not repeated."}

    def recover_publication(self, arguments, intent):
        canonical = {"column_mapping": None, "unit_evidence": None, **arguments}
        aliases = {key: canonical.pop(key, None) for key in ("number_format", "negative_format", "null_values")}
        canonical["contract"], _ = self._merge_numeric_arguments(canonical["contract"], **aliases)
        return self._find_publication(self._publication_key(canonical))

    def web_search(self, query, limit=5, *, _deadline=None):
        if self.searxng_url is False:
            return {"status": "unavailable", "code": "WEB_SEARCH_UNCONFIGURED", "results": []}
        if not isinstance(query, str) or not 1 <= len(query) <= 500 or type(limit) is not int or not 1 <= limit <= 10:
            raise DocumentError("Search query and limit exceed their bounds.")
        from agentic_analytics.agent.tools.search_pipeline import run_search
        return run_search(query, limit, configured_url=self.searxng_url, fetch=fetch_public_url, deadline=_deadline)

    def _contract_schema(self):
        """Describe the store contract without restricting its metadata extension fields."""
        nonblank = {"type": "string", "minLength": 1, "pattern": r"\S"}
        identifier = {"type": "string", "pattern": r"^[A-Za-z_][A-Za-z0-9_]{0,127}$"}
        kinds = sorted(KINDS)
        column = {"type": "object", "properties": {
            "dtype": {"type": "string", "enum": ["string", "integer", "float", "boolean", "date"],
                      "description": "Stored type: date for the calendar key, integer/float for numeric metrics, string for labels; boolean cells must be true/false."},
            "unit": {**nonblank, "description": "Explicit measurement unit, preferably base currency TRY/USD/EUR plus scale, or visits/persons/percent/calendar/label. Known aliases such as TRY_million with scale=1 are canonicalized to TRY with scale=1000000 without changing cells. Numeric columns need matching exact unit_evidence."},
            "kind": {"type": "string", "enum": kinds,
                     "description": "stock/count_stock are point-in-time balances; flow/count_flow are amounts/counts over a period; count is a generic count; ratio/rate are ratios/rates; index is an index; price is a unit price; dimension is a key/label. Use unknown when source semantics are not established."},
            "nullable": {"type": "boolean", "description": "Whether empty cells are allowed; key cells can never be null."},
            "scale": {"type": "number", "default": 1, "description": "Number of base units represented by a stored number, e.g. unit TRY with scale 1000000 for million TL. Do not double-apply a unit conversion."},
            "currency": {"type": ["string", "null"], "description": "Currency code for monetary metrics, e.g. TRY or USD; needed for monetary deflation."},
            "aggregation": {"type": "string", "enum": ["none", "last", "sum", "mean"],
                            "description": "Optional temporal aggregation literal: stocks last, flows sum, rates/indices mean where appropriate, otherwise none. Put prose scope explanations in separate metadata; never infer additivity across populations."},
            "index_role": {"type": ["string", "null"], "description": "price_deflator only when source evidence establishes a suitable price index; a production index is not an inflation deflator."},
            "deflator_currency": {"type": ["string", "null"], "description": "Currency whose monetary values this reviewed price deflator applies to."},
            "price_scope": {"type": ["string", "null"], "description": "Documented price scope such as consumer_prices."},
            "price_basis": {"type": ["string", "null"], "description": "For already inflation-adjusted source amounts, state the exact source purchasing-power date/basis. Omit or null for nominal values. These bases must match before comparison; never deflate an already-real source twice."},
            "temporal_semantics": {"type": "string", "description": "Preserve explicit source cumulative/YTD meaning, e.g. year_to_date_flow. A cumulative reported value is not a standalone monthly flow; it remains readable but arithmetic requires reviewed conversion."},
            "source_semantics": {"type": "string", "description": "Exact source wording establishing measurement meaning, population and time basis."},
        }, "required": ["dtype", "unit", "kind", "nullable"], "additionalProperties": True}
        keys = {"type": "array", "items": identifier, "minItems": 1, "uniqueItems": True}
        return {"type": "object", "properties": {
            "name": {**nonblank, "description": "Human-readable dataset name; used in discovered metric titles."},
            "columns": {"type": "object", "propertyNames": identifier, "minProperties": 1,
                        "maxProperties": min(self.max_columns, self.store.limits["max_columns"]), "additionalProperties": column,
                        "description": "Exactly all mapped candidate columns; each definition requires dtype/unit/kind/nullable."},
            "key": {**keys, "description": "Existing columns forming a unique nonmissing row key, including date_column for time data."},
            "grain": {**keys, "description": "Exactly the same set of columns as key, declaring what one row represents."},
            "frequency": {"type": "string", "enum": ["daily", "business_daily", "weekly", "monthly", "quarterly", "annual", "event", "static"],
                          "description": "Native frequency: monthly dates YYYY-MM; quarterly YYYY-Qn; annual YYYY; daily/business_daily/weekly/event YYYY-MM-DD. Static data has no date axis. Storage accepts event/static; regular time-series analytics require a supported calendar frequency."},
            "date_column": {"type": ["string", "null"], "description": "Required for non-static data: an existing dtype=date column included in key and grain. Static data must omit it or use null."},
            "source_date_format": {"type": "string", "enum": ["native", "iso_period_start", "iso_period_end"], "default": "native",
                                   "description": "Use native when dates already have the declared frequency's labels. For monthly/quarterly/annual source rows labeled with YYYY-MM-DD, explicitly choose iso_period_start or iso_period_end only when every date is that native period's first/last calendar day. Values and row counts are preserved; this never resamples daily data. Examples: monthly 2025-01-01 -> 2025-01 with iso_period_start; quarterly 2025-03-31 -> 2025-Q1 with iso_period_end."},
            "expected_rows": {"type": "integer", "minimum": 1, "maximum": self.store.limits["max_rows"], "description": "Optional exact expected row count; use the full candidate count, never the preview count."},
            "expected_periods": {"type": "array", "items": {"type": "string"}, "minItems": 1, "uniqueItems": True,
                                 "description": "Optional exact native period labels required for every entity; not allowed for static data."},
            "number_format": {"type": "string", "enum": ["decimal_dot", "decimal_dot_grouped", "decimal_comma"], "default": "decimal_dot",
                              "description": "Default decimal_dot uses 1234.56 without thousands separators. Explicit decimal_dot_grouped accepts 1,234.56 or 1234.56; decimal_comma accepts 1.234,56 or 1234,56. Never guess a separator convention."},
            "negative_format": {"type": "string", "enum": ["leading_minus", "accounting_parentheses"], "default": "leading_minus",
                                "description": "Explicit accounting_parentheses interprets (1,234) as a negative under the chosen number_format; formulas are never executed."},
            "null_values": {"type": "array", "items": {"type": "string", "minLength": 1, "maxLength": 30}, "maxItems": 10,
                            "description": "Only explicitly declared source strings are treated as missing, e.g. ['-','n/a']. Never equate these with zero; nullable and key checks still apply."},
        }, "required": ["name", "columns", "key", "grain", "frequency"], "additionalProperties": True,
            "allOf": [{"if": {"properties": {"frequency": {"const": "static"}}, "required": ["frequency"]},
                       "then": {"properties": {"date_column": {"type": "null"}}, "not": {"required": ["expected_periods"]}},
                       "else": {"required": ["date_column"], "properties": {"date_column": identifier}}}]}

    def extra_tools(self):
        import jsonschema
        definitions = {
            "inspect_source": (self.inspect_source, "Inspect an uploaded source ID or public URL; returns untrusted source text and candidate tables.",
                {"source_id": {"type": "string"}, "url": {"type": "string"},
                 "table_strategy": {"type": "string", "enum": ["lines", "text"], "default": "lines",
                                    "description": "PDF table extraction: lines uses drawn cell borders; text uses aligned words when a report has no cell borders or the default merges many numbers. Inspect a small page selection with text to obtain alternate candidates; original IDs remain unchanged."},
                 "page_numbers": {"type": "array", "items": {"type": "integer", "minimum": 1}, "minItems": 1,
                                  "maxItems": self.max_pages, "uniqueItems": True,
                                  "description": "Optional exact 1-based PDF pages. Default inspects the first bounded page batch, not the whole long report. Reuse source_id to inspect later pages; existing table IDs remain stable."}}, []),
            "publish_selected_table": (self.publish_selected_table, "Publish one inspected table with explicit dtype/unit/kind/grain contract and source unit quotes. Found valid web tables via research_web/inspect_source can also be published as an overlay dataset for auto-ingestion. Example monthly columns: month={dtype:date,unit:calendar,kind:dimension,nullable:false}, visits={dtype:integer,unit:visits,kind:count_flow,nullable:false}; name=clinic_visits, key=[month], grain=[month], date_column=month, frequency=monthly. Use actual candidate column names or explicit column_mapping, not these example names when different.",
                {"source_id": {"type": "string"}, "table_id": {"type": "string"}, "contract": self._contract_schema(),
                 "expected_version": {"type": "integer", "minimum": 0}, "column_mapping": {"type": "object",
                    "description": "Direction is SOURCE COLUMN -> OUTPUT CONTRACT COLUMN. Example actual candidate columns ['raw_month','raw_amount'] and contract columns ['month','amount'] require {'raw_month':'month','raw_amount':'amount'}. Cover every candidate column exactly once; mappings cannot select rows, drop columns or reshape a table. Omit only if candidate and contract column names already match.",
                    "additionalProperties": {"type": "string"}},
                 "number_format": {"type": "string", "enum": ["decimal_dot", "decimal_dot_grouped", "decimal_comma"],
                                   "description": "Optional alias for contract.number_format. Use decimal_dot_grouped for 1,234.56 or decimal_comma for 1.234,56. An explicit conflicting rule in contract is rejected."},
                 "negative_format": {"type": "string", "enum": ["leading_minus", "accounting_parentheses"],
                                     "description": "Optional alias for contract.negative_format, with identical conflict checks."},
                 "null_values": {"type": "array", "items": {"type": "string", "minLength": 1, "maxLength": 30}, "maxItems": 10,
                                 "description": "Optional alias for contract.null_values. These explicit source markers become missing values, never zeros."},
                 "unit_evidence": {"type": "object", "description": "Keys are mapped numeric column names; values must be exact source substrings such as staff (persons), copied from source_header_quotes or source text. Do not add a prefix, translation or explanation. Trusted user-reviewed unit declarations are reused automatically when omitted.",
                                   "additionalProperties": {"type": "string", "minLength": 1}}}, ["source_id", "table_id", "contract", "expected_version"]),
            "combine_source_tables": (self.combine_source_tables, "Create a new candidate from explicitly selected continuation tables on consecutive PDF pages. Original candidates remain unchanged. Headers must match exactly; verify the same population, period and units from page context, then explain the common scope. This does not publish a dataset or aggregate numbers.",
                {"source_id": {"type": "string"}, "table_ids": {"type": "array", "items": {"type": "string"}, "minItems": 2, "maxItems": 10, "uniqueItems": True},
                 "reason": {"type": "string", "minLength": 10, "maxLength": 1000}}, ["source_id", "table_ids", "reason"]),
            "prepare_source_table": (self.prepare_source_table, "Prepare an inspected table by selecting actual 1-based data rows/columns and optionally unpivoting period columns into a long table. For a report with metrics as rows and years as columns, select the relevant metric rows and unpivot year columns. Values are copied exactly from source cells, period labels from original headers; no numeric inputs or expressions allowed. Publish the new candidate with explicit units and grain after preparation.",
                {"source_id": {"type": "string"}, "table_id": {"type": "string"},
                 "selected_rows": {"type": "array", "items": {"type": "integer", "minimum": 1}, "minItems": 1, "maxItems": self.max_rows, "uniqueItems": True},
                 "selected_columns": {"type": "array", "items": {"type": "string"}, "minItems": 1, "maxItems": self.max_columns, "uniqueItems": True},
                 "unpivot": {"type": "object", "properties": {"columns": {"type": "array", "items": {"type": "string"}, "minItems": 1, "uniqueItems": True},
                              "period_column": {"type": "string"}, "value_column": {"type": "string"},
                              "header_row": {"type": "integer", "minimum": 1, "description": "Optional actual candidate row containing period labels when the extracted first header row was a report title. Must be outside selected_rows."},
                              "period_format": {"type": "string", "enum": ["source_header", "dmy", "parenthesized_dmy", "english_dmy", "year_month_slash", "year_quarter_space"], "default": "source_header",
                                                "description": "Explicit date-label parsing: dmy=30/06/2026, parenthesized_dmy=(30/06/2026), english_dmy=31 March 2026, year_month_slash=2026/6, year_quarter_space=2026 Q2. Output ISO day/month/quarter labels; numerical values are untouched."},
                              "period_sources": {"type": "object", "maxProperties": self.max_columns,
                                  "description": "When merged period headings are in different cells from numeric Total columns, map each value column to actual header cells. Date cells need not appear in selected_columns. If both dates share split cells, concatenate their actual source parts and use date_index=0 or 1 with period_format to select the actual first/second complete date. Never provide a literal date. Example amount_current -> {row:10,columns:[header_left,header_right],separator:' ',date_index:0}.",
                                  "additionalProperties": {"type": "object", "properties": {"row": {"type": "integer", "minimum": 1},
                                      "columns": {"type": "array", "items": {"type": "string"}, "minItems": 1, "maxItems": 4, "uniqueItems": True},
                                      "separator": {"type": "string", "enum": ["", " "]},
                                      "date_index": {"type": "integer", "minimum": 0, "maximum": 3,
                                                     "description": "Optional zero-based occurrence of a complete source date after joining the selected header cells. Requires explicit period_format (for example english_dmy). The selected literal date and all cell addresses are retained as evidence."}}, "required": ["row", "columns", "separator"], "additionalProperties": False}}},
                             "required": ["columns", "period_column", "value_column"], "additionalProperties": False},
                 "join_columns": {"type": "object", "properties": {"columns": {"type": "array", "items": {"type": "string"}, "minItems": 2, "maxItems": 10, "uniqueItems": True},
                                  "output": {"type": "string"}, "separator": {"type": "string", "enum": ["", " ", " / "]}},
                                  "required": ["columns", "output", "separator"], "additionalProperties": False,
                                  "description": "Optional reconstruction of fragmented source label columns after unpivot. Joined output is forced to remain a string dimension, never a new number."}}, ["source_id", "table_id"]),
            "read_source_table": (self.read_source_table, "Read actual candidate rows with stable 1-based row numbers before selecting financial report lines or a period header row. Supports bounded pagination beyond the initial eight-row preview.",
                {"source_id": {"type": "string"}, "table_id": {"type": "string"}, "row_start": {"type": "integer", "minimum": 1},
                 "row_limit": {"type": "integer", "minimum": 1, "maximum": 100}}, ["source_id", "table_id"]),
            "web_search": (self.web_search, "Find public sources with Bing RSS or configured SearXNG. Inspect result URLs before using them as citation evidence.",
                {"query": {"type": "string", "maxLength": 500}, "limit": {"type": "integer", "minimum": 1, "maximum": 10}}, ["query"]),
            "research_web": (self.research_web, "Search public web sources, read a bounded number of result URLs, and return source-grounded content with titles, dates and links. Do not dump raw text; synthesize the findings in complete, grammatically correct sentences or bullet points. Use this for current reports, official announcements and news; do not rely on search snippets alone.",
                {"query": {"type": "string", "maxLength": 500}, "limit": {"type": "integer", "minimum": 1, "maximum": 3},
                 "domains": {"type": "array", "maxItems": 5, "items": {"type": "string", "minLength": 1}}}, ["query"]),
        }
        registry = {}
        for name, (function, description, properties, required) in definitions.items():
            parameters = {"type": "object", "properties": properties, "required": required, "additionalProperties": False}
            validator = jsonschema.Draft202012Validator(parameters)
            def handler(arguments, function=function, validator=validator):
                try:
                    validator.validate(arguments)
                    return function(**arguments)
                except jsonschema.ValidationError as exc:
                    return {"status": "blocked", "code": "INVALID_ARGUMENTS", "message": exc.message}
                except ImportError:
                    return {"status": "blocked", "code": "DEPENDENCY_UNAVAILABLE", "message": "Required document parser is unavailable."}
                except (DocumentError, StoreError, ValueError, TypeError, KeyError, OSError, csv.Error, zipfile.BadZipFile) as exc:
                    recovery, message = getattr(exc, "recovery", None), str(exc)
                    if (recovery is None and getattr(exc, "code", None) in {"INVALID_UNPIVOT", "INVALID_SOURCE_DATE_FORMAT"}
                            and function.__name__ == "prepare_source_table" and isinstance(arguments.get("unpivot"), dict)):
                        try:
                            candidate = self.review_candidate(arguments["source_id"], arguments["table_id"])
                            recovery = self._period_header_recovery(candidate, arguments["unpivot"])
                            if recovery["suggested_unpivot_update"]:
                                message += " Retry with source-derived unpivot update: " + json.dumps(recovery["suggested_unpivot_update"])
                        except (DocumentError, ValueError, TypeError, KeyError):
                            pass
                    return {"status": "blocked", "code": getattr(exc, "code", "DOCUMENT_ERROR"), "message": message,
                            **({"recovery": recovery} if recovery else {}),
                            **({"artifact_ref": exc.artifact_ref} if getattr(exc, "artifact_ref", None) else {})}
            registry[name] = {"schema": {"type": "function", "function": {"name": name, "description": description,
                "parameters": parameters}},
                "handler": handler, "mutating": name == "publish_selected_table"}
        registry["publish_selected_table"]["recover"] = self.recover_publication
        return registry
